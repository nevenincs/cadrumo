"""Historical deduction groups retain live, independent workbook formulas."""

from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from cadrumo.adapters.outbound.workbook.calc_sheets_xlsx import materialize_export_plan
from cadrumo.application.storage.calc_sheets.engine import build_export_plan
from cadrumo.application.storage.calc_sheets.form_workbook import add_form_workbook
from cadrumo.application.storage.calc_sheets.records import OperatorInput, OperatorInputs, TabName

from ..workbook_demo import DEMO_CASES, DemoCase, build_demonstration_plan, demonstration_inputs, demonstration_snapshot

pytestmark = [pytest.mark.integration, pytest.mark.hex_application]


@pytest.mark.parametrize("year", (2022, 2023, 2024, 2025))
def test_full_deduction_total_and_regime_result_keep_all_printed_inputs_live(year: int) -> None:
    snapshot = demonstration_snapshot(DemoCase("390", str(year), None, None, filing_year=year, period="0A"))
    owners = {p.box_number: p.casilla_id for p in snapshot.revision.form_layouts[0].placements if p.box_number}
    components = ("49", "513", "51", "521", "53", "55", "57", "59", "598", "61", "661", "62", "652", "63", "522")
    casillas = {c.id: c for c in snapshot.revision.casillas}
    formulas = {f.id: f for f in snapshot.revision.formulas}
    pending = [owners[number] for number in components]
    leaves = set()
    visited = set()
    while pending:
        owner = pending.pop()
        if owner in visited:
            continue
        visited.add(owner)
        formula_id = casillas[owner].formula
        if formula_id is None:
            leaves.add(owner)
            continue
        expressions = [formulas[formula_id].expression]
        while expressions:
            expression = expressions.pop()
            if expression.casilla_id:
                pending.append(expression.casilla_id)
            expressions.extend(expression.args)
    inputs = OperatorInputs(
        values=tuple(
            OperatorInput(casilla_id=owner, value=Decimal(index * 100)) for index, owner in enumerate(sorted(leaves))
        )
    )
    plan = add_form_workbook(build_export_plan(snapshot, operator_inputs=inputs), snapshot)
    source_cells = {c.casilla_id: c for c in plan.value_cells if c.role == "operator_input"}
    form_cells = {c.casilla_id: c for c in plan.formula_cells if c.address.tab is TabName.FORM}
    for last_rate, subtotal in (
        ("606", "49"),
        ("610", "513"),
        ("614", "51"),
        ("618", "521"),
        ("622", "53"),
        ("626", "55"),
        ("630", "57"),
        ("634", "59"),
        ("638", "598"),
    ):
        detail_address = form_cells[owners[last_rate]].address
        total_address = form_cells[owners[subtotal]].address
        assert total_address.row == detail_address.row + 1
        assert total_address.column == detail_address.column
    adjustment_rows = [form_cells[owners[number]].address.row for number in ("61", "661", "62", "652")]
    assert adjustment_rows == list(range(adjustment_rows[0], adjustment_rows[0] + 4))
    calculated = {c.casilla_id: c for c in plan.formula_cells if c.address.tab is TabName.CALCULOS}
    total = calculated[owners["64"]]
    result = calculated[owners["65"]]
    assert owners["64"] not in source_cells and owners["65"] not in source_cells
    assert '"Sin dato"' in total.formula and '"Sin dato"' in result.formula
    assert owners["49"] not in source_cells
    assert all(
        source_cells[owner].address.qualified() in calculated[owners["49"]].formula
        for owner in (
            "iva.anual.deducible.interiores-corrientes.soportado.cuota",
            "iva.anual.autorepercutido.interior.deducible.cuota",
        )
    )
    for number in components:
        source = calculated[owners[number]] if owners[number] in calculated else source_cells[owners[number]]
        assert source.address.qualified() in total.formula
    assert total.address.qualified() in result.formula
    assert calculated[owners["47"]].address.qualified() in result.formula
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    for cell in (total, result):
        assert workbook[cell.address.tab.value][cell.address.a1].value == "=" + cell.formula
        mirror = next(
            c for c in plan.formula_cells if c.casilla_id == cell.casilla_id and c.address.tab is TabName.FORM
        )
        assert cell.address.qualified() in mirror.formula


@pytest.mark.parametrize("year", (2022, 2023, 2024, 2025))
def test_surcharge_form_uses_rate_inputs_and_total_ignores_aggregate_aliases(year: int) -> None:
    snapshot = demonstration_snapshot(DemoCase("390", str(year), None, None, filing_year=year, period="0A"))
    rates = ("0", "0-26", "0-5", "0-62", "1", "1-4", "5-2", "1-75")
    if year == 2022:
        rates = ("0-5", "1-4", "5-2", "1-75")
    elif year == 2023:
        rates = ("0", "0-5", "0-62", "1-4", "5-2", "1-75")
    prefix = "iva.anual.repercutido.recargo."
    amounts = {f"{prefix}tipo-{rate}.cuota": Decimal(index) for index, rate in enumerate(rates)}
    amounts.update({prefix + suffix: Decimal("9999") for suffix in ("general", "reducido", "super-reducido")})
    inputs = OperatorInputs(
        values=tuple(OperatorInput(casilla_id=owner, value=value) for owner, value in amounts.items())
    )
    plan = add_form_workbook(build_export_plan(snapshot, operator_inputs=inputs), snapshot)
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    source_cells = {c.casilla_id: c for c in plan.value_cells if c.role == "operator_input"}
    total = next(
        c
        for c in plan.formula_cells
        if c.casilla_id == "iva.anual.cuota-devengada-total" and c.address.tab is TabName.CALCULOS
    )
    assert '"Sin dato"' in total.formula
    auxiliary_heading = next(
        c.address.row
        for c in plan.value_cells
        if c.address.tab is TabName.FORM and c.value == "Datos auxiliares del cálculo"
    )
    for rate in rates:
        owner = f"{prefix}tipo-{rate}.cuota"
        source = source_cells[owner]
        assert source.address.qualified() in total.formula
        form_cells = [c for c in plan.formula_cells if c.casilla_id == owner and c.address.tab is TabName.FORM]
        assert len(form_cells) == 1 and source.address.qualified() in form_cells[0].formula
        assert form_cells[0].address.row < auxiliary_heading
        assert workbook[source.address.tab.value][source.address.a1].value == amounts[owner]
        assert workbook[TabName.FORM.value][form_cells[0].address.a1].value == "=" + form_cells[0].formula
    for suffix in ("general", "reducido", "super-reducido"):
        owner = prefix + suffix
        assert source_cells[owner].address.qualified() not in total.formula
        # Auxiliary calculation inputs remain reviewable after the official pages.
        cells = [c for c in plan.formula_cells if c.casilla_id == owner and c.address.tab is TabName.FORM]
        assert cells and all(c.address.row > auxiliary_heading for c in cells)


@pytest.mark.parametrize("year", (2024, 2025))
def test_historical_demo_seeds_only_the_declared_deduction_scenario(year: int) -> None:
    case = next(c for c in DEMO_CASES if c.modelo == "390" and c.filing_year == year)
    snapshot, plan = build_demonstration_plan(case)
    inputs, bindings = demonstration_inputs(snapshot, case)
    assert bindings == {} and len(inputs.values) == 27
    assert all("sector-diferenciado-" in value.casilla_id for value in inputs.values)
    for group in range(1, 4):
        members = [v for v in inputs.values if f"sector-diferenciado-{group}." in v.casilla_id]
        assert len(members) == 9
        assert {v.value for v in members} == {Decimal(group * 100)}
        assert all(not v.casilla_id.endswith(".suma-deducciones") for v in members)
    assert plan.guide and "siguen sin datos" in plan.guide.paragraphs[0]
    assert "48.000" not in plan.guide.paragraphs[0]


@pytest.mark.parametrize("year", (2022, 2023, 2024, 2025))
def test_deduction_totals_compile_as_three_independent_live_formulas(year: int) -> None:
    snapshot = demonstration_snapshot(DemoCase("390", str(year), None, None, filing_year=year, period="0A"))
    by_number = {c.number: c for c in snapshot.revision.casillas}
    groups = {
        "155": ("140", "142", "144", "146", "148", "150", "152", "153", "154"),
        "172": ("157", "159", "161", "163", "165", "167", "169", "170", "171"),
        "189": ("174", "176", "178", "180", "182", "184", "186", "187", "188"),
    }
    inputs = OperatorInputs(
        values=tuple(
            OperatorInput(casilla_id=by_number[number].id, value=Decimal(index * 100))
            for index, numbers in enumerate(groups.values(), 1)
            for number in numbers
        )
    )
    plan = add_form_workbook(build_export_plan(snapshot, operator_inputs=inputs), snapshot)
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    source_cells = {c.casilla_id: c for c in plan.value_cells if c.role == "operator_input"}
    for index, (number, components) in enumerate(groups.items(), 1):
        target = by_number[number].id
        assert target not in source_cells
        calculated = next(c for c in plan.formula_cells if c.casilla_id == target and c.address.tab is TabName.CALCULOS)
        assert '"Sin dato"' in calculated.formula
        for component in components:
            source = source_cells[by_number[component].id]
            assert source.address.qualified() in calculated.formula
            assert workbook[source.address.tab.value][source.address.a1].value == index * 100
        for other, numbers in groups.items():
            if other != number:
                assert all(source_cells[by_number[n].id].address.qualified() not in calculated.formula for n in numbers)
        projected = [c for c in plan.formula_cells if c.casilla_id == target and c.address.tab is TabName.FORM]
        assert len(projected) == 1 and calculated.address.qualified() in projected[0].formula
        for cell in (calculated, projected[0]):
            assert workbook[cell.address.tab.value][cell.address.a1].value == "=" + cell.formula
