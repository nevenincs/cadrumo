"""Each parallel OSS template projects its own declared detail owners."""

import re
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from cadrumo.adapters.outbound.workbook.calc_sheets_xlsx import materialize_export_plan
from cadrumo.application.storage.calc_sheets.layout import plan_layout
from cadrumo.application.storage.calc_sheets.records import TabName
from cadrumo.domain.calculations.registry.schema_form_layouts import FormGridBlock

from ..workbook_demo_oss import build_oss_detail_plan

pytestmark = [pytest.mark.integration, pytest.mark.hex_application]


@pytest.mark.parametrize("scheme", ["exterior", "union", "importacion"])
def test_oss_detail_owners_remain_distinct_and_missing_settlement_is_disclosed(scheme):
    source, plan = build_oss_detail_plan(scheme=scheme)
    layout = plan_layout(source.revision)
    values = {cell.address: cell.value for cell in plan.value_cells}
    grids = {
        block.id: block
        for form in source.revision.form_layouts
        for page in form.pages
        for section in page.sections
        for block in section.blocks
        if isinstance(block, FormGridBlock) and block.id != "resultado-estados"
    }
    for identifier, grid in grids.items():
        row = dict(zip((column.key for column in grid.columns), grid.rows[0].cells, strict=True))
        amount = row["correccion" if identifier == "correcciones" else "cuota"]
        assert amount.binding_id is not None
        address = layout.binding_cells[amount.binding_id]
        expected = (
            Decimal("-20") if identifier == "correcciones" else Decimal("380" if identifier == "union-envio" else "190")
        )
        assert values[address] == expected
        mirrors = [
            cell
            for cell in plan.formula_cells
            if cell.address.tab is TabName.FORM
            and not str(cell.casilla_id).startswith(
                ("iva.exterior.resultado.", "iva.union.resultado.", "iva.importacion.resultado.")
            )
            and re.search(re.escape(address.qualified()) + r"(?![0-9])", cell.formula)
        ]
        assert len(mirrors) == 1
        assert "ISBLANK" in mirrors[0].formula
        unused = grid.rows[1 if scheme == "union" else 2].cells[-1]
        assert unused.binding_id is not None
        assert values[layout.binding_cells[unused.binding_id]] is None
        if scheme != "union":
            second = grid.rows[1].cells[-1]
            assert second.binding_id is not None
            assert values[layout.binding_cells[second.binding_id]] == Decimal(
                "-140" if identifier == "correcciones" else "100"
            )
    text = "\n".join(str(cell.value) for cell in plan.value_cells)
    assert not re.search(r"\b[0-9a-f]{64}\b", text)
    if scheme != "union":
        assert "ingreso de 170 euros" in text
        assert "Falta el ingreso adicional" in text
        assert any(
            cell.casilla_id == f"iva.{scheme}.resultado.saldo-ingresar-espana" and cell.address.tab is TabName.FORM
            for cell in plan.formula_cells
        )
    else:
        assert "ingreso de 930 euros" in text
        assert "Falta el ingreso adicional" in text
        assert any(
            cell.casilla_id == "iva.union.resultado.saldo-ingresar-espana" and cell.address.tab is TabName.FORM
            for cell in plan.formula_cells
        )
    # Google Sheets rejects formulas longer than 50,000 characters.
    assert max(len(cell.formula) for cell in plan.formula_cells) < 50_000
    assert "entradas independientes" in text
    # Missing backend aggregate facts must not acquire the detail example's
    # values merely because they refer to the same destination country.
    for casilla in source.revision.casillas:
        if casilla.input_kind == "bound":
            assert values[layout.entradas_cells[casilla.id]] is None
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)))
    internal = {casilla.id: casilla for casilla in source.revision.casillas if casilla.internal_only}
    hidden = {(row.tab, row.row) for row in plan.hidden_rows}
    assert internal
    for cell in plan.formula_cells:
        if cell.casilla_id in internal:
            assert cell.address.tab is not TabName.FORM
            assert (cell.address.tab, cell.address.row) in hidden
            sheet = workbook[cell.address.tab.value]
            assert sheet.row_dimensions[cell.address.row].hidden
            assert sheet[cell.address.a1].value == "=" + cell.formula
        elif cell.casilla_id and cell.address.tab is TabName.CALCULOS:
            assert (cell.address.tab, cell.address.row) not in hidden
    reference_labels = {cell.value for cell in plan.value_cells if cell.address.tab is TabName.PROVENANCE}
    assert not {casilla.label for casilla in internal.values()} & reference_labels
    form = workbook[TabName.FORM.value]
    section_positions = {
        cell.value: row_number
        for row_number, row in enumerate(form.iter_rows(), start=1)
        for cell in row
        if cell.value
        in (
            "1. Declarante",
            "2. Ejercicio y período",
            "Datos del pago",
            "Datos complementarios de presentación",
        )
    }
    assert len(section_positions) == 4
    assert section_positions["1. Declarante"] < section_positions["2. Ejercicio y período"]
    assert section_positions["2. Ejercicio y período"] < section_positions["Datos del pago"]
    assert section_positions["Datos del pago"] < section_positions["Datos complementarios de presentación"]
    guide = workbook[TabName.GUIDE.value]
    assert tuple(guide.cell(row, 1).value for row in range(7, 11)) == plan.guide.paragraphs
    workbook.close()
