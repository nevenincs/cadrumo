"""Historical human workbooks come from the exact source and shared form renderer."""

from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from cadrumo.adapters.outbound.workbook.calc_sheets_xlsx import materialize_export_plan
from cadrumo.application.storage.calc_sheets.records import OperatorInput, OperatorInputs, SheetGuideContent, TabName

from ..workbook_template_preview import build_fictional_template_workbook
from ..workbook_template_source import load_workbook_template_source

pytestmark = [pytest.mark.integration, pytest.mark.hex_application]


@pytest.mark.parametrize(("casilla_id", "message"), [("invented", "absent"), ("19", "calculated")])
def test_fictional_inputs_cannot_invent_or_replace_calculated_casillas(casilla_id: str, message: str) -> None:
    source = load_workbook_template_source(
        "130",
        revision_id="2019-y-siguientes",
        source_ref="aeat-dr-130-2019-v12",
        preview_year=2020,
        preview_period="4T",
    )
    with pytest.raises(ValueError, match=message):
        build_fictional_template_workbook(
            source,
            operator_inputs=OperatorInputs(values=(OperatorInput(casilla_id=casilla_id, value=Decimal(1)),)),
            guide=SheetGuideContent(title="Ejemplo ficticio", paragraphs=("No válido para presentar.",)),
        )


def test_historical_232_human_form_retains_independent_amounts_and_unknowns() -> None:
    source = load_workbook_template_source(
        "232", revision_id="2016-2017", source_ref="aeat-dr-232-2016", preview_year=2017, preview_period="0A"
    )
    figures = {
        "decl.ejercicio": Decimal(2017),
        "vinculada-1-nif": "B00000000",
        "vinculada-1-razon-social": "Entidad vinculada ficticia A",
        "vinculada-1-fjo": "J",
        "vinculada-1-ingreso-pago": "I",
        "vinculada-1-importe": Decimal(250000),
        "vinculada-2-nif": "B00000001",
        "vinculada-2-razon-social": "Entidad vinculada ficticia B",
        "vinculada-2-fjo": "J",
        "vinculada-2-ingreso-pago": "P",
        "vinculada-2-importe": Decimal(175000),
    }
    plan = build_fictional_template_workbook(
        source,
        operator_inputs=OperatorInputs(values=tuple(OperatorInput(casilla_id=k, value=v) for k, v in figures.items())),
        guide=SheetGuideContent(
            title="Modelo 232 · Ejemplo ficticio 2017",
            paragraphs=("EJEMPLO FICTICIO. Importes declarados independientes; no calcula una cuota tributaria.",),
        ),
    )
    assert plan.human_presentation
    assert plan.metadata.template_digest == source.template_digest
    assert plan.tabs[0] is TabName.FORM
    mirrors = {cell.casilla_id: cell for cell in plan.formula_cells if cell.address.tab is TabName.FORM}
    assert mirrors["vinculada-1-importe"].address.row != mirrors["vinculada-2-importe"].address.row
    text = "\n".join(str(cell.value) for cell in plan.value_cells)
    assert "EJEMPLO FICTICIO" in text
    assert "2017-01-01" in text and "2017-12-31" in text
    assert "Importe sin IVA" in text
    assert source.template_digest not in text
    assert "dr23201" not in text
    book = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    for casilla, value in figures.items():
        cell = next(c for c in plan.value_cells if c.address.tab is TabName.ENTRADAS and c.casilla_id == casilla)
        assert book[cell.address.tab.value].cell(cell.address.row, cell.address.column).value == value
        if casilla in mirrors:
            assert cell.address.qualified() in mirrors[casilla].formula
    missing = next(
        c for c in plan.value_cells if c.casilla_id == "vinculada-3-importe" and c.address.tab is TabName.ENTRADAS
    )
    assert missing.value is None
    assert '"Sin dato"' in mirrors["vinculada-3-importe"].formula
