"""The Modelo 369 wrapper retains its two unusual printed marker forms."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from cadrumo.core.resources.bundled_data import bundled_path
from cadrumo.domain.calculations.registry.errors import RegistryValidationError
from dev.registry.compiler.loader import load_catalogue_file
from dev.registry.compiler.record_design_schema import RecordDesign369RelativeClosing
from dev.registry.compiler.record_design_workbook import extract_sheet
from dev.registry.pipeline.record_design_intermediate import (
    RecordDesignIntermediateCompositeRelativeClosing,
    load_record_design_intermediate,
)

pytestmark = [pytest.mark.unit, pytest.mark.hex_core]

_SOURCE_REF = "aeat-dr-369-2021"
_SOURCE_SHA256 = "b59ade58821e8e0988a1aa4e2a7f52c97b21375fc0a6720d76ca0601a7c8b1a3"
_SOURCE_PATH = bundled_path(
    "corpus",
    "aeat_official",
    "disenos_registro",
    "modelo_369",
    "files",
    "01-369-regimenes-especiales-aplicables-a-los-servicios-prestados-a-personas-que-no-tengan-la-co.xlsx",
)


def test_official_369_wrapper_keeps_relative_positions_and_absent_total_length() -> None:
    catalogues = load_catalogue_file(bundled_path("registry", "aeat", "legal", "iva.toml"))
    intermediate = load_record_design_intermediate(
        bundled_path(), catalogues.sources, source_ref=_SOURCE_REF, filing_year=2021, design_epoch="2021"
    )

    assert intermediate.source.source_sha256 == _SOURCE_SHA256
    assert len(intermediate.sheets) == 13
    assert len(intermediate.variable_envelopes) == 1
    envelope = intermediate.variable_envelopes[0]
    assert envelope.sheet == "T3690 Estruc. gral"
    assert envelope.prefix_extent == 328
    assert envelope.body_offset == 329
    assert isinstance(envelope.closing, RecordDesignIntermediateCompositeRelativeClosing)
    assert tuple(part.offset for part in envelope.closing.parts) == (1, 4, 7, 8, 12, 14)
    assert tuple(part.length for part in envelope.closing.parts) == (3, 3, 1, 4, 2, 5)
    assert envelope.total_source_row == 23
    assert envelope.total_source_cell == "A23"
    assert envelope.total_length is None


def test_369_relative_closing_refuses_a_changed_printed_position() -> None:
    workbook = load_workbook(_SOURCE_PATH, data_only=True)
    sheet = workbook["T3690 Estruc. gral"]
    sheet["B18"] = 5

    with pytest.raises(RegistryValidationError, match="malformed composite relative closing"):
        extract_sheet(sheet)


def test_369_bare_total_refuses_an_unstated_length_value() -> None:
    workbook = load_workbook(_SOURCE_PATH, data_only=True)
    sheet = workbook["T3690 Estruc. gral"]
    sheet["D23"] = "unknown"

    with pytest.raises(RegistryValidationError, match="incomplete variable-envelope composition"):
        extract_sheet(sheet)


def test_official_369_closing_is_source_typed_before_projection() -> None:
    workbook = load_workbook(_SOURCE_PATH, data_only=True)
    parsed = extract_sheet(workbook["T3690 Estruc. gral"])
    assert parsed.variable_envelope is not None
    assert isinstance(parsed.variable_envelope.closing, RecordDesign369RelativeClosing)
