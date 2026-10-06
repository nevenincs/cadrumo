"""Real registry acceptance for the shared form compiler, using fictional data."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from cadrumo.adapters.outbound.workbook.calc_sheets_xlsx import materialize_export_plan
from cadrumo.application.modelo.work_form_context_values import form_context_value
from cadrumo.application.modelo.work_form_records import saved_form_records
from cadrumo.application.storage.calc_sheets.engine import build_export_plan
from cadrumo.application.storage.calc_sheets.form_workbook import add_form_workbook
from cadrumo.application.storage.calc_sheets.layout import plan_layout
from cadrumo.application.storage.calc_sheets.records import OperatorInputs, TabName
from cadrumo.core.authority_grade import RegistryAuthorityGrade
from cadrumo.core.config import override_settings
from cadrumo.core.i18n.render import lookup_translation, tr
from cadrumo.core.period import Period
from cadrumo.domain.calculations.registry.afiliado_contribution_bindings import AfiliadoContributionProvider
from cadrumo.domain.calculations.registry.errors import RegistryValidationError
from cadrumo.domain.calculations.registry.formula_runtime import evaluate_expression
from cadrumo.domain.calculations.registry.formula_runtime_ops import apply_rounding
from cadrumo.domain.calculations.registry.governed_fact_scope import validating_governed_facts
from cadrumo.domain.calculations.registry.schema import RegistrySnapshot
from cadrumo.domain.calculations.registry.schema_form_layouts import FormContextFieldBlock, FormRepeatingGroupBlock
from cadrumo.domain.calculations.registry.schema_formula import FormulaExpression

from .. import workbook_demo
from ..compiler.authority import compiled_bundled_authority
from ..workbook_demo import DEMO_CASES, DemoCase, build_demonstration_plan, demonstration_inputs
from ..workbook_demo_records import (
    annual_capital_records,
    annual_employment_records,
    annual_rent_records,
    financial_asset_records,
    member_attribution_records,
)

pytestmark = [pytest.mark.integration, pytest.mark.hex_application]


@pytest.mark.parametrize("case", [case for case in DEMO_CASES if case.modelo == "194"], ids=lambda case: case.revision)
def test_194_summary_and_form_share_five_transactions(case: DemoCase) -> None:
    snapshot, plan = build_demonstration_plan(case)
    saved = financial_asset_records(snapshot)
    inputs, bindings = demonstration_inputs(snapshot, case)
    values = _numeric_inputs(inputs)
    assert values == {"01": Decimal("2"), "04": Decimal("3")}
    for box, expected in (("02", "125"), ("03", "23.75"), ("05", "50")):
        assert _expected_result(snapshot, values, bindings, box) == Decimal(expected)
    observed: dict[int, dict[str, object]] = {}
    for page in snapshot.revision.form_layouts[0].pages:
        for section in page.sections:
            for block in section.blocks:
                if not isinstance(block, FormRepeatingGroupBlock):
                    continue
                columns = tuple(column.casilla_id for column in block.columns)
                known, rows = saved_form_records(
                    snapshot=snapshot, revision=saved, block=block, column_casillas=columns
                )
                assert known and len(rows) == 5
                for row in rows:
                    for column, value in zip(columns, row.values, strict=True):
                        assert column is not None
                        observed.setdefault(row.index, {})[str(column)] = value
    assert len({row["perc.nif"] for row in observed.values()}) == 1
    bases: list[Decimal] = []
    retentions: list[Decimal] = []
    for row in observed.values():
        base, retention = row["perc.base"], row["perc.retenciones"]
        assert isinstance(base, Decimal) and isinstance(retention, Decimal)
        bases.append(base)
        retentions.append(retention)
    assert sorted(bases) == [
        Decimal("-40"),
        Decimal("-10"),
        Decimal("0"),
        Decimal("25"),
        Decimal("100"),
    ]
    assert sum(retentions) == Decimal("23.75")
    form_values = [cell.value for cell in plan.value_cells if cell.address.tab is TabName.FORM]
    assert "Entidad financiera ficticia Ejemplo" in form_values
    assert "Perceptora ficticia Ejemplo" in form_values
    assert "Presentador ficticio" not in form_values


@pytest.mark.parametrize("case", [case for case in DEMO_CASES if case.modelo == "193"], ids=lambda case: case.revision)
def test_193_expense_rows_do_not_leak_into_recipient_sections(case: DemoCase) -> None:
    snapshot, plan = build_demonstration_plan(case)
    saved = annual_capital_records(snapshot)
    for page in snapshot.revision.form_layouts[0].pages:
        for section in page.sections:
            for block in section.blocks:
                if not isinstance(block, FormRepeatingGroupBlock):
                    continue
                columns = tuple(column.casilla_id for column in block.columns)
                known, rows = saved_form_records(
                    snapshot=snapshot, revision=saved, block=block, column_casillas=columns
                )
                assert known
                expected_rows = [1] if block.export_record_id == "modelo-193-gastos" else [1, 2]
                assert [row.index for row in rows] == expected_rows
                for row in rows:
                    for column, value in zip(columns, row.values, strict=True):
                        if column in {"perc.nif-representante", "gasto.nif-representante"}:
                            assert value is None
                        if column == "gasto.importe":
                            assert value == Decimal("25")
    inputs, bindings = demonstration_inputs(snapshot, case)
    values = _numeric_inputs(inputs)
    assert _expected_result(snapshot, values, bindings, "decl.base-total") == Decimal("3000")
    _, withholding = workbook_demo._scenario_leaves(snapshot, "decl.retenciones-total")
    assert len(withholding) == 1
    changed = {**bindings, next(iter(withholding)): Decimal("600")}
    assert _expected_result(snapshot, values, changed, "decl.retenciones-total") == Decimal("600")
    assert _expected_result(snapshot, values, changed, "decl.base-total") == Decimal("3000")
    form_values = [cell.value for cell in plan.value_cells if cell.address.tab is TabName.FORM]
    assert "Perceptor ficticio A" in form_values and "Perceptor ficticio B" in form_values
    assert "Empresa ficticia Ejemplo" in form_values and "B12345674" in form_values
    assert "00000000T" not in form_values and "Presentador ficticio" not in form_values
    assert "08" in form_values


@pytest.mark.parametrize("case", [case for case in DEMO_CASES if case.modelo == "190"], ids=lambda case: case.revision)
def test_190_sections_keep_recipient_identity_and_unknown_circumstances(case: DemoCase) -> None:
    snapshot, plan = build_demonstration_plan(case)
    revision = annual_employment_records(snapshot)
    sections = snapshot.revision.form_layouts[0].pages[1].sections
    values_by_recipient: dict[int, dict[str, object]] = {}
    for section in sections:
        for block in section.blocks:
            assert isinstance(block, FormRepeatingGroupBlock)
            columns = tuple(column.casilla_id for column in block.columns)
            known, rows = saved_form_records(snapshot=snapshot, revision=revision, block=block, column_casillas=columns)
            assert known and [row.index for row in rows] == [1, 2]
            for row in rows:
                for column, value in zip(columns, row.values, strict=True):
                    assert column is not None
                    values_by_recipient.setdefault(row.index, {})[str(column)] = value
            assert saved_form_records(snapshot=snapshot, revision=None, block=block, column_casillas=columns) == (
                False,
                (),
            )
    assert [
        (v["perc.nombre"], v["perc.provincia"], v["perc.percepcion-dineraria"]) for v in values_by_recipient.values()
    ] == [
        ("Persona ficticia A", "28", Decimal("24000")),
        ("Persona ficticia B", "08", Decimal("18000")),
    ]
    assert all(v["perc.nif-representante-legal"] is None for v in values_by_recipient.values())
    inputs, bindings = demonstration_inputs(snapshot, case)
    assert _expected_result(snapshot, _numeric_inputs(inputs), bindings, "decl.percepciones-total") == Decimal("42000")
    form_values = [c.value for c in plan.value_cells if c.address.tab is TabName.FORM]
    assert "Persona ficticia A" in form_values and "Persona ficticia B" in form_values
    assert "Empresa ficticia Ejemplo" in form_values
    assert "B12345674" in form_values
    assert case.filing_year in form_values
    assert "Presentador ficticio" not in form_values
    assert "00000000T" not in form_values
    assert "08" in form_values
    assert not any("modelo-190-perceptor-row" in str(value) for value in form_values)
    assert plan.metadata.registry_sha not in "\n".join(str(v) for v in form_values)


@pytest.mark.parametrize("case", [case for case in DEMO_CASES if case.modelo == "189"], ids=lambda case: case.revision)
def test_189_valuation_is_not_invented_from_nominal_value(case: DemoCase) -> None:
    snapshot, plan = build_demonstration_plan(case)
    assert not snapshot.revision.formulas
    inputs = {c.casilla_id: c for c in plan.value_cells if c.address.tab is TabName.ENTRADAS}
    mirrors = {c.casilla_id: c for c in plan.formula_cells if c.address.tab is TabName.FORM}
    assert inputs["ejercicio-declaracion"].value == Decimal(case.filing_year)
    assert inputs["codigo-provincia"].value == "08"
    assert inputs["numero-valores"].value == Decimal("10")
    assert inputs["nominal-unitario-valores"].value == Decimal("100")
    assert inputs["valoracion"].value == inputs["valoracion-total"].value == Decimal("1500")
    assert inputs["valoracion"].address.qualified() in mirrors["valoracion"].formula
    assert inputs["nominal-unitario-valores"].address.qualified() not in mirrors["valoracion"].formula
    assert inputs["valoracion"].address.qualified() not in mirrors["valoracion-total"].formula
    for identifier in ("persona-relacion", "numero-justificante", "identificacion-valores", "codigo-pais"):
        assert inputs[identifier].value is None
    assert {c.id for c in snapshot.revision.casillas} <= set(mirrors)
    assert not any(c.address.tab is TabName.CALCULOS for c in plan.formula_cells)


def test_188_recipient_projection_preserves_codes_and_independent_declared_values() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "188")
    snapshot, plan = build_demonstration_plan(case)
    assert not snapshot.revision.formulas
    inputs = {c.casilla_id: c for c in plan.value_cells if c.address.tab is TabName.ENTRADAS}
    mirrors = {c.casilla_id: c for c in plan.formula_cells if c.address.tab is TabName.FORM}
    assert inputs["perceptor.provincia"].value == "08"
    assert inputs["perceptor.porcentaje-retencion"].value == Decimal("19")
    assert inputs["perceptor.representante-nif"].value is None
    assert inputs["perceptor.renta-vitalicia-fecha"].value is None
    assert inputs["decl.ejercicio"].value == Decimal(case.filing_year)
    assert inputs["decl.nombre"].value == "Aseguradora ficticia Ejemplo"
    for identifier in ("numero-justificante", "justificante-anterior", "telefono", "complementaria", "sustitutiva"):
        assert inputs[f"decl.{identifier}"].value is None
        assert inputs[f"decl.{identifier}"].address.qualified() in mirrors[f"decl.{identifier}"].formula
    assert inputs["perceptor.retenciones"].address.qualified() in mirrors["perceptor.retenciones"].formula
    assert inputs["03"].address.qualified() in mirrors["03"].formula
    assert inputs["perceptor.retenciones"].address.qualified() not in mirrors["03"].formula
    assert not any(c.address.tab is TabName.CALCULOS for c in plan.formula_cells)


def test_188_code_constraints_reach_xlsx_even_when_input_is_unknown() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "188")
    _, plan = build_demonstration_plan(case)
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    for identifier, allowed in (
        ("perceptor.modalidad", ("1", "2")),
        ("perceptor.clave", ("B",)),
        ("perceptor.renta-vitalicia-clave", ("A", "B")),
    ):
        constraint = next(c for c in plan.cell_constraints if c.casilla_id == identifier)
        assert constraint.allowed_values == allowed
        validation = next(
            v for v in workbook["Entradas"].data_validations.dataValidation if constraint.address.a1 in v.sqref
        )
        assert validation.type == "custom"
        assert validation.errorStyle == "stop"
        assert constraint.text_validation_formula() == validation.formula1
        assert "valores admitidos" in constraint.grounding_message()


def test_188_empty_annuity_identifier_keeps_its_24_character_limit_in_xlsx() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "188")
    _, plan = build_demonstration_plan(case)
    constraint = next(c for c in plan.cell_constraints if c.casilla_id == "perceptor.renta-vitalicia-identificacion")
    assert constraint.max_length == 24
    assert "24 caracteres" in constraint.grounding_message()
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    validation = next(
        v for v in workbook["Entradas"].data_validations.dataValidation if constraint.address.a1 in v.sqref
    )
    assert workbook["Entradas"][constraint.address.a1].value is None
    assert f"LEN({constraint.address.a1})<=24" in validation.formula1


def test_185_three_months_keep_independent_sources_and_affiliation_zeros() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "185")
    snapshot, plan = build_demonstration_plan(case)
    assert snapshot.filing_year == 2026 and not snapshot.revision.formulas
    inputs = {c.casilla_id: c for c in plan.value_cells if c.address.tab is TabName.ENTRADAS}
    mirrors = {c.casilla_id: c for c in plan.formula_cells if c.address.tab is TabName.FORM}
    for suffix, days in (("", "31"), ("-1", "28"), ("-2", "15")):
        target = f"declarado.dias-alta-mes{suffix}"
        assert inputs[target].value == Decimal(days)
        assert inputs[target].address.qualified() in mirrors[target].formula
    assert inputs["declarado.numero-afiliacion"].value == "000000000001"
    assert inputs["decl.periodo"].value == "03"
    assert inputs["decl.numero-justificante"].value is None
    assert not any(c.address.tab is TabName.CALCULOS for c in plan.formula_cells)


def test_181_compact_date_has_a_human_projection_without_changing_source() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "181")
    snapshot, plan = build_demonstration_plan(case)
    assert not snapshot.revision.formulas
    source = next(
        c for c in plan.value_cells if c.address.tab is TabName.ENTRADAS and c.casilla_id == "fecha-operacion"
    )
    assert source.value == "20240115"
    date_cell = next(
        c for c in plan.formula_cells if c.address.tab is TabName.FORM and c.casilla_id == "fecha-operacion"
    )
    assert source.address.qualified() in date_cell.formula
    assert '"dd/mm/yyyy"' in date_cell.formula and '"yyyymmdd"' in date_cell.formula
    assert "Fecha no válida" in date_cell.formula
    identifier = next(
        c for c in plan.formula_cells if c.address.tab is TabName.FORM and c.casilla_id == "identificacion-operacion"
    )
    assert "DATE(" not in identifier.formula


@pytest.mark.parametrize("case", [case for case in DEMO_CASES if case.modelo == "184"], ids=lambda case: case.revision)
def test_184_member_amounts_are_saved_declarations_not_invented_formulas(case: DemoCase) -> None:
    snapshot, plan = build_demonstration_plan(case)
    assert not snapshot.revision.formulas
    revision = member_attribution_records(snapshot)
    block = snapshot.revision.form_layouts[0].pages[2].sections[0].blocks[0]
    assert isinstance(block, FormRepeatingGroupBlock)
    columns = tuple(column.casilla_id for column in block.columns)
    known, rows = saved_form_records(snapshot=snapshot, revision=revision, block=block, column_casillas=columns)
    assert known and len(rows) == 2
    values = [dict(zip(columns, row.values, strict=True)) for row in rows]
    assert [(v["tipo3.miembro-nombre"], v["tipo3.porcentaje-participacion"], v["tipo3.importe"]) for v in values] == [
        ("Miembro ficticio A", Decimal("60"), Decimal("3600")),
        ("Miembro ficticio B", Decimal("40"), Decimal("2400")),
    ]
    assert values[1]["tipo3.codigo-provincia"] == "08"
    assert all(v["tipo3.clave-pais"] is None for v in values)
    assert all(v["tipo3.referencia-catastral"] is None for v in values)
    form = [cell.value for cell in plan.value_cells if cell.address.tab is TabName.FORM]
    assert "Comunidad ficticia Ejemplo" in form and "E00000000" in form
    assert "Miembro ficticio A" in form and "Miembro ficticio B" in form
    formats = {directive.address: directive.pattern for directive in plan.number_formats}
    amounts = [
        cell
        for cell in plan.value_cells
        if cell.address.tab is TabName.FORM and cell.value in (Decimal("3600"), Decimal("2400"))
    ]
    assert len(amounts) == 2
    assert all(formats[cell.address] == "#,##0.00" for cell in amounts)
    shares = [
        cell
        for cell in plan.value_cells
        if cell.address.tab is TabName.FORM and cell.value in (Decimal("60"), Decimal("40"))
    ]
    assert len(shares) == 2
    assert all(formats[cell.address] == "0.00####" for cell in shares)
    assert not any(c.address.tab is TabName.CALCULOS for c in plan.formula_cells)
    assert any("no cambian al editar la renta" in paragraph for paragraph in plan.guide.paragraphs)
    excluded = {"decl.tipo-soporte", "tipo2.tipo-hoja", "tipo3.tipo-hoja"}
    assert not any(str(c.casilla_id) in excluded for c in (*plan.value_cells, *plan.formula_cells))


@pytest.mark.parametrize(
    "raw, expected", [("0", Decimal("0")), ("-125.50", Decimal("-125.50")), ("NaN", None), ("invalid", None)]
)
def test_184_saved_amount_uses_numeric_casilla_despite_text_wire_slot(raw: str, expected: Decimal | None) -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "184")
    snapshot, _ = build_demonstration_plan(case)
    revision = member_attribution_records(snapshot)
    bindings = {binding: dict(values) for binding, values in revision.row_binding_values.items()}
    bindings["modelo-184-member-row-base-assigned"]["1"] = raw
    revision = revision.model_copy(update={"row_binding_values": bindings})
    block = snapshot.revision.form_layouts[0].pages[2].sections[0].blocks[0]
    assert isinstance(block, FormRepeatingGroupBlock)
    columns = tuple(column.casilla_id for column in block.columns)
    known, rows = saved_form_records(snapshot=snapshot, revision=revision, block=block, column_casillas=columns)
    assert known and len(rows) == 2
    assert rows[0].values[columns.index("tipo3.importe")] == expected
    assert rows[1].values[columns.index("tipo3.importe")] == Decimal("2400")


@pytest.mark.parametrize("case", [case for case in DEMO_CASES if case.modelo == "180"], ids=lambda case: case.revision)
def test_180_saved_recipients_remain_separate_and_missing_fields_unknown(case: DemoCase) -> None:
    snapshot, plan = build_demonstration_plan(case)
    revision = annual_rent_records(snapshot)
    block = snapshot.revision.form_layouts[0].pages[1].sections[0].blocks[0]
    assert isinstance(block, FormRepeatingGroupBlock)
    columns = tuple(column.casilla_id for column in block.columns)
    known, rows = saved_form_records(snapshot=snapshot, revision=revision, block=block, column_casillas=columns)
    assert known and [row.index for row in rows] == [1, 2]
    values = [dict(zip(columns, row.values, strict=True)) for row in rows]
    assert [(row["perc.nombre"], row["perc.base"], row["perc.retenciones"]) for row in values] == [
        ("Arrendador ficticio A", Decimal("12000"), Decimal("2280")),
        ("Arrendador ficticio B", Decimal("6000"), Decimal("1140")),
    ]
    assert all(row["perc.referencia-catastral"] is None for row in values)
    assert all(row["perc.nif-representante-legal"] is None for row in values)
    assert [row["perc.provincia"] for row in values] == ["28", "08"]
    assert [row["perc.inmueble-provincia"] for row in values] == ["28", "08"]
    # Omitting the saved channel must not invent records from aggregate totals.
    assert saved_form_records(snapshot=snapshot, revision=None, block=block, column_casillas=columns) == (False, ())
    form = [cell.value for cell in plan.value_cells if cell.address.tab is TabName.FORM]
    assert "Arrendador ficticio A" in form and "Arrendador ficticio B" in form
    assert "Ana Ejemplo" in form and case.filing_year in form
    assert Decimal("12000") in form and Decimal("6000") in form
    assert "08" in form
    formats = {directive.address: directive.pattern for directive in plan.number_formats}
    money = [
        cell
        for cell in plan.value_cells
        if cell.address.tab is TabName.FORM
        and cell.value in (Decimal("12000"), Decimal("6000"), Decimal("2280"), Decimal("1140"))
    ]
    assert len(money) == 4
    assert all(formats[cell.address] == "#,##0.00" for cell in money)
    codes = [cell for cell in plan.value_cells if cell.address.tab is TabName.FORM and cell.value == "08"]
    assert len(codes) == 2 and all(formats[cell.address] == "@" for cell in codes)
    assert not any("[78]" in str(value) or "[114]" in str(value) for value in form)
    assert not any("modelo-180-perceptor-row" in str(value) for value in form)
    assert any("no modifica ni vuelve a sumar" in paragraph for paragraph in plan.guide.paragraphs)


@pytest.mark.parametrize("modelo", ["117", "126", "128", "222"])
def test_demonstration_does_not_promote_calculation_authority_to_filing(modelo: str) -> None:
    case = next(case for case in DEMO_CASES if case.modelo == modelo)
    assert case.authority_grade is RegistryAuthorityGrade.CALCULATION
    with pytest.raises(RegistryValidationError, match="cannot satisfy the requested 'filing'"):
        compiled_bundled_authority().snapshot(
            modelo, filing_year=case.filing_year, period=case.period, revision_id=case.revision
        )


def test_authored_131_activity_table_has_live_readable_bound_cells() -> None:
    authority = compiled_bundled_authority()
    snapshot = authority.snapshot("131", filing_year=2023, period="4T", revision_id="2019-2023")
    fictional_values = {
        "epigrafe": "659.4",
        "rendimiento-neto": Decimal("20000"),
        "porcentaje": Decimal("2"),
        "resultado": Decimal("400"),
    }
    with override_settings(cadrumo_output_language="es"), validating_governed_facts(authority):
        plan = build_export_plan(snapshot)
        layout = plan_layout(snapshot.revision, bracket_filter_date=date(2023, 12, 31))
        supplied = {
            layout.binding_cells[f"modelo-131.page1.actividad-1-{key}"]: value
            for key, value in fictional_values.items()
        }
        plan = plan.model_copy(
            update={
                "value_cells": tuple(
                    cell.model_copy(update={"value": supplied[cell.address]}) if cell.address in supplied else cell
                    for cell in plan.value_cells
                )
            }
        )
        plan = add_form_workbook(plan, snapshot)
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    form_formulas = [cell for cell in plan.formula_cells if cell.address.tab is TabName.FORM]
    projected = []
    for address, value in supplied.items():
        assert workbook[address.tab.value].cell(address.row, address.column).value == value
        mirrors = [cell for cell in form_formulas if address.qualified() in cell.formula]
        assert len(mirrors) == 1
        projected.append(mirrors[0].address)
    assert len({address.row for address in projected}) == 1
    assert len({address.column for address in projected}) == 4
    text = "\n".join(str(cell.value) for cell in plan.value_cells)
    assert "Actividad 1: epígrafe IAE" in text
    assert "m131-modulos-coeficientes" not in text
    assert plan.metadata.registry_sha not in text


def test_232_populated_related_parties_remain_separate_rows() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "232")
    snapshot, plan = build_demonstration_plan(case)
    assert case.result_casilla is None
    assert not snapshot.revision.formulas
    inputs, _ = demonstration_inputs(snapshot, case)
    values = {item.casilla_id: item.value for item in inputs.values}
    assert values["vinculada-1-importe"] == Decimal("250000")
    assert values["vinculada-2-importe"] == Decimal("175000")
    assert values["vinculada-1-ingreso-pago"] == "I"
    assert values["vinculada-2-ingreso-pago"] == "P"
    assert values["vinculada-1-fjo"] == values["vinculada-2-fjo"] == "J"
    assert "vinculada-3-importe" not in values
    mirrors = {
        str(cell.casilla_id): cell.address
        for cell in plan.formula_cells
        if cell.address.tab is TabName.FORM and cell.casilla_id
    }
    for i in (1, 2):
        assert mirrors[f"vinculada-{i}-nif"].row == mirrors[f"vinculada-{i}-importe"].row
    assert mirrors["vinculada-1-importe"].row != mirrors["vinculada-2-importe"].row
    text = "\n".join(str(cell.value) for cell in plan.value_cells)
    assert "no calcula una cuota tributaria" in text
    assert "Datos declarados" in text
    assert "I: ingreso · P: pago" in text
    assert "Importe sin IVA (€)" in text
    assert "Renta antes de la reducción, sin IVA (€)" in text
    assert "Escenario de cálculo" not in text
    assert "2025-01-01" in text
    assert "2025-12-31" in text
    assert plan.metadata.registry_sha not in text
    assert "dr23201" not in text
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    for casilla, expected in values.items():
        source = next(
            cell for cell in plan.value_cells if cell.casilla_id == casilla and cell.address.tab is TabName.ENTRADAS
        )
        assert workbook[source.address.tab.value].cell(source.address.row, source.address.column).value == expected


@pytest.mark.parametrize(
    "code,expected",
    [
        ("0A", (date(2024, 1, 1), date(2024, 12, 31))),
        ("4T", (date(2024, 10, 1), date(2024, 12, 31))),
        ("02", (date(2024, 2, 1), date(2024, 2, 29))),
        ("1P", (None, None)),
        (None, (None, None)),
    ],
)
def test_232_date_context_uses_bound_period_and_never_guesses_event_dates(code, expected) -> None:
    snapshot = compiled_bundled_authority().snapshot("232", filing_year=2024, period="0A")
    # Exercise the shared date resolver with synthetic alternate periods;
    # this does not admit those periods for a Modelo 232 filing.
    snapshot = snapshot.model_copy(update={"filing_period": Period.from_year_and_code(2024, code) if code else None})
    fields = {
        block.id: block
        for page in snapshot.revision.form_layouts[0].pages
        for section in page.sections
        for block in section.blocks
        if isinstance(block, FormContextFieldBlock)
    }
    assert tuple(form_context_value(snapshot, fields[key]) for key in ("period_start", "period_end")) == expected


def _numeric_inputs(inputs: OperatorInputs) -> dict[str, Decimal]:
    """Pass numeric operands to the arithmetic oracle, excluding typed context."""
    return {item.casilla_id: item.value for item in inputs.values if isinstance(item.value, Decimal)}


def test_122_uses_annual_frame_and_discloses_manual_result() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "122")
    snapshot, plan = build_demonstration_plan(case)
    assert case.period == snapshot.period == "0A"
    assert case.authority_grade is RegistryAuthorityGrade.APPLICABILITY
    result = next(c for c in snapshot.revision.casillas if c.id == "resultado.a-ingresar")
    assert result.formula is None
    inputs, _ = demonstration_inputs(snapshot, case)
    values = _numeric_inputs(inputs)
    assert values["resultado.a-ingresar"] == Decimal("300")
    assert values["familia-numerosa.abono-anticipado"] - values["familia-numerosa.deduccion"] == Decimal("300")
    assert not any(c.casilla_id == result.id and c.address.tab is TabName.CALCULOS for c in plan.formula_cells)
    text = "\n".join(str(c.value) for c in plan.value_cells)
    assert "no se recalcula automáticamente" in text
    assert "diseño de 2018" in text
    assert "12345678Z" in text
    assert "Ana" in text
    assert "Ejemplo" in text
    form_text = "\n".join(str(c.value) for c in plan.value_cells if c.address.tab is TabName.FORM)
    assert "Página 1 · Regularización de deducciones · diseño de 2018" in form_text
    by_id = {
        c.casilla_id: c.address.row
        for c in plan.value_cells
        if c.address.tab is TabName.ENTRADAS and c.role == "operator_input"
    }
    numbers = {
        c.address.row: c.value for c in plan.value_cells if c.address.tab is TabName.ENTRADAS and c.address.column == 2
    }
    assert "decl.pagina-complementaria" not in by_id
    assert "decl.tipo-declaracion" not in by_id
    assert "Indicador de página complementaria" not in text
    assert "Tipo de declaración: clave del resultado" not in text
    assert numbers[by_id["resultado.a-ingresar"]] == "595"
    assert numbers[by_id["familia-numerosa.deduccion"]] == "588"


@pytest.mark.parametrize("case", [case for case in DEMO_CASES if case.modelo == "136"], ids=lambda c: c.revision)
def test_136_context_is_generated_from_inputs_and_payment_remains_unknown(case: DemoCase) -> None:
    snapshot, plan = build_demonstration_plan(case)
    inputs, _ = demonstration_inputs(snapshot, case)
    supplied = {item.casilla_id: item.value for item in inputs.values}
    expected = {
        "declarante-nif": "12345678Z",
        "declarante-nombre": "Ana Ejemplo",
        "provincia-residencia-irpf": "28",
        "fecha-cobro-premio": f"{case.filing_year}-12-20",
        "ejercicio-declaracion": Decimal(case.filing_year),
        "periodo-declaracion": "4T",
        "organizador-denominacion": "Organizador ficticio",
        "organizador-pais": "Francia",
        "organizador-codigo-pais": "FR",
        "loteria-apuesta-denominacion": "Sorteo ficticio de diciembre",
        "fecha-celebracion": f"{case.filing_year}-12-15",
        "precio-unitario": Decimal("20"),
    }
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    for identifier, value in expected.items():
        assert supplied[identifier] == value
        source = next(c for c in plan.value_cells if c.casilla_id == identifier and c.role == "operator_input")
        assert source.value == value
        assert workbook[source.address.tab.value].cell(source.address.row, source.address.column).value == value
        mirrors = [c for c in plan.formula_cells if c.address.tab is TabName.FORM and c.casilla_id == identifier]
        assert len(mirrors) == 1
        assert f"ISBLANK({source.address.qualified()})" in mirrors[0].formula
    for identifier in ("09", "forma-pago", "iban", "justificante-anterior"):
        assert identifier not in supplied
        source = next(c for c in plan.value_cells if c.casilla_id == identifier and c.role == "operator_input")
        assert source.value is None


def test_145_family_cells_reference_distinct_inputs_and_do_not_invent_receipt() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "145")
    snapshot, plan = build_demonstration_plan(case)
    inputs, _ = demonstration_inputs(snapshot, case)
    assert not snapshot.revision.formulas
    assert case.period == "comunicacion"
    assert all(c.casilla_id != "comunicacion.pagina-complementaria" for c in plan.value_cells)
    assert all(c.casilla_id != "comunicacion.pagina-complementaria" for c in plan.formula_cells)
    form_text = [str(c.value) for c in plan.value_cells if c.address.tab is TabName.FORM]
    assert not any("Escenario de cálculo" in value for value in form_text)
    assert not any("comunicacion" in value for value in form_text)
    raw_plan = build_export_plan(snapshot)
    assert raw_plan.guide is not None
    assert tr("application.storage.calc_sheets.engine.guide.pull_command") not in raw_plan.guide.paragraphs
    with override_settings(cadrumo_output_language="en"):
        spanish_form = add_form_workbook(raw_plan, snapshot)
    for key in ("communication_caption", "communication_guide"):
        expected = lookup_translation(f"application.storage.calc_sheets.form.{key}", locale="es")
        assert expected is not None and any(c.value == expected for c in spanish_form.value_cells)
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    for supplied in inputs.values:
        source = next(c for c in plan.value_cells if c.casilla_id == supplied.casilla_id and c.role == "operator_input")
        assert (
            workbook[source.address.tab.value].cell(source.address.row, source.address.column).value == supplied.value
        )
        mirrors = [
            c for c in plan.formula_cells if c.address.tab is TabName.FORM and c.casilla_id == supplied.casilla_id
        ]
        assert len(mirrors) == 1
        assert f"ISBLANK({source.address.qualified()})" in mirrors[0].formula
    for identifier in ("comunicacion.firma-tipo", "acuse-recibo.empresa-entidad", "descendiente-3.anio-nacimiento"):
        source = next(c for c in plan.value_cells if c.casilla_id == identifier and c.role == "operator_input")
        assert source.value is None
    assert plan.guide and "no reproduce íntegramente el formulario vigente" in plan.guide.paragraphs[0]


def test_202_four_rate_rows_feed_the_result_independently() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "202")
    snapshot, plan = build_demonstration_plan(case)
    inputs, bindings = demonstration_inputs(snapshot, case)
    values = _numeric_inputs(inputs)
    for target, expected in (("22", "1000"), ("25", "3000"), ("63", "6000"), ("66", "10000")):
        assert _expected_result(snapshot, values, bindings, target) == Decimal(expected)
        assert len([c for c in plan.formula_cells if c.address.tab is TabName.FORM and c.casilla_id == target]) == 1
    changed = {**values, "64": Decimal("44000")}
    assert _expected_result(snapshot, changed, bindings, "34") == Decimal("20200")
    assert _expected_result(snapshot, changed, bindings, "63") == Decimal("6000")
    assert values["19"] == Decimal("100000")
    assert plan.guide and "cuatro porcentajes son ficticios" in plan.guide.paragraphs[0]


def test_222_preserves_pre_group_adjustments_in_live_result_formulas() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "222")
    snapshot, plan = build_demonstration_plan(case)
    inputs, bindings = demonstration_inputs(snapshot, case)
    values = _numeric_inputs(inputs)
    assert "19" not in values
    assert values["04"] == Decimal("0")
    assert _expected_result(snapshot, values, bindings, "19") == Decimal("100000")
    assert _expected_result(snapshot, values, bindings, "34") == Decimal("19200")
    assert _expected_result(snapshot, {**values, "57": Decimal("300")}, bindings, "34") == Decimal("19500")
    assert _expected_result(snapshot, {**values, "43": Decimal("300")}, bindings, "34") == Decimal("18900")
    assert _expected_result(snapshot, {**values, "64": Decimal("44000")}, bindings, "34") == Decimal("20200")
    for identifier in ("57", "43", "64", "34"):
        assert len([c for c in plan.formula_cells if c.address.tab is TabName.FORM and c.casilla_id == identifier]) == 1
    assert plan.guide and "no aporta un perfil del grupo fiscal" in plan.guide.paragraphs[0]
    assert values["decl.ejercicio"] == Decimal(case.filing_year)
    assert not any(c.casilla_id == "decl.tipo-declaracion" for c in (*plan.value_cells, *plan.formula_cells))
    assert not any("clave del resultado" in str(c.value) for c in plan.value_cells)


@pytest.mark.parametrize(
    "case", [c for c in DEMO_CASES if c.modelo == "222" and c.filing_year < 2025], ids=lambda c: c.revision
)
def test_historical_222_keeps_its_year_rates_and_group_adjustments(case: DemoCase) -> None:
    snapshot, plan = build_demonstration_plan(case)
    inputs, bindings = demonstration_inputs(snapshot, case)
    values = _numeric_inputs(inputs)
    calculated = {str(c.id) for c in snapshot.revision.casillas if c.formula is not None}
    assert not calculated & set(values)
    assert values["04"] == Decimal("30000")
    assert _expected_result(snapshot, values, bindings, "19") == Decimal("30000")
    assert _expected_result(snapshot, values, bindings, "23") == Decimal("20000")
    assert values["decl.ejercicio"] == Decimal(case.filing_year)
    assert _expected_result(snapshot, values, bindings, "34") == Decimal("3200")
    assert _expected_result(snapshot, {**values, "57": Decimal("300")}, bindings, "34") == Decimal("3500")
    assert _expected_result(snapshot, {**values, "43": Decimal("300")}, bindings, "34") == Decimal("2900")
    assert not {"61", "62", "63", "64", "65", "66", "67"} & set(values)
    assert not any(c.casilla_id == "decl.tipo-declaracion" for c in (*plan.value_cells, *plan.formula_cells))
    assert plan.guide and "dos porcentajes son ficticios" in plan.guide.paragraphs[0]


@pytest.mark.parametrize("case", [c for c in DEMO_CASES if c.modelo == "202"], ids=lambda c: c.revision)
def test_202_regime_details_do_not_invent_unavailable_profile_values(case: DemoCase) -> None:
    snapshot, plan = build_demonstration_plan(case)
    section = next(s for s in snapshot.revision.form_layouts[0].pages[0].sections if s.id == "datos-adicionales")
    form_cells = [c for c in plan.value_cells if c.address.tab is TabName.FORM]
    for block in section.blocks:
        assert isinstance(block, FormContextFieldBlock)
        label = lookup_translation(block.heading_key, locale="es")
        assert label is not None
        label_cell = next(c for c in form_cells if c.value == label)
        value_cell = next(c for c in form_cells if c.address.row == label_cell.address.row and c.address.column == 9)
        assert value_cell.value == "Sin dato"
        for locale in ("en", "ca", "hu"):
            assert lookup_translation(block.heading_key, locale=locale) is not None


@pytest.mark.parametrize(
    "case", [c for c in DEMO_CASES if c.modelo == "202" and c.revision != "2025-y-siguientes"], ids=lambda c: c.revision
)
def test_202_historical_rate_rows_do_not_import_later_rates(case: DemoCase) -> None:
    snapshot, plan = build_demonstration_plan(case)
    inputs, bindings = demonstration_inputs(snapshot, case)
    values = _numeric_inputs(inputs)
    assert _expected_result(snapshot, values, bindings, "34") == Decimal("3200")
    assert _expected_result(snapshot, {**values, "23": Decimal("22000")}, bindings, "34") == Decimal("3500")
    assert not {"61", "62", "63", "64", "65", "66", "67"} & {str(c.casilla_id) for c in plan.formula_cells}
    assert plan.guide and "dos porcentajes son ficticios" in plan.guide.paragraphs[0]


def _expected_result(
    snapshot: RegistrySnapshot, values: dict[str, Decimal], bindings: dict[str, Decimal], target: str
) -> Decimal:
    """Check fictional arithmetic through the domain interpreter, not sheet formulas."""
    values = values.copy()
    casillas = {c.id: c for c in snapshot.revision.casillas}
    if any(casillas[key].formula is not None for key in values):
        raise ValueError("oracle inputs cannot override calculated casillas")
    formulas = {f.id: f for f in snapshot.revision.formulas}
    parameters = {p.id: p for p in snapshot.revision.parameters}

    def resolve(identifier: str) -> Decimal:
        if identifier in values:
            return values[identifier]
        formula_id = casillas[identifier].formula
        assert formula_id is not None
        formula = formulas[formula_id]
        visit(formula.expression)
        unresolved: set[str] = set()
        result = evaluate_expression(
            formula.expression,
            values=values,
            binding_values=bindings,
            parameters=parameters,
            date_context={"filing_period": date(snapshot.filing_year, 12, 31)},
            relation_values={},
            unresolved_relation_ids=frozenset(),
            unresolved_casilla_ids=unresolved,
            operand_refs=[],
            operand_casilla_refs=[],
            operand_values=[],
            filing_year=snapshot.filing_year,
        )
        assert not unresolved
        values[identifier] = apply_rounding(result, formula.rounding)
        return values[identifier]

    def visit(expression: FormulaExpression) -> None:
        if expression.casilla_id is not None:
            resolve(expression.casilla_id)
        for argument in expression.args:
            visit(argument)

    return resolve(target)


def test_demo_and_oracle_refuse_calculated_input_overrides(monkeypatch: pytest.MonkeyPatch) -> None:
    case = next(c for c in DEMO_CASES if c.modelo == "222")
    snapshot, _ = build_demonstration_plan(case)
    inputs, bindings = demonstration_inputs(snapshot, case)
    with pytest.raises(ValueError, match="oracle inputs cannot override"):
        _expected_result(snapshot, {**_numeric_inputs(inputs), "19": Decimal("100000")}, bindings, "34")
    leaves, binding_ids = workbook_demo._scenario_leaves(snapshot, "34")
    monkeypatch.setattr(workbook_demo, "_scenario_leaves", lambda *_: (leaves | {"34"}, binding_ids))
    with pytest.raises(ValueError, match="demonstration inputs cannot override"):
        demonstration_inputs(snapshot, case)


@pytest.mark.parametrize("case", DEMO_CASES, ids=lambda case: case.modelo)
def test_current_registry_generates_form_geometry_and_xlsx(case: DemoCase) -> None:
    snapshot, plan = build_demonstration_plan(case)
    assert plan.tabs[0] == TabName.FORM
    assert plan.merged_ranges
    form_text = [str(cell.value) for cell in plan.value_cells if cell.address.tab == TabName.FORM]
    assert any("EJEMPLO FICTICIO" in value for value in form_text)
    assert any("pendiente de revisión" in value for value in form_text)
    for page in snapshot.revision.form_layouts[0].pages:
        heading = lookup_translation(page.heading_key, locale="es") or page.official_heading
        if heading:
            assert any(heading in value for value in form_text)
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    form = workbook[TabName.FORM.value]
    assert len(form.merged_cells.ranges) == len([r for r in plan.merged_ranges if r.tab == TabName.FORM])
    for formula in plan.formula_cells:
        assert workbook[formula.address.tab.value].cell(formula.address.row, formula.address.column).value == (
            "=" + formula.formula
        )
    if case.modelo == "349":
        assert "Empresa ficticia A" in form_text
        assert "Empresa ficticia B" in form_text
        assert "12000" in form_text
        assert "4500" in form_text
        assert len(plan.row_sets) == 2
    elif case.result_casilla is not None:
        inputs, bindings = demonstration_inputs(snapshot, case)
        values = _numeric_inputs(inputs)
        assert not {str(c.id) for c in snapshot.revision.casillas if c.formula is not None} & set(values)
        assert _expected_result(snapshot, values, bindings, case.result_casilla) == case.expected_result


def test_fictional_303_does_not_turn_unrelated_unknown_inputs_into_zero() -> None:
    snapshot, plan = build_demonstration_plan(DEMO_CASES[1])
    inputs, _ = demonstration_inputs(snapshot, DEMO_CASES[1])
    supplied = inputs.by_casilla_id()
    unseeded = [
        cell
        for cell in plan.value_cells
        if cell.role == "operator_input" and cell.casilla_id is not None and cell.casilla_id not in supplied
    ]
    assert unseeded
    assert all(cell.value is None for cell in unseeded)


@pytest.mark.parametrize(
    "case",
    [case for case in DEMO_CASES if case.modelo in {"111", "115", "117", "123", "126", "128", "130", "131", "216"}],
    ids=lambda case: f"{case.modelo}-{case.revision}",
)
def test_fictional_form_projects_only_declared_identity_facts(case: DemoCase) -> None:
    _, plan = build_demonstration_plan(case)
    form_values = [cell.value for cell in plan.value_cells if cell.address.tab is TabName.FORM]
    assert "12345678Z" in form_values
    if case.modelo in {"123", "216"} and case.revision == "2024-y-siguientes":
        assert "Ana Ejemplo" in form_values
    else:
        assert "Ana" in form_values
        assert "Ejemplo" in form_values
    assert "Presentador ficticio" not in form_values
    assert "00000000T" not in form_values
    assert case.filing_year in form_values
    assert case.period in form_values


def test_216_nonwithheld_income_changes_its_totals_without_changing_amount_payable() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "216")
    snapshot, _ = build_demonstration_plan(case)
    inputs, bindings = demonstration_inputs(snapshot, case)
    values = _numeric_inputs(inputs)
    assert _expected_result(snapshot, values, bindings, "16") == Decimal("3")
    assert _expected_result(snapshot, values, bindings, "19") == Decimal("2000")
    values["14"] = Decimal("4")
    values["17"] = Decimal("2500")
    assert _expected_result(snapshot, values, bindings, "16") == Decimal("6")
    assert _expected_result(snapshot, values, bindings, "19") == Decimal("4000")
    assert _expected_result(snapshot, values, bindings, "21") == Decimal("1800")


def test_fictional_130_recalculates_after_an_income_change() -> None:
    snapshot, _ = build_demonstration_plan(DEMO_CASES[0])
    inputs, bindings = demonstration_inputs(snapshot, DEMO_CASES[0])
    values = _numeric_inputs(inputs)
    values["01"] += Decimal("100")
    assert _expected_result(snapshot, values, bindings, "19") == Decimal("6393")


def test_156_calendar_keeps_status_and_money_as_independent_facts() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "156")
    snapshot, plan = build_demonstration_plan(case)
    inputs, bindings = demonstration_inputs(snapshot, case)
    values = {str(value.casilla_id): value.value for value in inputs.values}
    assert bindings == {} and snapshot.revision.formulas == ()
    saved = workbook_demo.demonstration_records(snapshot)
    assert saved is not None and len(saved.detail_rows) == 2
    assert values["numero-total-afiliados"] == Decimal(len(saved.detail_rows))
    assert not any(key.startswith("cotizacion-") for key in values)
    member = snapshot.revision.form_layouts[0].pages[1].sections[0].blocks[0]
    assert isinstance(member, FormRepeatingGroupBlock)
    known, rows = saved_form_records(
        snapshot=snapshot,
        revision=saved,
        block=member,
        column_casillas=tuple(column.casilla_id for column in member.columns),
    )
    assert known and len(rows) == 2
    first, second = [dict(zip((c.casilla_id for c in member.columns), row.values, strict=True)) for row in rows]
    assert first["cotizacion-enero"] == Decimal("100.25")
    assert first["cotizacion-enero-situacion"] == "S"
    assert first["cotizacion-abril"] == Decimal("0")
    assert first["cotizacion-abril-situacion"] == "N"
    assert second["cotizacion-enero"] == Decimal("75.50")
    assert second["cotizacion-diciembre"] is None
    assert second["cotizacion-diciembre-situacion"] is None
    assert "numero-justificante" not in values
    form_values = [cell.value for cell in plan.value_cells if cell.address.tab is TabName.FORM]
    assert values["declarante-razon-social"] == "Mutualidad ficticia Ejemplo"
    assert "Presentador ficticio" not in form_values
    mirrors = {str(cell.casilla_id) for cell in plan.formula_cells if cell.address.tab is TabName.FORM}
    assert "declarante-razon-social" in mirrors
    assert "Ejemplo Prueba Ana" in form_values and "Ejemplo Prueba Luis" in form_values
    assert "001234567890" in form_values and "001234567891" in form_values
    assert form_values.count("Enero") == 2 and form_values.count("Diciembre") == 2
    casillas = {casilla.id: casilla for casilla in snapshot.revision.casillas}
    expected_headings = {
        binding.id: casillas[binding.provider.target_casilla_id].label
        for binding in snapshot.revision.bindings
        if isinstance(binding.provider, AfiliadoContributionProvider)
    }
    actual_headings = {column.binding: column.header_label for row_set in plan.row_sets for column in row_set.columns}
    assert len(expected_headings) == 27
    assert actual_headings == expected_headings


def test_156_amendment_marks_share_one_source_and_preserve_unknown() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "156")
    _snapshot, plan = build_demonstration_plan(case)
    marks = [
        c
        for c in plan.formula_cells
        if c.address.tab is TabName.FORM and c.casilla_id == "declaracion-complementaria-sustitutiva"
    ]
    assert len(marks) == 2
    source = next(
        c
        for c in plan.value_cells
        if c.role == "operator_input" and c.casilla_id == "declaracion-complementaria-sustitutiva"
    )
    assert source.value is None
    for mark in marks:
        assert f"ISBLANK({source.address.qualified()})" in mark.formula
        assert '"Sin dato"' in mark.formula and '"Valor no válido"' in mark.formula
        assert f'EXACT({source.address.qualified()},"C ")' in mark.formula
        assert f'EXACT({source.address.qualified()}," S")' in mark.formula
    book = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    for mark in marks:
        assert book["Modelo"][mark.address.a1].value == "=" + mark.formula


def test_156_receipt_checks_are_derived_from_registered_predicates() -> None:
    case = next(case for case in DEMO_CASES if case.modelo == "156")
    snapshot, plan = build_demonstration_plan(case)
    checks = [c for c in plan.formula_cells if c.address.tab is TabName.FORM and '"Falta dato"' in c.formula]
    assert len(checks) == len(snapshot.revision.verification_predicates) == 3
    inputs = {str(c.casilla_id): c.address.qualified() for c in plan.value_cells if c.role == "operator_input"}
    for check in checks:
        assert check.casilla_id is None
        assert inputs["declaracion-complementaria-sustitutiva"] in check.formula
        assert inputs["numero-identificativo-anterior"] in check.formula
        assert "ISNUMBER(" in check.formula and '"Sin dato"' in check.formula
        assert '"No aplica"' in check.formula
        assert '"Dato aportado"' in check.formula or '"Cero indicado"' in check.formula
    book = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    assert all(book["Modelo"][c.address.a1].value == "=" + c.formula for c in checks)


def test_308_refund_form_compiles_declared_quotas_and_preserves_unknown_branches() -> None:
    case = next(c for c in DEMO_CASES if c.modelo == "308")
    snapshot, plan = build_demonstration_plan(case)
    inputs, bindings = demonstration_inputs(snapshot, case)
    values = _numeric_inputs(inputs)
    target = "decl.req-iva-devolver-17"
    assert target not in values
    assert _expected_result(snapshot, values, bindings, target) == Decimal("250.35")
    assert _expected_result(snapshot, {**values, "decl.req-cuota-16": Decimal("20")}, bindings, target) == Decimal(
        "270.35"
    )
    assert "decl.mtn-iva-devolver-eib" not in values
    assert any(str(cell.casilla_id) == target for cell in plan.formula_cells)
    form_values = [cell.value for cell in plan.value_cells if cell.address.tab is TabName.FORM]
    assert "12345678Z" in form_values
    assert next(v.value for v in inputs.values if v.casilla_id == "decl.nombre") == "Ana"
    assert "IVA a devolver por recargo de equivalencia" in form_values
    assert "Actividades exclusivamente en recargo de equivalencia" in form_values
    assert next(v.value for v in inputs.values if v.casilla_id == "decl.tipo-tributacion") == "2"
    assert not any("MTN -" in str(v) or "REQ -" in str(v) for v in form_values)
    assert plan.metadata.registry_sha not in "\n".join(str(v) for v in form_values)
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    assert any(cell.data_type == "f" for row in workbook["Modelo"] for cell in row)


def test_309_payment_copy_and_liquidation_share_one_live_result() -> None:
    case = next(c for c in DEMO_CASES if c.modelo == "309")
    snapshot, plan = build_demonstration_plan(case)
    inputs, bindings = demonstration_inputs(snapshot, case)
    values = _numeric_inputs(inputs)
    target = "decl.resultado-24"
    assert target not in values
    assert _expected_result(snapshot, values, bindings, target) == Decimal("210")
    assert _expected_result(snapshot, {**values, "decl.rg-cuota-27": Decimal("20")}, bindings, target) == Decimal("230")
    repeated = [cell for cell in plan.formula_cells if cell.address.tab is TabName.FORM and cell.casilla_id == target]
    assert len(repeated) == 2
    assert repeated[0].formula == repeated[1].formula
    form_values = [cell.value for cell in plan.value_cells if cell.address.tab is TabName.FORM]
    assert "12345678Z" in form_values
    assert "00000000T" not in form_values
    assert "Forma de pago" in form_values
    assert "Firma (11)" in form_values
    assert "Persona física no empresario o profesional" in form_values
    assert not any("trigger" in str(value) for value in form_values)
    assert plan.metadata.registry_sha not in "\n".join(str(value) for value in form_values)
    book = load_workbook(BytesIO(materialize_export_plan(plan)), data_only=False)
    assert all(book["Modelo"][cell.address.a1].value == "=" + cell.formula for cell in repeated)


def test_309_historical_example_keeps_seven_rows_and_cannot_acquire_filing_authority() -> None:
    case = next(c for c in DEMO_CASES if c.modelo == "309" and c.revision == "2018-2022")
    snapshot, plan = build_demonstration_plan(case)
    inputs, bindings = demonstration_inputs(snapshot, case)
    values = _numeric_inputs(inputs)
    assert snapshot.revision.authority_grade is RegistryAuthorityGrade.APPLICABILITY
    with pytest.raises(RegistryValidationError, match="authority"):
        compiled_bundled_authority().snapshot("309", filing_year=2022, period="AD-HOC", revision_id="2018-2022")
    assert {value.casilla_id for value in inputs.values}.isdisjoint(
        {"decl.rg-base-25", "decl.rg-tipo-26", "decl.rg-cuota-27"}
    )
    assert _expected_result(snapshot, values, bindings, "decl.resultado-24") == Decimal("210")
    assert _expected_result(
        snapshot, {**values, "decl.rg-cuota-03": Decimal("20")}, bindings, "decl.resultado-24"
    ) == Decimal("230")
    form_values = [cell.value for cell in plan.value_cells if cell.address.tab is TabName.FORM]
    assert "12345678Z" in form_values
    assert "00000000T" not in form_values
    results = [
        cell
        for cell in plan.formula_cells
        if cell.address.tab is TabName.FORM and cell.casilla_id == "decl.resultado-24"
    ]
    assert len(results) == 2 and results[0].formula == results[1].formula
