"""The 360 preview preserves independent declared amounts and missing identity."""

import re
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from cadrumo.adapters.outbound.workbook.calc_sheets_xlsx import materialize_export_plan
from cadrumo.application.storage.calc_sheets.records import TabName

from ..workbook_demo_eu_refund import build_eu_refund_plan

pytestmark = [pytest.mark.integration, pytest.mark.hex_application]


def test_eu_refund_example_keeps_both_invoices_and_honest_absences() -> None:
    source, plan = build_eu_refund_plan()
    inputs = {
        cell.casilla_id: cell.value
        for cell in plan.value_cells
        if cell.address.tab is TabName.ENTRADAS and cell.casilla_id
    }
    assert inputs["devolucion.importe-solicitado"] == Decimal("285")
    assert inputs["decl.op1-cuota"] == Decimal("190")
    assert inputs["decl.op2-cuota"] == Decimal("95")
    assert inputs["decl.solicitante.direccion.codigo-postal"] == "08001"
    assert inputs["decl.bancario-iban"] is None
    assert inputs["decl.bancario-bic"] is None
    assert not {"decl.presentacion-pruebas", "decl.calidad-datos"}.intersection(inputs)
    assert not source.revision.formulas
    assert not any(cell.address.tab is TabName.CALCULOS for cell in plan.formula_cells)
    mirrors = {cell.casilla_id for cell in plan.formula_cells if cell.address.tab is TabName.FORM}
    assert {"devolucion.importe-solicitado", "decl.op1-cuota", "decl.op2-cuota"} <= mirrors
    text = "\n".join(str(cell.value) for cell in plan.value_cells)
    assert not re.search(r"\b[0-9a-f]{64}\b", text)
    assert plan.guide is not None
    assert "entradas independientes" in " ".join(plan.guide.paragraphs)
    guide_text = " ".join(str(cell.value) for cell in plan.value_cells if cell.address.tab is TabName.GUIDE)
    assert "entradas independientes" in guide_text
    assert "datos bancarios obligatorios" in guide_text
    assert "Ejemplo ficticio" in guide_text
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)))
    guide = workbook[TabName.GUIDE.value]
    assert guide.cell(1, 1).value == plan.guide.title
    for row, paragraph in enumerate(plan.guide.paragraphs, 7):
        assert guide.cell(row, 1).value == paragraph
        assert guide.cell(row, 1).alignment.wrap_text
        assert guide.row_dimensions[row].height >= 33
    workbook.close()
