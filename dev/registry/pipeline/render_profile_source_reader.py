"""Canonical binary-backed reader for exact official render-profile evidence."""

from __future__ import annotations

import re
from collections.abc import Callable, Generator
from contextlib import AbstractContextManager, contextmanager
from pathlib import Path
from typing import Final

from cadrumo.core.hashing import sha256_file
from cadrumo.core.link_safety import is_link_like
from cadrumo.domain.calculations.registry.errors import RegistryValidationError

from .render_profile_evidence import (
    OfficialSourceEvidence,
    RenderProfileSourceEvidence,
    RenderProfileSourceEvidenceEntry,
)
from .render_profile_model import RenderProfile


def load_render_profile_source_evidence(
    source_path: Path,
    profile: RenderProfile,
) -> RenderProfileSourceEvidence:
    """Read every claimed official cell from the hash-verified binary itself.

    A profile may rest entirely on reviewed policy decisions, which cite no
    official cell: the design identity is still verified against the exact
    binary, the coverage gate still demands one reviewed rule per eligible
    anchor, and there is then no cell to read -- a PDF-sourced design has no
    Contenido column at all. The workbook requirement applies only when a rule
    DOES claim an official cell, because only a workbook can carry one.
    """
    _require_render_profile_source_binary(source_path, profile)
    official_evidence = _official_profile_source_evidence(profile)
    if not official_evidence:
        return RenderProfileSourceEvidence(design_identity=profile.design_identity, entries=())
    read_cell = _source_evidence_reader(source_path)
    locators = tuple(dict.fromkeys((item.source_sheet, item.source_cell) for item in official_evidence))
    entries = _read_profile_source_entries(read_cell, locators)
    return RenderProfileSourceEvidence(design_identity=profile.design_identity, entries=entries)


def _require_render_profile_source_binary(source_path: Path, profile: RenderProfile) -> None:
    if not source_path.is_file() or is_link_like(source_path):
        raise RegistryValidationError(f"render profile source must be a regular file: {source_path}")
    actual_sha256 = sha256_file(source_path)
    if actual_sha256 != profile.design_identity.source_sha256:
        raise RegistryValidationError(
            "render profile source binary SHA-256 does not match the exact design identity",
        )


def _official_profile_source_evidence(profile: RenderProfile) -> tuple[OfficialSourceEvidence, ...]:
    return tuple(
        evidence
        for evidence in (
            *(rule.evidence for rule in profile.width_17_rules),
            *(rule.evidence for rule in profile.singleton_rules),
            *(rule.evidence for rule in profile.literal_numeric_rules),
        )
        if isinstance(evidence, OfficialSourceEvidence)
    )


def _source_evidence_reader(source_path: Path) -> AbstractContextManager[Callable[[str, str], object]]:
    suffix = source_path.suffix.lower()
    if suffix not in {".xlsx", ".xlsm", ".xls"}:
        raise RegistryValidationError(
            f"render profile source evidence requires a spreadsheet workbook: {source_path}",
        )
    return _legacy_xls_cell_reader(source_path) if suffix == ".xls" else _ooxml_cell_reader(source_path)


def _read_profile_source_entries(
    read_cell: AbstractContextManager[Callable[[str, str], object]],
    locators: tuple[tuple[str, str], ...],
) -> tuple[RenderProfileSourceEvidenceEntry, ...]:
    with read_cell as cell_value:
        entries: list[RenderProfileSourceEvidenceEntry] = []
        for sheet_name, cell in locators:
            value = cell_value(sheet_name, cell)
            normalized = _normalize_source_statement(value)
            if not normalized:
                raise RegistryValidationError(
                    f"render profile evidence locator does not contain source text: {(sheet_name, cell)!r}",
                )
            entries.append(
                RenderProfileSourceEvidenceEntry(
                    sheet=sheet_name,
                    cell=cell,
                    normalized_statement=normalized,
                ),
            )
    return tuple(entries)


_CELL_REFERENCE_RE: Final[re.Pattern[str]] = re.compile(r"^(?P<column>[A-Z]+)(?P<row>[1-9][0-9]*)$")


def _cell_indices(cell: str) -> tuple[int, int]:
    match = _CELL_REFERENCE_RE.fullmatch(cell)
    if match is None:
        raise RegistryValidationError(f"render profile evidence locator is not an A1 cell reference: {cell!r}")
    column = 0
    for character in match.group("column"):
        column = column * 26 + (ord(character) - ord("A") + 1)
    return int(match.group("row")) - 1, column - 1


@contextmanager
def _ooxml_cell_reader(source_path: Path) -> Generator[Callable[[str, str], object]]:
    from openpyxl import load_workbook

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:

        def read(sheet_name: str, cell: str) -> object:
            if sheet_name not in workbook.sheetnames:
                raise RegistryValidationError(f"render profile evidence sheet does not exist: {sheet_name!r}")
            return workbook[sheet_name][cell].value

        yield read
    finally:
        workbook.close()


@contextmanager
def _legacy_xls_cell_reader(source_path: Path) -> Generator[Callable[[str, str], object]]:
    """Read evidence cells out of a legacy binary XLS design."""
    import xlrd

    workbook = xlrd.open_workbook(str(source_path), on_demand=True)
    try:

        def read(sheet_name: str, cell: str) -> object:
            if sheet_name not in workbook.sheet_names():
                raise RegistryValidationError(f"render profile evidence sheet does not exist: {sheet_name!r}")
            sheet = workbook.sheet_by_name(sheet_name)
            row, column = _cell_indices(cell)
            if row >= sheet.nrows or column >= sheet.ncols:
                raise RegistryValidationError(
                    f"render profile evidence locator does not exist: {(sheet_name, cell)!r}",
                )
            return sheet.cell_value(row, column)

        yield read
    finally:
        workbook.release_resources()


def _normalize_source_statement(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())
