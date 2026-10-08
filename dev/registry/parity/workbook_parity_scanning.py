"""Discover, inspect, and classify official workbook artefacts."""

from __future__ import annotations

import re
import time
from collections.abc import Iterable
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING
from zipfile import BadZipFile

from cadrumo.core.decimal.coercion import coerce_decimal
from cadrumo.core.directory_scan import DirectoryEntryKind, scan_directory
from cadrumo.core.external_constants import XLS_EXTENSION as _XLS_EXTENSION
from cadrumo.core.external_constants import XLSX_EXTENSION as _XLSX_EXTENSION
from cadrumo.core.hashing import hash_file as _hash_file
from cadrumo.core.logging import get_logger
from cadrumo.domain.calculations.registry.errors import RegistryValidationError
from cadrumo.domain.calculations.registry.schema_base import EvidenceTier

from .workbook_parity_models import (
    WorkbookArtefactReport,
    WorkbookBackendVerificationReport,
    WorkbookCellRef,
    WorkbookExtension,
    WorkbookModeloCoverage,
)
from .workbook_parity_types import (
    WorkbookKind,
    WorkbookScanStatus,
)

if TYPE_CHECKING:
    # Annotation-only: ``from __future__ import annotations`` above makes every
    # annotation a string, so these need not exist at runtime. openpyxl is one of
    # the heaviest third-party imports in the tree and this module is imported
    # eagerly by the registry facade, so the symbols that ARE needed at runtime
    # (``load_workbook``, ``Tokenizer``, and the ``TokenizerError`` /
    # ``InvalidFileException`` handler types) are imported inside the functions
    # that use them -- a taxpayer calculation must not load a spreadsheet engine.
    from openpyxl.worksheet.worksheet import Worksheet


_log = get_logger(__name__)


_WORKBOOK_SUFFIXES = {_XLSX_EXTENSION, _XLS_EXTENSION}


_MODELO_PATTERN = re.compile(r"(?:^|[\\/])modelo[_-](?P<modelo>\d{3})(?:[\\/]|$)", re.IGNORECASE)


_CELL_REF_PATTERN = re.compile(r"(?<![A-Z0-9_])(?:'[^']+'!)?\$?[A-Z]{1,3}\$?\d+(?![A-Z0-9_])")


_CELL_REF_VALUE_PATTERN = re.compile(r"^(?:(?P<sheet>'[^']+'|[^!]+)!)?(?P<coordinate>\$?[A-Z]{1,3}\$?\d+)$")


@dataclass(frozen=True)
class WorkbookScanOptions:
    """Controls for bounded workbook discovery.

    ``per_file_timeout_seconds`` is a budget of CPU time spent by the scanning
    thread, counted from the moment the workbook is opened. The budget exists to
    stop a runaway parse of a pathological workbook, which is CPU work; wall-clock
    time would also count the time the thread waits for a processor, so the same
    workbook would pass on an idle host and time out on a busy one.
    """

    per_file_timeout_seconds: float = field(
        default=15.0,
    )
    max_formula_refs: int = 500


def discover_workbooks(root: Path) -> tuple[Path, ...]:
    """Return every official workbook artefact below ``root``."""
    resolved = root.resolve()
    if not resolved.exists():
        raise RegistryValidationError(f"workbook root does not exist: {root}")
    return tuple(
        p
        for p in scan_directory(resolved, recursive=True, select=DirectoryEntryKind.FILES)
        if p.suffix.lower() in _WORKBOOK_SUFFIXES
    )


def scan_workbook(path: Path, *, root: Path, options: WorkbookScanOptions | None = None) -> WorkbookArtefactReport:
    """Scan one workbook and classify formula coverage.

    Returns:
        A :class:`WorkbookArtefactReport` describing the workbook's formula coverage.
    """
    # Imported before the ``try`` below: ``InvalidFileException`` is named in an
    # ``except`` clause, which is evaluated as an exception propagates, so it must
    # already be a real object by then -- a TYPE_CHECKING guard would turn a
    # handled parse failure into a NameError on the failure path.
    from openpyxl.utils.exceptions import InvalidFileException

    opts = options or WorkbookScanOptions()
    started = time.monotonic()
    resolved_root = root.resolve()
    resolved_path = path.resolve()
    if resolved_root not in resolved_path.parents and resolved_root != resolved_path:
        raise RegistryValidationError(f"workbook path escapes scan root: {path}")

    relative = resolved_path.relative_to(resolved_root).as_posix()
    digest, byte_count = _hash_file(resolved_path)
    suffix = resolved_path.suffix.lower()
    modelo = _infer_modelo(relative)

    if suffix == _XLS_EXTENSION:
        return _unsupported_binary_xls_report(
            relative=relative,
            modelo=modelo,
            byte_count=byte_count,
            digest=digest,
            started=started,
        )

    try:
        sheets, formulas, references = _scan_xlsx_contents(resolved_path, relative, opts)
    except TimeoutError as exc:
        return _failed_report(
            relative=relative,
            modelo=modelo,
            suffix=_XLSX_EXTENSION,
            byte_count=byte_count,
            digest=digest,
            status=WorkbookScanStatus.TIMEOUT,
            error=str(exc),
            started=started,
        )
    except (InvalidFileException, BadZipFile, OSError) as exc:
        _log.warning(
            "workbook parity scan failed for %s: %s",
            relative,
            exc,
            exc_info=True,
        )
        return _failed_report(
            relative=relative,
            modelo=modelo,
            suffix=_XLSX_EXTENSION,
            byte_count=byte_count,
            digest=digest,
            status=WorkbookScanStatus.FAILED,
            error=f"{type(exc).__name__}: {exc}",
            started=started,
        )
    except Exception as exc:
        # Deliberately broad and deliberately not rendering ``exc``'s own text:
        # this branch is everything openpyxl's read-only XML parser can raise
        # beyond the three types measured positional-only above, and it is not
        # a small closed set -- confirmed empirically, a malformed <row r="..">
        # attribute drives openpyxl's own row-number coercion into a bare
        # ``ValueError`` whose message embeds that raw XML attribute verbatim,
        # while a different malformation (non-well-formed XML) raises a
        # ``ParseError`` carrying only a line/column position. Composing from
        # the exception's type name and the already-known relative path keeps
        # the message useful without gambling on which of those two shapes (or
        # an unenumerated third) fired. ``exc_info=True`` still attaches the
        # full exception to the log record for local diagnosis.
        _log.warning(
            "workbook parity scan unexpected error for %s: %s",
            relative,
            type(exc).__name__,
            exc_info=True,
        )
        raise RegistryValidationError(
            f"Unexpected error scanning workbook {relative}: {type(exc).__name__}",
        ) from exc
    kind = _classify_xlsx(relative, formulas)
    evidence_tier, not_evidence_for = _evidence_for_workbook_kind(kind)
    return WorkbookArtefactReport(
        path=relative,
        modelo=modelo,
        extension=_XLSX_EXTENSION,
        bytes=byte_count,
        sha256=digest,
        sheets=tuple(sheets),
        formula_cells=len(formulas),
        input_candidates=tuple(_dedupe_cells(references)),
        output_candidates=tuple(formulas),
        workbook_kind=kind,
        evidence_tier=evidence_tier,
        not_evidence_for=not_evidence_for,
        scan_status=WorkbookScanStatus.SCANNED,
        elapsed_seconds=_elapsed_decimal(started),
    )


def _unsupported_binary_xls_report(
    *,
    relative: str,
    modelo: str | None,
    byte_count: int,
    digest: str,
    started: float,
) -> WorkbookArtefactReport:
    """Stable .xls short-circuit report — binary XLS requires conversion before scanning."""
    evidence_tier, not_evidence_for = _evidence_for_workbook_kind(WorkbookKind.UNSUPPORTED_BINARY_XLS)
    return WorkbookArtefactReport(
        path=relative,
        modelo=modelo,
        extension=_XLS_EXTENSION,
        bytes=byte_count,
        sha256=digest,
        workbook_kind=WorkbookKind.UNSUPPORTED_BINARY_XLS,
        evidence_tier=evidence_tier,
        not_evidence_for=not_evidence_for,
        scan_status=WorkbookScanStatus.UNSUPPORTED,
        formula_cells=0,
        error="binary XLS requires isolated conversion before workbook formula inspection",
        elapsed_seconds=_elapsed_decimal(started),
    )


def _scan_xlsx_contents(
    resolved_path: Path,
    relative: str,
    opts: WorkbookScanOptions,
) -> tuple[list[str], list[WorkbookCellRef], list[WorkbookCellRef]]:
    """Open the workbook in read-only mode and collect (sheet titles, formulas, references)."""
    from openpyxl import load_workbook

    cpu_started = time.thread_time()
    workbook = load_workbook(resolved_path, data_only=False, read_only=True)
    sheets: list[str] = []
    formulas: list[WorkbookCellRef] = []
    references: list[WorkbookCellRef] = []
    try:
        for worksheet in workbook.worksheets:
            _raise_if_timed_out(cpu_started, opts.per_file_timeout_seconds, relative)
            sheets.append(worksheet.title)
            _scan_worksheet_cells(
                worksheet,
                relative=relative,
                opts=opts,
                cpu_started=cpu_started,
                formulas=formulas,
                references=references,
            )
    finally:
        workbook.close()
    return sheets, formulas, references


def _scan_worksheet_cells(
    worksheet: Worksheet,
    *,
    relative: str,
    opts: WorkbookScanOptions,
    cpu_started: float,
    formulas: list[WorkbookCellRef],
    references: list[WorkbookCellRef],
) -> None:
    """Walk one worksheet's cells, appending formulas and bounded references in place."""
    for row in worksheet.iter_rows(values_only=False):
        _raise_if_timed_out(cpu_started, opts.per_file_timeout_seconds, relative)
        for cell in row:
            value = cell.value
            if not (isinstance(value, str) and value.startswith("=")):
                continue
            ref = WorkbookCellRef(sheet=worksheet.title, coordinate=cell.coordinate, formula=value)
            formulas.append(ref)
            if len(references) < opts.max_formula_refs:
                references.extend(_formula_references(worksheet.title, value, opts.max_formula_refs - len(references)))


def inventory_workbook_coverage(
    root: Path,
    *,
    options: WorkbookScanOptions | None = None,
    limit: int | None = None,
    previous_reports: Iterable[WorkbookArtefactReport] = (),
) -> tuple[WorkbookArtefactReport, ...]:
    """Scan official workbook artefacts and return :class:`WorkbookArtefactReport` coverage records."""
    paths = discover_workbooks(root)
    if limit is not None:
        paths = paths[:limit]
    previous_by_path = {report.path: report for report in previous_reports}
    reports: list[WorkbookArtefactReport] = []
    resolved_root = root.resolve()
    for path in paths:
        relative = path.resolve().relative_to(resolved_root).as_posix()
        previous = previous_by_path.get(relative)
        if previous is not None and previous.sha256 == _hash_file(path)[0]:
            reports.append(previous)
            continue
        reports.append(scan_workbook(path, root=root, options=options))
    return tuple(reports)


def assert_workbook_scan_clean(report: WorkbookBackendVerificationReport) -> None:
    """Raise when discovery could not inspect every workbook artefact."""
    failed_statuses = {WorkbookScanStatus.FAILED, WorkbookScanStatus.TIMEOUT}
    failed = tuple(item for item in report.reports if item.scan_status in failed_statuses)
    if failed:
        details = "\n".join(f" - {item.path}: {item.error}" for item in failed)
        raise RegistryValidationError(f"workbook verification failed to scan {len(failed)} artefact(s):\n{details}")


def _build_modelo_coverage(reports: Iterable[WorkbookArtefactReport]) -> tuple[WorkbookModeloCoverage, ...]:
    buckets: dict[str, list[WorkbookArtefactReport]] = {}
    for report in reports:
        modelo = report.modelo or "unknown"
        buckets.setdefault(modelo, []).append(report)
    return tuple(
        WorkbookModeloCoverage(
            modelo=modelo,
            workbook_count=len(modelo_reports),
            formula_workbook_count=sum(
                1 for report in modelo_reports if report.workbook_kind == WorkbookKind.FORMULA_FORM
            ),
            unsupported_xls_count=sum(
                1 for report in modelo_reports if report.workbook_kind == WorkbookKind.UNSUPPORTED_BINARY_XLS
            ),
            failed_count=sum(
                1
                for report in modelo_reports
                if report.scan_status in {WorkbookScanStatus.FAILED, WorkbookScanStatus.TIMEOUT}
            ),
        )
        for modelo, modelo_reports in sorted(buckets.items())
    )


def _infer_modelo(relative_path: str) -> str | None:
    match = _MODELO_PATTERN.search(relative_path)
    if match is None:
        return None
    modelo = match.group("modelo")
    if not isinstance(modelo, str):
        raise TypeError("the named group always participates in this pattern")
    return modelo


def _raise_if_timed_out(cpu_started: float, timeout_seconds: float, relative: str) -> None:
    """Raise once the scanning thread has spent more than ``timeout_seconds`` of CPU time."""
    if time.thread_time() - cpu_started > timeout_seconds:
        raise TimeoutError(f"workbook scan of {relative!r} exceeded its {timeout_seconds:.1f}s CPU budget")


def _elapsed_decimal(started: float) -> Decimal:
    return coerce_decimal(round(time.monotonic() - started, 6), default=Decimal("0"))


def _classify_xlsx(relative: str, formulas: Iterable[WorkbookCellRef]) -> WorkbookKind:
    formula_count = sum(1 for _ in formulas)
    lowered = relative.lower()
    if "valid" in lowered or "valida" in lowered:
        return WorkbookKind.VALIDATION_HINTS
    if _is_record_design_path(lowered):
        return WorkbookKind.RECORD_DESIGN_LAYOUT
    if formula_count > 0:
        return WorkbookKind.FORMULA_FORM
    return WorkbookKind.STATIC_LAYOUT


def _is_record_design_path(lowered_relative_path: str) -> bool:
    return any(marker in lowered_relative_path for marker in ("disenos_registro", "diseños_registro", "registro"))


def _evidence_for_workbook_kind(kind: WorkbookKind) -> tuple[EvidenceTier | None, tuple[EvidenceTier, ...]]:
    legal, guidance = EvidenceTier.LEGAL_AUTHORITY, EvidenceTier.OFFICIAL_SOURCE_GUIDANCE
    parity, layout = EvidenceTier.EXECUTABLE_PARITY_EVIDENCE, EvidenceTier.LAYOUT_AUTHORITY
    if kind == WorkbookKind.FORMULA_FORM:
        return parity, (legal, layout)
    if kind in {WorkbookKind.RECORD_DESIGN_LAYOUT, WorkbookKind.UNSUPPORTED_BINARY_XLS, WorkbookKind.STATIC_LAYOUT}:
        return layout, (legal, parity)
    if kind == WorkbookKind.VALIDATION_HINTS:
        return guidance, (legal, parity, layout)
    return None, (legal, guidance, parity, layout)


def _formula_references(sheet: str, formula: str, remaining: int) -> tuple[WorkbookCellRef, ...]:
    if remaining <= 0:
        return ()
    # Both imported before the ``try``: ``TokenizerError`` is named in the
    # ``except`` clause and must be a real object when the exception propagates.
    from openpyxl.formula import Tokenizer
    from openpyxl.formula.tokenizer import TokenizerError

    refs: list[WorkbookCellRef] = []
    try:
        tokens = Tokenizer(formula).items
        token_values = (token.value for token in tokens)
    except TokenizerError as exc:
        _log.debug(
            "workbook parity: openpyxl Tokenizer failed on formula %r; falling back to regex (%s)",
            formula[:80],
            exc,
            exc_info=True,
        )
        token_values = (match.group(0) for match in _CELL_REF_PATTERN.finditer(formula))
    for value in token_values:
        for match in _CELL_REF_PATTERN.finditer(value):
            if len(refs) >= remaining:
                return tuple(refs)
            ref_sheet = sheet
            coordinate = match.group(0).replace("$", "")
            if "!" in coordinate:
                raw_sheet, coordinate = coordinate.rsplit("!", 1)
                ref_sheet = raw_sheet.strip("'")
            refs.append(WorkbookCellRef(sheet=ref_sheet, coordinate=coordinate))
    return tuple(refs)


def _dedupe_cells(cells: Iterable[WorkbookCellRef]) -> tuple[WorkbookCellRef, ...]:
    seen: set[tuple[str, str]] = set()
    deduped: list[WorkbookCellRef] = []
    for cell in cells:
        key = (cell.sheet, cell.coordinate)
        if key in seen:
            continue
        seen.add(key)
        deduped.append(cell)
    return tuple(deduped)


def _failed_report(
    *,
    relative: str,
    modelo: str | None,
    suffix: WorkbookExtension,
    byte_count: int,
    digest: str,
    status: WorkbookScanStatus,
    error: str,
    started: float,
) -> WorkbookArtefactReport:
    return WorkbookArtefactReport(
        path=relative,
        modelo=modelo,
        extension=suffix,
        bytes=byte_count,
        sha256=digest,
        workbook_kind=WorkbookKind.UNREADABLE,
        formula_cells=0,
        evidence_tier=None,
        not_evidence_for=(
            EvidenceTier.LEGAL_AUTHORITY,
            EvidenceTier.OFFICIAL_SOURCE_GUIDANCE,
            EvidenceTier.EXECUTABLE_PARITY_EVIDENCE,
            EvidenceTier.LAYOUT_AUTHORITY,
        ),
        scan_status=status,
        error=error,
        elapsed_seconds=_elapsed_decimal(started),
    )
