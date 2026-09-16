"""Legal/source evidence and source-citation validation helpers."""

from __future__ import annotations

import json
import logging
import os
import warnings
from collections.abc import Callable, Iterable, Mapping
from datetime import date
from functools import lru_cache
from pathlib import Path
from typing import TYPE_CHECKING
from zipfile import BadZipFile

from pydantic import ConfigDict, TypeAdapter, ValidationError

from cadrumo.core.atomic_write import atomic_write_best_effort_text
from cadrumo.core.corpus_text import normalise_corpus_text
from cadrumo.core.hashing import sha256_hex
from cadrumo.core.manual_corpus_sidecar import (
    MANUAL_CORPUS_TEXT_CORPUS_PATH_PREFIX,
    MANUAL_CORPUS_TEXT_SIDECAR_SUFFIX,
    ManualCorpusTextSidecar,
)
from cadrumo.core.resources.bundled_data import resolve_companion_binary
from cadrumo.domain.calculations.registry.schema_base import RegistrySourceKind, SourceCitation
from cadrumo.domain.calculations.registry.schema_references import LegalReference, SourceReference

if TYPE_CHECKING:
    import pypdfium2
    from openpyxl.workbook import Workbook

type WorkbookOpener = Callable[[str], Workbook]
"""Opens one XLSX source for read-only cell extraction."""

type PdfDocumentOpener = Callable[[str], pypdfium2.PdfDocument]
"""Opens one PDF source for page-text extraction."""

_SourceTextCacheKey = tuple[str, str, str, str, int, int]
_NORMALISED_SOURCE_TEXT_CACHE: dict[_SourceTextCacheKey, str] = {}
_LOGGER = logging.getLogger(__name__)

CORPUS_TEXT_CACHE_DIR_ENV = "CADRUMO_CORPUS_TEXT_CACHE_DIR"
"""Environment variable that relocates the corpus-text validation cache."""

_CORPUS_TEXT_CACHE_FILENAME = "cadrumo_corpus_text_cache.json"

# Shipped sidecar constants (written by the corpus extraction tooling).
# Sidecars live at _data/manual_corpus_text/<path-relative-to-corpus>.corpus_text.json
# where the path is source.corpus_path with the leading "corpus/" prefix stripped.
# The payload shape itself is the shared ManualCorpusTextSidecar contract in core.
_MANUAL_CORPUS_TEXT_DIR = "manual_corpus_text"
_MAX_XLSX_CITATION_CELLS = 1_000_000
_MAX_XLSX_CITATION_TEXT_CHARS = 16 * 1024 * 1024


def _validated_sidecar_text(raw: str, corpus_path: str, actual_sha256: str) -> str | None:
    """Return the sidecar's text when it satisfies the shared contract, else ``None``.

    The whole admission decision for a manual-PDF sidecar lives here: the
    payload must validate against :class:`ManualCorpusTextSidecar` (pinned
    schema version, prefixed corpus path, hex-64 content key, stamped
    extraction platform, non-empty text), must claim the very
    ``corpus_path`` it was addressed under, and must carry the content key of
    the deployed PDF bytes. Any refusal is logged and returns ``None``, which
    the caller answers by re-extracting from the real bytes -- so a tampered,
    truncated, foreign, or older-schema sidecar costs a slow path, never a
    wrong ``required_text`` verdict.

    Args:
        raw: The sidecar file's decoded JSON text.
        corpus_path: The :attr:`SourceReference.corpus_path` the sidecar was
            addressed under.
        actual_sha256: Hex SHA-256 of the deployed source PDF's bytes.
    """
    try:
        sidecar = ManualCorpusTextSidecar.model_validate_json(raw)
    except ValidationError:
        _LOGGER.warning(
            "Manual PDF sidecar for %s does not satisfy the sidecar contract; falling back to on-demand extraction",
            corpus_path,
            exc_info=True,
        )
        return None
    if not sidecar.addresses(corpus_path):
        _LOGGER.warning(
            "Manual PDF sidecar for %s claims corpus_path %r; falling back to on-demand extraction",
            corpus_path,
            sidecar.corpus_path,
        )
        return None
    if sidecar.source_sha256 != actual_sha256:
        _LOGGER.warning(
            "Manual PDF sidecar sha256 mismatch for %s; falling back to on-demand extraction",
            corpus_path,
        )
        return None
    return sidecar.normalised_text


def _manual_sidecar_path(corpus_path: str, source_root: Path) -> Path:
    relative = corpus_path.removeprefix(MANUAL_CORPUS_TEXT_CORPUS_PATH_PREFIX)
    return source_root / _MANUAL_CORPUS_TEXT_DIR / (relative + MANUAL_CORPUS_TEXT_SIDECAR_SUFFIX)


def _read_manual_pdf_sidecar(
    corpus_path: str,
    source_path: Path,
    *,
    source_root: Path | None = None,
) -> str | None:
    """Return the shipped normalised text for a manual-PDF source, or ``None``.

    Locates the content-keyed sidecar committed under
    ``_data/manual_corpus_text/`` that was built by
    the corpus extraction tooling, then hands the payload to
    :func:`_validated_sidecar_text`, which owns the whole admission decision
    against the shared :class:`ManualCorpusTextSidecar` contract the extractor
    writes through.

    Returns ``None`` when the sidecar is absent (not yet generated or path
    unexpected), unreadable, or refused by that contract. The caller falls
    back to on-demand pypdfium2 extraction on ``None``. End-user machines
    should never reach that path: the shipped sidecar covers every
    ``manual_pdf`` source the registry declares.

    Args:
        corpus_path: The source's :attr:`SourceReference.corpus_path`,
            e.g. ``"corpus/manuals/renta/2020/part1/source.pdf"``.
        source_path: Resolved on-disk path to the source PDF bytes.
        source_root: Candidate evidence root; inferred from corpus_path when omitted.
    """
    if not corpus_path.startswith(MANUAL_CORPUS_TEXT_CORPUS_PATH_PREFIX):
        return None
    if source_root is None:
        source_root = source_path.parents[len(Path(corpus_path).parts) - 1]
    try:
        raw = _manual_sidecar_path(corpus_path, source_root).read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError):
        return None
    return _validated_sidecar_text(raw, corpus_path, sha256_hex(source_path.read_bytes()))


@lru_cache(maxsize=4096)
def _normalise_required_text(text: str) -> str:
    return normalise_corpus_text(text)


def _resolve_source_path(source: SourceReference, source_root: Path) -> Path:
    """Resolve a source path and retain the mirrored-binary fallback."""
    resolved_root = source_root.expanduser().resolve()
    source_path = (resolved_root / source.corpus_path).expanduser().resolve()
    if resolved_root not in source_path.parents and source_path != resolved_root:
        raise OSError(f"source {source.id!r} escapes source root")
    if not source_path.is_file():
        # The command-bearing wheel sheds corpus source binaries; the
        # mandatory cadrumo_data namespace supplies the same bytes at the
        # mirrored relative path, keeping required_text verification
        # byte-identical to a full checkout.
        companion_path = resolve_companion_binary(*source.corpus_path.split("/"))
        if companion_path is not None:
            source_path = companion_path
    return source_path


def _read_source_text(source: SourceReference, source_path: Path, *, source_root: Path | None = None) -> str:
    """Read and normalise a source, using the shipped PDF text when present."""
    if source.kind is RegistrySourceKind.MANUAL_PDF:
        # Try the shipped content-keyed sidecar first; verify sha256 before
        # using it so a modified source PDF never serves stale text.  The
        # sidecar is generated once at build time by the corpus extraction
        # tooling and shipped with the cadrumo wheel — end-user machines
        # should never reach the fallback.
        sidecar_text = _read_manual_pdf_sidecar(source.corpus_path, source_path, source_root=source_root)
        if sidecar_text is not None:
            return sidecar_text
        return normalise_corpus_text(_extract_pdf_text_impl(str(source_path)))
    if source_path.suffix.casefold() == ".xlsx":
        return normalise_corpus_text(_extract_xlsx_text_impl(str(source_path)))
    return normalise_corpus_text(source_path.read_text(encoding="utf-8", errors="replace"))


def _extract_xlsx_text_impl(path: str, *, open_workbook: WorkbookOpener | None = None) -> str:
    """Return cell text from an enrolled XLSX record-design authority.

    Only a workbook the reader itself rejects is relabelled as an unreadable
    XLSX source. The extraction limits below raise ``OSError`` deliberately and
    keep their own message, and every other exception keeps its own subject.

    Args:
        path: Filesystem path of the XLSX source to read.
        open_workbook: Reader override used by the tests to exercise the
            refusal paths; defaults to openpyxl's read-only loader.
    """
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:  # pragma: no cover - dependency is required by pyproject.
        raise OSError("openpyxl is required to validate XLSX source citations") from exc

    def default_opener(source: str) -> Workbook:
        return load_workbook(source, read_only=True, data_only=False)

    opener = open_workbook or default_opener
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = opener(path)
            try:
                cells: list[str] = []
                text_chars = 0
                for worksheet in workbook.worksheets:
                    for row in worksheet.iter_rows(values_only=True):
                        for value in row:
                            if value is None:
                                continue
                            if len(cells) >= _MAX_XLSX_CITATION_CELLS:
                                raise OSError(
                                    f"XLSX source {path} exceeds the citation extraction cell limit "
                                    f"({_MAX_XLSX_CITATION_CELLS})"
                                )
                            rendered = str(value)
                            text_chars += len(rendered) + 1
                            if text_chars > _MAX_XLSX_CITATION_TEXT_CHARS:
                                raise OSError(
                                    f"XLSX source {path} exceeds the citation extraction text limit "
                                    f"({_MAX_XLSX_CITATION_TEXT_CHARS} characters)"
                                )
                            cells.append(rendered)
                return "\n".join(cells)
            finally:
                workbook.close()
    except (InvalidFileException, BadZipFile, KeyError, TypeError, ValueError) as exc:
        raise OSError(f"could not extract text from XLSX source {path}: {exc}") from exc


_disk_cache: dict[str, str] | None = None
_disk_cache_dirty: bool = False
_DISK_CACHE_ADAPTER: TypeAdapter[dict[str, str]] = TypeAdapter(dict[str, str], config=ConfigDict(strict=True))


def corpus_text_cache_dir() -> Path:
    """Resolve the runner-local corpus-text validation cache directory.

    Follows the development cache convention: an explicit
    ``CADRUMO_CORPUS_TEXT_CACHE_DIR`` wins, otherwise ``~/.cadrumo`` holds it.
    The cache lives outside the application's storage root, so validating
    evidence never changes the application state that root fingerprints, and
    it stays scoped per user rather than shared through an OS temp directory.

    Returns:
        The directory holding the corpus-text cache file.
    """
    override = os.environ.get(CORPUS_TEXT_CACHE_DIR_ENV)
    if override:
        return Path(override)
    return Path.home() / ".cadrumo" / "corpus-text"


def _corpus_text_cache_path() -> Path:
    return corpus_text_cache_dir() / _CORPUS_TEXT_CACHE_FILENAME


def flush_corpus_text_cache() -> None:
    """Persist accumulated corpus-text entries in one write.

    Cache misses only mutate the in-process mapping and mark it dirty; the
    validation entry points flush once when they finish. Writing per miss was
    accidentally quadratic: every miss re-read and fully rewrote a JSON file
    that grows to tens of megabytes, which alone cost ~13 seconds of the
    first-touch registry validation on an end-user machine.
    """
    global _disk_cache_dirty
    if not _disk_cache_dirty or _disk_cache is None:
        return
    _write_disk_cache(_disk_cache)
    _disk_cache_dirty = False


def _load_disk_cache() -> dict[str, str]:
    global _disk_cache
    if _disk_cache is not None:
        return _disk_cache
    cache_path = _corpus_text_cache_path()
    if not cache_path.is_file():
        _disk_cache = {}
        return _disk_cache
    try:
        with open(cache_path, encoding="utf-8") as f:
            loaded = _DISK_CACHE_ADAPTER.validate_python(json.load(f))
            _disk_cache = loaded
            return loaded
    except Exception:
        # Degrade to a cache miss (the entries recompute deterministically),
        # but surface the anomaly rather than swallowing it silently.
        _LOGGER.warning("Ignoring unreadable corpus text cache at %s; recomputing", cache_path, exc_info=True)
        _disk_cache = {}
        return _disk_cache


def _write_disk_cache(data: dict[str, str]) -> None:
    cache_path = _corpus_text_cache_path()
    try:
        # Read-merge-before-write narrows the multi-process last-writer-wins
        # window that the atomic replace alone does not close: fold in any
        # entries a concurrent writer committed since our in-memory copy
        # loaded, so a parallel writer's new key is not dropped. The residual
        # race (two writers merging the same pre-image) can only cost a
        # recompute, never a wrong value, because every key embeds the source
        # file's size and mtime -- a stale entry cannot match a changed file.
        merged: dict[str, str] = {}
        if cache_path.is_file():
            try:
                with open(cache_path, encoding="utf-8") as f:
                    on_disk = _DISK_CACHE_ADAPTER.validate_python(json.load(f))
                merged.update(on_disk)
            except Exception:
                _LOGGER.debug("Ignoring unreadable corpus text cache while merging at %s", cache_path, exc_info=True)
        merged.update(data)
        # Compact separators: this is a machine cache that reaches tens of
        # megabytes; indentation only inflates every read and write.
        atomic_write_best_effort_text(
            cache_path,
            json.dumps(merged, ensure_ascii=False, separators=(",", ":")),
            encoding="utf-8",
        )
    except Exception:
        _LOGGER.warning("Could not write corpus text cache at %s", cache_path, exc_info=True)


def _extract_pdf_text_impl(path: str, *, open_document: PdfDocumentOpener | None = None) -> str:
    """Return page text from an enrolled manual PDF authority.

    Only a document pdfium itself rejects, or a filesystem failure reading it,
    is relabelled as an unreadable PDF source; every other exception keeps its
    own subject.

    Args:
        path: Filesystem path of the PDF source to read.
        open_document: Reader override used by the tests to exercise the
            refusal paths; defaults to pdfium's document constructor.
    """
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:  # pragma: no cover - dependency is required by pyproject.
        raise OSError("pypdfium2 is required to validate manual PDF citations") from exc
    opener = open_document or pdfium.PdfDocument
    try:
        pdf = opener(path)
        pages: list[str] = []
        try:
            for index in range(len(pdf)):
                page = pdf[index]
                try:
                    text_page = page.get_textpage()
                    try:
                        pages.append(text_page.get_text_range())
                    finally:
                        text_page.close()
                finally:
                    page.close()
        finally:
            pdf.close()
        return "\n".join(pages)
    except (pdfium.PdfiumError, OSError, ValueError) as exc:
        raise OSError(f"could not extract text from manual PDF {path}: {exc}") from exc


class EvidenceValidator:
    """Validate legal authority, source tiers, and source-citation text evidence."""

    def __init__(
        self,
        *,
        legal_refs: Mapping[str, LegalReference],
        source_refs: Mapping[str, SourceReference],
        source_root: Path | None,
    ) -> None:
        """Bind the candidate catalogues and its evidence root."""
        self._legal = legal_refs
        self._sources = source_refs
        self._source_root = source_root
        self._source_text_cache: dict[str, str] = {}

    def require_source_tier(
        self,
        scope: str,
        owner: str,
        refs: Iterable[str],
        required_tier: str,
    ) -> list[str]:
        """Require a cited source to carry the requested tier."""
        if any(
            (source := self._sources.get(ref)) is not None and source.evidence_tier == required_tier for ref in refs
        ):
            return []
        return [f"{scope}: {owner} requires {required_tier} source evidence"]

    def require_any_source_tier(
        self,
        scope: str,
        owner: str,
        refs: Iterable[str],
        allowed_tiers: Iterable[str],
    ) -> list[str]:
        """Require at least one source in the permitted tiers."""
        allowed = tuple(allowed_tiers)
        if any((source := self._sources.get(ref)) is not None and source.evidence_tier in allowed for ref in refs):
            return []
        if len(allowed) == 1:
            requirement = f"{allowed[0]} source evidence"
        else:
            requirement = f"one of {', '.join(allowed)} source evidence"
        return [f"{scope}: {owner} requires {requirement}"]

    def require_procedural_evidence(
        self,
        scope: str,
        owner: str,
        refs: Iterable[str],
        legal_refs: Iterable[str],
        *,
        valid_from: date,
        valid_to: date | None,
    ) -> list[str]:
        """Accept guidance or a cited, in-window BOE clause in the same source.

        Full compilation separately verifies source hashes and the legal clause's
        required text. A layout source alone cannot establish procedural authority:
        the declaration must cite the actual legal clause in that document.
        """
        refs = tuple(refs)
        failures = self.require_source_tier(scope, owner, refs, "official_source_guidance")
        if not failures:
            return []
        if self.procedural_legal_clauses(refs, legal_refs, valid_from=valid_from, valid_to=valid_to):
            return []
        return failures

    def procedural_legal_clauses(
        self,
        refs: Iterable[str],
        legal_refs: Iterable[str],
        *,
        valid_from: date,
        valid_to: date | None,
    ) -> tuple[LegalReference, ...]:
        """Return cited BOE clauses whose document and governed span match."""
        legal_refs = tuple(legal_refs)
        matched: dict[str, LegalReference] = {}
        for ref in refs:
            source = self._sources.get(ref)
            if source is None or source.authority != "boe" or source.kind != "form_spec":
                continue
            if source.applies_from is None or source.applies_from > valid_from:
                continue
            if source.applies_to is not None and (valid_to is None or source.applies_to < valid_to):
                continue
            for legal_id in legal_refs:
                legal = self._legal.get(legal_id)
                if legal is None or legal.authority != "boe" or not legal.article:
                    continue
                if legal.corpus_ref.partition("#")[0] != source.corpus_path:
                    continue
                governed_from = legal.governs_periods_from or legal.effective_from
                governed_to = legal.governs_periods_to if legal.governs_periods_from else legal.effective_to
                if governed_from > valid_from:
                    continue
                if governed_to is not None and (valid_to is None or governed_to < valid_to):
                    continue
                matched[legal.id] = legal
        return tuple(matched.values())

    def validate_source_citations(
        self,
        scope: str,
        owner: str,
        refs: Iterable[str],
        citations: Iterable[SourceCitation],
        required_tier: str,
    ) -> list[str]:
        """Verify citation membership, tier, and quoted source text."""
        failures: list[str] = []
        refs_set = set(refs)
        citations_tuple = tuple(citations)
        if not citations_tuple:
            return [f"{scope}: {owner} requires source citations"]
        for citation in citations_tuple:
            failures.extend(self._validate_source_citation(scope, owner, refs_set, citation, required_tier))
        return failures

    def _validate_source_citation(
        self,
        scope: str,
        owner: str,
        refs: set[str],
        citation: SourceCitation,
        required_tier: str,
    ) -> list[str]:
        """Validate one citation while retaining the gate's refusal order."""
        if citation.source_ref not in refs:
            return [
                f"{scope}: {owner} source citation {citation.source_ref!r} is not listed in source_refs",
            ]
        source = self._sources.get(citation.source_ref)
        if source is None:
            return []
        if source.evidence_tier != required_tier:
            return [
                f"{scope}: {owner} source citation {citation.source_ref!r} is not {required_tier} evidence",
            ]
        if self._source_root is None:
            return []
        try:
            source_text = self._source_text(source)
        except FileNotFoundError as exc:
            return [f"{scope}: {owner} source citation {citation.source_ref!r} cannot be read: {exc}"]
        except OSError as exc:
            return [f"{scope}: {owner} source citation {citation.source_ref!r} cannot be read: {exc}"]
        return [
            f"{scope}: {owner} source citation {citation.source_ref!r} missing text {required!r}"
            for required in citation.required_text
            if _normalise_required_text(required) not in source_text
        ]

    def source_text(self, source: SourceReference) -> str | None:
        """Return ``source``'s normalised bundled text, or ``None`` when unreachable.

        Public wrapper around :meth:`_source_text` for content-shape checks that
        need the same PDF-sidecar-aware, disk-cached text resolution this class
        already owns -- e.g. a vocabulary probe over a deadline window's
        ``official_source_guidance`` sources -- without duplicating the manual-PDF
        sidecar lookup, the on-disk cache, or the companion-binary fallback path.
        Swallows the read failures :meth:`validate_source_citations` reports as
        failures itself; a caller with no per-citation failure slot to put that in
        treats an unreadable source the same way :func:`validate_layout_authority_content`
        does -- skipped rather than silently accepted.
        """
        if self._source_root is None:
            return None
        try:
            return self._source_text(source)
        except OSError:
            return None

    def _source_text(self, source: SourceReference) -> str:
        cached = self._source_text_cache.get(source.id)
        if cached is not None:
            return cached
        source_root = self._source_root
        if source_root is None:
            return ""
        source_path = _resolve_source_path(source, source_root)
        stat = source_path.stat()
        extraction_contract = f"{source.kind}:xlsx-text-v1" if source_path.suffix.casefold() == ".xlsx" else source.kind
        sidecar_digest = ""
        if source.kind is RegistrySourceKind.MANUAL_PDF:
            sidecar = _manual_sidecar_path(source.corpus_path, source_root)
            if sidecar.is_file():
                sidecar_digest = sha256_hex(sidecar.read_bytes())
        source_key = (
            extraction_contract,
            str(source_path),
            str(source.sha256),
            sidecar_digest,
            stat.st_size,
            stat.st_mtime_ns,
        )
        global_cached = _NORMALISED_SOURCE_TEXT_CACHE.get(source_key)
        if global_cached is not None:
            self._source_text_cache[source.id] = global_cached
            return global_cached

        # Check disk cache
        cache_key_str = json.dumps(source_key)
        disk_cache = _load_disk_cache()
        if cache_key_str in disk_cache:
            normalised = disk_cache[cache_key_str]
            _NORMALISED_SOURCE_TEXT_CACHE[source_key] = normalised
            self._source_text_cache[source.id] = normalised
            return normalised

        normalised = _read_source_text(source, source_path, source_root=source_root)

        _NORMALISED_SOURCE_TEXT_CACHE[source_key] = normalised
        self._source_text_cache[source.id] = normalised
        disk_cache[cache_key_str] = normalised
        global _disk_cache_dirty
        _disk_cache_dirty = True
        return normalised
