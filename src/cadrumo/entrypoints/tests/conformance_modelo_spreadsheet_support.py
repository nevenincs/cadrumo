"""Registered-executor conformance scenarios for the modelo spreadsheet family."""

from __future__ import annotations

import hashlib
import io
import zipfile

from openpyxl import load_workbook

from ...adapters.outbound.google import session_store
from ...adapters.outbound.google.records import DriveConfig
from ...adapters.persistence.storage.tests.profile_capsule_runtime import upsert_test_profile_facts
from ...application.evidence.models import EvidenceBundle
from ...application.evidence.service import EvidenceBundleService
from ...application.modelo.audit_operation import (
    ModeloAuditExportProjection,
    ModeloAuditExportRequest,
    ModeloAuditQueryProjection,
    ModeloAuditReadProjection,
    ModeloAuditReadRequest,
)
from ...application.modelo.modelo_spreadsheet_operation_contracts import (
    MODELO_SPREADSHEET_CALCULATE_OPERATION_DEFINITION_ID,
    MODELO_SPREADSHEET_EXPORT_OPERATION_DEFINITION_ID,
    MODELO_SPREADSHEET_PULL_OPERATION_DEFINITION_ID,
    MODELO_SPREADSHEET_VERIFY_OPERATION_DEFINITION_ID,
    ModeloSpreadsheetCalculateRequest,
    ModeloSpreadsheetExportOutcome,
    ModeloSpreadsheetExportRequest,
    ModeloSpreadsheetPullRequest,
    ModeloSpreadsheetVerifyRequest,
)
from ...application.modelo.modelo_spreadsheet_operation_projections import ModeloSpreadsheetExportProjection
from ...application.operations.public_period import PublicPeriod
from ...application.storage.calc_sheets.records import TabName
from ...application.user_profile.capabilities import resolve_active_capability
from ...application.workflow.persistence import workflow_state_repository
from ...core.capabilities import ServiceCapability
from ...core.config import load_settings
from ...core.hashing import sha256_hex
from ...core.operations import OperationEffect, OperationTerminalCondition, profile_operation_subject
from ...domain.buckets.event import BucketEventObjectType
from ...domain.user_profile.values import UserProfileFact
from ..modelo_audit_operation_composition import build_modelo_audit_operation_ports
from .conformance_family_contract import (
    ConformanceFamily,
    ConformanceFamilyContext,
    ConformanceOutcome,
    ConformancePreparation,
    RegisteredExecutorConformanceCase,
)
from .modelo_operation_test_support import seeded_modelo_work_unit

_MODELO = "303"
_PERIOD = PublicPeriod(filing_year=2025, code="1T")
_SPREADSHEET_ID = "conformance-spreadsheet"


def _prepare_export(context: ConformanceFamilyContext) -> ConformancePreparation:
    profile_id = context.profile_id
    output = (context.input_root / "modelo-303-2025-1T.xlsx").resolve()
    revision = context.operation.snapshot(_MODELO, filing_year=_PERIOD.filing_year, period=_PERIOD.code).revision.id

    def verify(outcome: ConformanceOutcome) -> None:
        result = outcome.resolve_result(ModeloSpreadsheetExportOutcome)
        assert result.result is not None
        landed = output.read_bytes()
        with zipfile.ZipFile(output) as workbook:
            assert workbook.testzip() is None
        assert result.result.casilla_count > 0
        assert result == ModeloSpreadsheetExportOutcome(
            profile_id=profile_id,
            modelo=_MODELO,
            revision=revision,
            period=_PERIOD,
            outcome="succeeded",
            result=ModeloSpreadsheetExportProjection(
                profile_id=profile_id,
                modelo=_MODELO,
                revision=revision,
                period=_PERIOD,
                output_path=str(output),
                byte_size=len(landed),
                sha256=sha256_hex(landed),
                tab_names=tuple(tab.value for tab in TabName),
                casilla_count=result.result.casilla_count,
                prefill_relations=False,
            ),
        )

    return ConformancePreparation(
        subject_ref=profile_operation_subject(str(profile_id)),
        request=ModeloSpreadsheetExportRequest(
            profile_id=profile_id, modelo=_MODELO, period=_PERIOD, output_path=str(output)
        ),
        verify=verify,
    )


def _prepare(context: ConformanceFamilyContext) -> ConformancePreparation:
    definition_id = context.definition.definition_id
    profile_id = context.profile_id
    subject_ref = profile_operation_subject(str(profile_id))
    if definition_id == MODELO_SPREADSHEET_EXPORT_OPERATION_DEFINITION_ID:
        return _prepare_export(context)
    if definition_id == MODELO_SPREADSHEET_PULL_OPERATION_DEFINITION_ID:
        return ConformancePreparation(
            subject_ref=subject_ref,
            request=ModeloSpreadsheetPullRequest(
                profile_id=profile_id, modelo=_MODELO, period=_PERIOD, spreadsheet_id=_SPREADSHEET_ID
            ),
        )
    if definition_id == MODELO_SPREADSHEET_CALCULATE_OPERATION_DEFINITION_ID:
        return ConformancePreparation(
            subject_ref=subject_ref,
            request=ModeloSpreadsheetCalculateRequest(
                profile_id=profile_id, modelo=_MODELO, period=_PERIOD, spreadsheet_id=_SPREADSHEET_ID
            ),
        )
    if definition_id == MODELO_SPREADSHEET_VERIFY_OPERATION_DEFINITION_ID:
        return ConformancePreparation(
            subject_ref=subject_ref,
            request=ModeloSpreadsheetVerifyRequest(profile_id=profile_id, modelo=_MODELO, period=_PERIOD),
        )
    raise AssertionError(f"no modelo spreadsheet conformance scenario for {definition_id}")


def _drive_root_refusal(definition_id: str) -> RegisteredExecutorConformanceCase:
    return RegisteredExecutorConformanceCase(
        definition_id,
        OperationTerminalCondition.REFUSED,
        OperationEffect.NONE,
        (definition_id,),
        expected_refusal_ref="REFUSED_OUTBOUND_STORAGE_VALIDATION",
    )


MODELO_SPREADSHEET_CONFORMANCE_FAMILY = ConformanceFamily(
    cases=(
        RegisteredExecutorConformanceCase(
            MODELO_SPREADSHEET_EXPORT_OPERATION_DEFINITION_ID,
            OperationTerminalCondition.SUCCEEDED,
            OperationEffect.UPDATED,
            (MODELO_SPREADSHEET_EXPORT_OPERATION_DEFINITION_ID,),
        ),
        _drive_root_refusal(MODELO_SPREADSHEET_PULL_OPERATION_DEFINITION_ID),
        _drive_root_refusal(MODELO_SPREADSHEET_CALCULATE_OPERATION_DEFINITION_ID),
        _drive_root_refusal(MODELO_SPREADSHEET_VERIFY_OPERATION_DEFINITION_ID),
    ),
    prepare=_prepare,
)
_RETAINED_MODELO_REPORTS_PRIVATE_ADDRESS = "audit-private-source-address"
_RETAINED_MODELO_REPORTS_PRIVATE_NOTE = "synthetic private audit note"
_RETAINED_MODELO_REPORTS_EVIDENCE = b"synthetic audit evidence"
_RETAINED_MODELO_REPORTS_PERIOD = PublicPeriod(filing_year=2025, code="1T")


def _retained_modelo_reports_prepare_audit(context: ConformanceFamilyContext) -> ConformancePreparation:
    unit = seeded_modelo_work_unit(context.profile_id, operation=context.operation)
    ports = build_modelo_audit_operation_ports(profile_id=context.profile_id, operation=context.operation)
    bundle = EvidenceBundleService(ports=ports.evidence).build(
        bucket_id=str(context.profile_id),
        work_unit_id=unit.work_unit_id,
        record_payloads={
            (
                BucketEventObjectType.WORK_UNIT.value,
                _RETAINED_MODELO_REPORTS_PRIVATE_ADDRESS,
            ): _RETAINED_MODELO_REPORTS_EVIDENCE
        },
        notes=_RETAINED_MODELO_REPORTS_PRIVATE_NOTE,
    )
    assert ports.evidence.repository.load(bundle.bundle_id) == bundle
    output = context.input_root / "audit.zip"
    definition_id = context.definition.definition_id

    def verify(outcome: ConformanceOutcome) -> None:
        assert ports.evidence.repository.load(bundle.bundle_id) == bundle
        if definition_id == "modelo.audit.read":
            projection = outcome.resolve_result(ModeloAuditReadProjection)
            assert projection.profile_id == context.profile_id
            assert projection.kind == "view"
            assert projection.bundle is not None
            assert projection.bundle.to_bundle() == bundle
        elif definition_id == "modelo.audit.query":
            projection = outcome.resolve_result(ModeloAuditQueryProjection)
            assert projection.profile_id == context.profile_id
            assert projection.bundle is not None
            assert projection.bundle.bundle_id == bundle.bundle_id
            assert projection.bundle.work_unit_id == unit.work_unit_id
            assert projection.bundle.notes_present is True
            assert len(projection.bundle.records) == 1
            record = projection.bundle.records[0]
            assert record.object_type is BucketEventObjectType.WORK_UNIT
            assert record.content_sha256 == hashlib.sha256(_RETAINED_MODELO_REPORTS_EVIDENCE).hexdigest()
            assert record.payload_size_bytes == len(_RETAINED_MODELO_REPORTS_EVIDENCE)
            serialized = projection.model_dump_json()
            assert _RETAINED_MODELO_REPORTS_PRIVATE_NOTE not in serialized
            assert _RETAINED_MODELO_REPORTS_PRIVATE_ADDRESS not in serialized
        else:
            projection = outcome.resolve_result(ModeloAuditExportProjection)
            assert projection.profile_id == context.profile_id
            assert projection.bundle_id == bundle.bundle_id
            assert projection.output == str(output)
            assert projection.records == 1
            assert projection.verification_state == bundle.verification_state
            with zipfile.ZipFile(output) as archive:
                assert archive.namelist() == ["manifest.json"]
                assert EvidenceBundle.model_validate_json(archive.read("manifest.json")) == bundle

    request = (
        ModeloAuditExportRequest(
            profile_id=context.profile_id, bundle_id=bundle.bundle_id, output=output, force_incomplete=True
        )
        if definition_id == "modelo.audit.export"
        else ModeloAuditReadRequest(profile_id=context.profile_id, kind="view", bundle_id=bundle.bundle_id)
    )
    return ConformancePreparation(
        subject_ref=profile_operation_subject(str(context.profile_id)), request=request, verify=verify
    )


def _retained_modelo_reports_prepare_spreadsheet(context: ConformanceFamilyContext) -> ConformancePreparation:
    profile = str(context.profile_id)
    definition_id = context.definition.definition_id
    assert session_store.load_client(profile) is None
    assert session_store.load_token(profile) is None
    assert session_store.load_drive_config(profile) is None
    assert not load_settings().cadrumo_google_drive_root_folder_id
    output = context.input_root / "modelo-130.xlsx"
    if definition_id == "modelo.spreadsheet.calculate":
        session_store.save_drive_config(profile, DriveConfig(root_folder_id="conformance-selected-folder"))
        request = ModeloSpreadsheetCalculateRequest(
            profile_id=context.profile_id,
            modelo="130",
            period=_RETAINED_MODELO_REPORTS_PERIOD,
            spreadsheet_id="conformance-workbook",
        )
    elif definition_id == "modelo.spreadsheet.pull":
        request = ModeloSpreadsheetPullRequest(
            profile_id=context.profile_id,
            modelo="130",
            period=_RETAINED_MODELO_REPORTS_PERIOD,
            spreadsheet_id="conformance-workbook",
        )
    elif definition_id == "modelo.spreadsheet.verify":
        upsert_test_profile_facts(profile, (UserProfileFact(path="capabilities.google_export", value=False),))
        assert resolve_active_capability(ServiceCapability.GOOGLE_EXPORT).enabled is False
        request = ModeloSpreadsheetVerifyRequest(
            profile_id=context.profile_id, modelo="130", period=_RETAINED_MODELO_REPORTS_PERIOD
        )
    else:
        request = ModeloSpreadsheetExportRequest(
            profile_id=context.profile_id, modelo="130", period=_RETAINED_MODELO_REPORTS_PERIOD, output_path=str(output)
        )
    before_drive = session_store.load_drive_config(profile)
    workflow = workflow_state_repository()
    seed_workflow = workflow.load()
    workflow.save(seed_workflow)
    before_workflow = workflow.load()
    assert before_workflow.model_copy(update={"updated_at": seed_workflow.updated_at}) == seed_workflow
    assert before_workflow.updated_at >= seed_workflow.updated_at
    snapshot = context.operation.snapshot("130", filing_year=2025, period="1T")

    def verify(outcome: ConformanceOutcome) -> None:
        assert session_store.load_client(profile) is None
        assert session_store.load_token(profile) is None
        assert session_store.load_drive_config(profile) == before_drive
        assert workflow_state_repository().load() == before_workflow
        if definition_id != "modelo.spreadsheet.export":
            assert not output.exists()
            if definition_id == "modelo.spreadsheet.verify":
                assert resolve_active_capability(ServiceCapability.GOOGLE_EXPORT).enabled is False
            return
        settled = outcome.resolve_result(ModeloSpreadsheetExportOutcome)
        assert settled.outcome == "succeeded"
        assert settled.refusal is None
        projection = settled.result
        assert projection is not None
        payload = output.read_bytes()
        assert projection.profile_id == context.profile_id
        assert projection.modelo == "130"
        assert projection.revision == snapshot.revision.id
        assert projection.period == _RETAINED_MODELO_REPORTS_PERIOD
        assert projection.output_path == str(output)
        assert projection.byte_size == len(payload)
        assert projection.sha256 == hashlib.sha256(payload).hexdigest()
        assert projection.prefill_relations is False
        assert projection.casilla_count > 0
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
        try:
            assert tuple(workbook.sheetnames) == tuple(tab.value for tab in TabName)
            assert projection.tab_names == tuple(workbook.sheetnames)
            assert any(cell.data_type == "f" for row in workbook[TabName.CALCULOS.value] for cell in row)
            guide_values = {
                str(cell.value) for row in workbook[TabName.GUIDE.value] for cell in row if cell.value is not None
            }
            assert "130" in guide_values
            assert str(snapshot.revision.id) in guide_values
            assert "1T / 2025" in guide_values
            assert workbook[TabName.PROVENANCE.value].max_row > 1
        finally:
            workbook.close()

    return ConformancePreparation(subject_ref=profile_operation_subject(profile), request=request, verify=verify)


def _retained_modelo_reports_prepare(context: ConformanceFamilyContext) -> ConformancePreparation:
    if context.definition.definition_id.startswith("modelo.audit."):
        return _retained_modelo_reports_prepare_audit(context)
    return _retained_modelo_reports_prepare_spreadsheet(context)


MODELO_REPORTS_MATERIAL_CONFORMANCE_FAMILY = ConformanceFamily(
    cases=tuple(
        RegisteredExecutorConformanceCase(
            definition_id,
            OperationTerminalCondition.SUCCEEDED,
            OperationEffect.UPDATED if definition_id.endswith("export") else OperationEffect.NONE,
            (definition_id,),
        )
        for definition_id in (
            "modelo.audit.read",
            "modelo.audit.query",
            "modelo.audit.export",
            "modelo.spreadsheet.export",
        )
    )
    + tuple(
        RegisteredExecutorConformanceCase(
            definition_id,
            OperationTerminalCondition.REFUSED,
            OperationEffect.NONE,
            (definition_id,),
            "REFUSED_PROFILE_ACCESS" if definition_id.endswith("verify") else "REFUSED_OUTBOUND_STORAGE_VALIDATION",
        )
        for definition_id in ("modelo.spreadsheet.pull", "modelo.spreadsheet.calculate", "modelo.spreadsheet.verify")
    ),
    prepare=_retained_modelo_reports_prepare,
)
