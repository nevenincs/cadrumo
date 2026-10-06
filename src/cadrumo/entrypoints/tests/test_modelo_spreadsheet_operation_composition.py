"""Local workbook composition has no remote ingress or parity capabilities."""

from __future__ import annotations

from dataclasses import fields
from io import BytesIO
from typing import NoReturn
from uuid import UUID

import pytest
from openpyxl import load_workbook

from ...application.storage.calc_sheets.records import OperatorInputs
from ...core.config import override_settings
from ...domain.calculations.registry.authority import PinnedAuthorityOperation
from ...domain.calculations.registry.governed_fact_scope import validating_governed_facts
from ..modelo_spreadsheet_operation_composition import build_modelo_spreadsheet_operation_ports

pytestmark = [pytest.mark.unit, pytest.mark.hex_entrypoint]
_PROFILE = UUID("5aa00000-0000-4000-8000-0000000000aa")


def test_local_workbook_materializes_without_google_capability(
    authority_operation: PinnedAuthorityOperation, monkeypatch: pytest.MonkeyPatch
) -> None:
    def credentials(*, profile: str) -> NoReturn:
        pytest.fail("local XLSX must never hydrate Google credentials")

    monkeypatch.setattr("cadrumo.adapters.outbound.storage.factory.build_google_credentials", credentials)
    with override_settings(cadrumo_active_profile=str(_PROFILE)), validating_governed_facts(authority_operation):
        ports = build_modelo_spreadsheet_operation_ports(profile_id=_PROFILE, operation=authority_operation)
        assert {field.name for field in fields(ports)} == {"profile_id", "operation", "materialize", "plan_builder"}
        snapshot = authority_operation.snapshot("130", filing_year=2026, period="1T")
        plan = ports.plan_builder(snapshot, operator_inputs=OperatorInputs())
        payload = ports.materialize(plan)
    workbook = load_workbook(BytesIO(payload))
    try:
        assert workbook.sheetnames
        assert any(sheet.max_row > 1 for sheet in workbook)
    finally:
        workbook.close()
