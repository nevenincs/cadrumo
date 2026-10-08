"""Registered-executor conformance scenarios for saved-review exports.

The local workbook exports publish a real calculation or reconciliation
selection through the guarded local sink, and each outcome is judged against
the bytes that actually landed. The Google review publication stays on the
local side of the provider boundary: the isolated profile records no Drive
root folder, so preparation refuses before any credential or network use.
"""

from __future__ import annotations

import hashlib
from collections.abc import Callable
from io import BytesIO
from pathlib import Path
from uuid import NAMESPACE_URL, UUID, uuid5

from openpyxl import load_workbook

from ...adapters.outbound.google.session_store import load_drive_config, load_token
from ...adapters.persistence.profile.modelos_calculation import CalculationRevisionCatalogueRepository
from ...application.export.calculation_review_xlsx_operation import (
    CALCULATION_REVIEW_XLSX_OPERATION_DEFINITION_ID,
    CalculationReviewXlsxRequest,
    CalculationReviewXlsxResult,
)
from ...application.export.google_review_operation_contracts import (
    GOOGLE_REVIEW_OPERATION_DEFINITION_ID,
    GoogleReviewRequest,
)
from ...application.export.review_snapshot import CalculationReviewSelection
from ...application.modelo.reconciliation_export_operation import (
    RECONCILIATION_EXPORT_XLSX_OPERATION_DEFINITION_ID,
    ReconciliationExportXlsxProjection,
    ReconciliationExportXlsxRequest,
)
from ...application.operations.frontend_requests import OperationPublicEffectEventV1
from ...application.storage.calc_sheets.records import TabName
from ...core.operations import OperationEffect, OperationTerminalCondition, profile_operation_subject
from ...domain.modelos.calculation_revision import CalculationRevisionState
from ..calculation_review_snapshot_composition import load_calculation_review_snapshot
from .conformance_family_contract import (
    ConformanceFamily,
    ConformanceFamilyContext,
    ConformanceOutcome,
    ConformancePreparation,
    RegisteredExecutorConformanceCase,
)
from .modelo_operation_test_support import seeded_modelo_calculation_revision

_PUBLICATION_ID_SEED = "conformance-google-review-publication"


def _landed(path_text: str) -> tuple[int, str]:
    data = Path(path_text).read_bytes()
    return len(data), hashlib.sha256(data).hexdigest()


def _prepare_calculation_review(context: ConformanceFamilyContext) -> ConformancePreparation:
    revision_id = seeded_modelo_calculation_revision(context.profile_id, operation=context.operation)
    revision = CalculationRevisionCatalogueRepository().load(operation=context.operation).get(revision_id)
    assert revision is not None
    output = (context.input_root / "saved-calculation-review.xlsx").resolve()
    request = CalculationReviewXlsxRequest(
        profile_id=context.profile_id,
        calculation_revision_id=revision_id,
        output_path=str(output),
    )

    def verify(outcome: ConformanceOutcome) -> None:
        result = outcome.resolve_result(CalculationReviewXlsxResult)
        byte_size, sha256 = _landed(result.output_path)
        assert result.output_path == str(output)
        assert (result.byte_size, result.file_sha256) == (byte_size, sha256)
        assert result.profile_id == context.profile_id
        assert result.calculation_revision_id == revision_id
        assert result.work_unit_id == revision.work_unit_id
        assert result.calculation_state is CalculationRevisionState.BORRADOR
        workbook = load_workbook(BytesIO(output.read_bytes()), data_only=False)
        assert TabName.CALCULOS.value in workbook.sheetnames
        assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)

    return ConformancePreparation(
        subject_ref=profile_operation_subject(str(context.profile_id)), request=request, verify=verify
    )


def _prepare_google_review(context: ConformanceFamilyContext) -> ConformancePreparation:
    profile = str(context.profile_id)
    revision_id = seeded_modelo_calculation_revision(context.profile_id, operation=context.operation)
    # The selected revision resolves locally, so the refusal can only come from
    # the missing Drive root, not from the saved selection.
    snapshot = load_calculation_review_snapshot(revision_id, profile_id=context.profile_id, operation=context.operation)
    assert isinstance(snapshot.selection, CalculationReviewSelection)
    assert snapshot.selection.calculation_revision_id == revision_id
    assert load_drive_config(profile) is None
    request = GoogleReviewRequest(
        profile_id=context.profile_id,
        calculation_revision_id=revision_id,
        publication_id=_publication_id(),
    )

    def verify(outcome: ConformanceOutcome) -> None:
        assert load_drive_config(profile) is None
        assert load_token(profile) is None
        # Nothing crossed the provider boundary: no handoff ever opened an effect in doubt.
        assert {
            event.effect
            for event in outcome.observed.event_page.events
            if isinstance(event, OperationPublicEffectEventV1)
        } == {OperationEffect.NONE}

    return ConformancePreparation(subject_ref=profile_operation_subject(profile), request=request, verify=verify)


def _publication_id() -> UUID:
    return uuid5(NAMESPACE_URL, _PUBLICATION_ID_SEED)


def _prepare_reconciliation_export(context: ConformanceFamilyContext) -> ConformancePreparation:
    output = (context.input_root / "reconciliation-history.xlsx").resolve()
    request = ReconciliationExportXlsxRequest(
        profile_id=context.profile_id,
        all_history=True,
        output_path=str(output),
    )

    def verify(outcome: ConformanceOutcome) -> None:
        result = outcome.resolve_result(ReconciliationExportXlsxProjection)
        byte_size, sha256 = _landed(result.output_path)
        assert result.output_path == str(output)
        assert (result.byte_size, result.file_sha256) == (byte_size, sha256)
        assert result.profile_id == context.profile_id
        assert result.all_history is True
        assert (result.event_id, result.work_unit_id) == (None, None)
        # A fresh profile has no saved reconciliation, and the export says so.
        assert (result.reconciliation_count, result.difference_count, result.advisory_count) == (0, 0, 0)
        workbook = load_workbook(BytesIO(output.read_bytes()), data_only=False)
        assert workbook.sheetnames
        assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)

    return ConformancePreparation(
        subject_ref=profile_operation_subject(str(context.profile_id)), request=request, verify=verify
    )


_PREPARERS: dict[str, Callable[[ConformanceFamilyContext], ConformancePreparation]] = {
    CALCULATION_REVIEW_XLSX_OPERATION_DEFINITION_ID: _prepare_calculation_review,
    GOOGLE_REVIEW_OPERATION_DEFINITION_ID: _prepare_google_review,
    RECONCILIATION_EXPORT_XLSX_OPERATION_DEFINITION_ID: _prepare_reconciliation_export,
}


def _prepare(context: ConformanceFamilyContext) -> ConformancePreparation:
    preparer = _PREPARERS.get(context.definition.definition_id)
    if preparer is None:
        raise AssertionError(f"no saved-review export conformance scenario for {context.definition.definition_id}")
    return preparer(context)


SAVED_REVIEW_EXPORT_CONFORMANCE_FAMILY = ConformanceFamily(
    cases=(
        RegisteredExecutorConformanceCase(
            CALCULATION_REVIEW_XLSX_OPERATION_DEFINITION_ID,
            OperationTerminalCondition.SUCCEEDED,
            OperationEffect.UPDATED,
            (CALCULATION_REVIEW_XLSX_OPERATION_DEFINITION_ID,),
        ),
        RegisteredExecutorConformanceCase(
            GOOGLE_REVIEW_OPERATION_DEFINITION_ID,
            OperationTerminalCondition.REFUSED,
            OperationEffect.NONE,
            (GOOGLE_REVIEW_OPERATION_DEFINITION_ID + ".prepare",),
            expected_refusal_ref="REFUSED_PROFILE_ACCESS",
        ),
        RegisteredExecutorConformanceCase(
            RECONCILIATION_EXPORT_XLSX_OPERATION_DEFINITION_ID,
            OperationTerminalCondition.SUCCEEDED,
            OperationEffect.UPDATED,
            (RECONCILIATION_EXPORT_XLSX_OPERATION_DEFINITION_ID,),
        ),
    ),
    prepare=_prepare,
)
