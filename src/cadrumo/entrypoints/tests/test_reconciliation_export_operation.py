"""The registered export selects saved detail and writes through the guarded local sink."""

from __future__ import annotations

import asyncio
from collections.abc import Iterator
from contextlib import asynccontextmanager
from datetime import UTC, datetime
from pathlib import Path
from types import SimpleNamespace
from typing import cast
from uuid import UUID

import pytest
from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError

from ...adapters.outbound.workbook.calc_sheets_xlsx import materialize_export_plan
from ...application.modelo.reconciliation_export_labels import ReconciliationWorkbookLabels
from ...application.modelo.reconciliation_export_operation import (
    RECONCILIATION_EXPORT_XLSX_OPERATION_DEFINITION_ID,
    ReconciliationExportXlsxExecutor,
    ReconciliationExportXlsxPorts,
    ReconciliationExportXlsxProjection,
    ReconciliationExportXlsxRequest,
    ReconciliationExportXlsxResult,
    build_reconciliation_export_xlsx_definition,
    build_reconciliation_export_xlsx_registration,
)
from ...application.modelo.reconciliation_records import (
    ModeloReconciliationRecord,
    bind_modelo_reconciliation_persistence_factory,
)
from ...application.modelo.tests.reconciliation_fixture import RECONCILIATION_PROFILE_FIXTURE as _PROFILE
from ...application.modelo.tests.reconciliation_fixture import reconciliation_history_entry_fixture as _history_entry
from ...application.operations.models import OperationIdentity, OperationRequest, OperationTerminalReceipt
from ...application.operations.owner import OperationExecutorContext
from ...application.operations.registry import OperationRegistry
from ...application.user_profile.access_errors import ProfileAccessRefusedError
from ...core.config import override_settings
from ...core.external_constants import OutputLanguage
from ...core.operations import OperationEffect, OperationTerminalCondition, profile_operation_subject
from ...domain.calculations.registry.authority import PinnedAuthorityOperation
from ...domain.calculations.registry.schema_references import RegistrySnapshotRef
from ...domain.modelos.errors import ModeloExportError

pytestmark = [pytest.mark.unit, pytest.mark.hex_entrypoint]


class _Records:
    def iter_records(self) -> Iterator[ModeloReconciliationRecord]:
        for event_id, evidence_value in (("a" * 64, "101.00"), ("b" * 64, "999.00")):
            entry = _history_entry(event_id=event_id)
            yield ModeloReconciliationRecord(
                bucket_event_id=entry.event_id,
                bucket_id=entry.bucket_id,
                work_unit_id=entry.work_unit_id,
                registry_snapshot_ref=RegistrySnapshotRef(
                    modelo="303", modelo_year=2026, period="1T", revision_id="saved-2026"
                ),
                source_kind=entry.source_kind,
                source_ref=entry.source_path,
                verdict=entry.verdict,
                diffs=(entry.diffs[0].model_copy(update={"evidence_value": evidence_value}),),
                actor=entry.actor,
                reconciled_at=entry.reconciled_at,
            )

    def persist_with_event(self, record, event) -> None:
        pytest.fail("review export must not mutate saved history")


def _run(
    payload: ReconciliationExportXlsxRequest, *, foreign: bool = False, change_profile_before_publish: bool = False
):
    operation = cast(PinnedAuthorityOperation, object())
    results: list[BaseModel] = []
    effects: list[OperationEffect] = []

    def factory(*, profile_id: UUID, operation: PinnedAuthorityOperation):
        return ReconciliationExportXlsxPorts(
            profile_id=UUID(int=1) if foreign else profile_id, operation=operation, materialize=materialize_export_plan
        )

    class Events:
        async def phase(self, _phase: str) -> None:
            pass

        async def effect(self, effect: OperationEffect) -> None:
            effects.append(effect)

    class Operands:
        async def put(self, result: BaseModel, *, written_at: datetime) -> str:
            results.append(result)
            return "f" * 64

    class Cancellation:
        @asynccontextmanager
        async def irreversible_section(self):
            with override_settings(
                cadrumo_active_profile=str(UUID(int=1)) if change_profile_before_publish else str(_PROFILE)
            ):
                yield

    context = cast(
        OperationExecutorContext,
        cast(
            object,
            SimpleNamespace(
                identity=OperationIdentity(
                    operation_id="c" * 64,
                    definition_id=RECONCILIATION_EXPORT_XLSX_OPERATION_DEFINITION_ID,
                    subject_ref=profile_operation_subject(str(_PROFILE)),
                ),
                authority_operation=operation,
                events=Events(),
                operands=Operands(),
                cancellation=Cancellation(),
            ),
        ),
    )
    request = OperationRequest[BaseModel](
        definition_id=RECONCILIATION_EXPORT_XLSX_OPERATION_DEFINITION_ID,
        subject_ref=profile_operation_subject(str(_PROFILE)),
        payload=payload,
    )
    with (
        override_settings(cadrumo_active_profile=str(_PROFILE)),
        bind_modelo_reconciliation_persistence_factory(_Records),
    ):
        reference = asyncio.run(ReconciliationExportXlsxExecutor(factory).execute(request, context))
    return reference, results, effects


def test_registered_export_writes_only_the_selected_saved_comparison(tmp_path: Path) -> None:
    output = tmp_path / "reconciliation.xlsx"
    payload = ReconciliationExportXlsxRequest(profile_id=_PROFILE, event_id="a" * 64, output_path=str(output))
    reference, results, effects = _run(payload)
    assert reference == "f" * 64
    assert effects == [OperationEffect.UNKNOWN, OperationEffect.UPDATED]
    assert isinstance(results[0], ReconciliationExportXlsxResult)
    assert results[0].event_id == "a" * 64
    assert results[0].reconciliation_count == 1
    workbook = load_workbook(output)
    assert workbook["Detalle"]["E2"].value == "101.00"
    assert workbook["Detalle"].max_row == 2
    assert all(cell.value != "999.00" for row in workbook["Detalle"] for cell in row)


@pytest.mark.parametrize("defect", ["missing_record", "foreign_ports", "profile_changed", "existing_file"])
def test_export_refuses_before_publication_for_selection_identity_and_destination_defects(
    tmp_path: Path, defect: str
) -> None:
    output = tmp_path / "review.xlsx"
    if defect == "existing_file":
        output.write_bytes(b"operator artifact")
    payload = ReconciliationExportXlsxRequest(
        profile_id=_PROFILE, event_id=("e" * 64 if defect == "missing_record" else "a" * 64), output_path=str(output)
    )
    with pytest.raises((ProfileAccessRefusedError, ModeloExportError)):
        _run(payload, foreign=defect == "foreign_ports", change_profile_before_publish=defect == "profile_changed")
    assert output.read_bytes() == b"operator artifact" if defect == "existing_file" else not output.exists()


def test_request_requires_explicit_narrow_or_whole_history_selection(tmp_path: Path) -> None:
    for selectors in ({}, {"event_id": "a" * 64, "all_history": True}, {"work_unit_id": "1" * 64, "all_history": True}):
        with pytest.raises(ValidationError):
            ReconciliationExportXlsxRequest.model_validate(
                {"profile_id": _PROFILE, "output_path": str(tmp_path / "review.xlsx"), **selectors}
            )


def test_public_receipt_refuses_another_operation_and_unsettled_effect(tmp_path: Path) -> None:
    _, results, _ = _run(
        ReconciliationExportXlsxRequest(
            profile_id=_PROFILE,
            event_id="a" * 64,
            output_path=str(tmp_path / "review.xlsx"),
        )
    )
    value = results[0]
    assert isinstance(value, ReconciliationExportXlsxResult)
    receipt = OperationTerminalReceipt(
        identity=value.operation_identity,
        revision=1,
        condition=OperationTerminalCondition.SUCCEEDED,
        effect=OperationEffect.UPDATED,
        settled_at=datetime(2026, 10, 7, tzinfo=UTC),
        result_ref="f" * 64,
    )

    def unused_factory(*, profile_id: UUID, operation: PinnedAuthorityOperation) -> ReconciliationExportXlsxPorts:
        raise AssertionError("result projection must not construct execution ports")

    definition = build_reconciliation_export_xlsx_definition(unused_factory)
    projector = build_reconciliation_export_xlsx_registration(definition).result_projector
    assert projector is not None
    public = projector(value, receipt)
    assert isinstance(public, ReconciliationExportXlsxProjection)
    assert "publication_id" not in public.model_dump()
    with pytest.raises(ValueError):
        projector(
            value,
            receipt.model_copy(update={"identity": receipt.identity.model_copy(update={"operation_id": "e" * 64})}),
        )
    with pytest.raises(ValueError):
        projector(value, receipt.model_copy(update={"effect": OperationEffect.UNKNOWN}))


@pytest.mark.parametrize("scope", ["all", "work_unit"])
def test_export_larger_sets_only_when_the_operator_explicitly_selects_them(tmp_path: Path, scope: str) -> None:
    payload = ReconciliationExportXlsxRequest(
        profile_id=_PROFILE,
        all_history=scope == "all",
        work_unit_id="1" * 64 if scope == "work_unit" else None,
        output_path=str(tmp_path / "selected.xlsx"),
    )
    _, results, _ = _run(payload)
    assert isinstance(results[0], ReconciliationExportXlsxResult)
    assert results[0].reconciliation_count == 2


def test_enrollment_and_required_copy_are_complete_for_every_locale() -> None:
    def unused_factory(*, profile_id: UUID, operation: PinnedAuthorityOperation) -> ReconciliationExportXlsxPorts:
        raise AssertionError("registration must not construct execution ports")

    definition = build_reconciliation_export_xlsx_definition(unused_factory)
    registration = build_reconciliation_export_xlsx_registration(definition)
    registry = OperationRegistry(definitions=(definition,), public_registrations=(registration,))
    assert registry is not None
    for locale in OutputLanguage:
        labels = ReconciliationWorkbookLabels(locale)
        for key in ("title", "saved_comparison_notice", "saved_value", "evidence_value", "advisory"):
            assert labels(key).strip()
