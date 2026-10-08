"""Opt-in live export acceptance inside an already owned synthetic-profile runtime.

No runtime, OAuth, or provider client is created here. The owning live runner
calls this helper after proving runtime ownership. Every action uses real CLI
commands and the approved inherited HANDLE channel for the DPAPI passphrase.
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import os
import subprocess
import sys
from collections.abc import Sequence
from decimal import Decimal
from pathlib import Path
from uuid import UUID, uuid4

from pydantic import TypeAdapter

_PROFILE = "Session02 Google Review"
_JSON_OBJECT: TypeAdapter[dict[str, object]] = TypeAdapter(dict[str, object])
_JSON_OBJECT_ROWS: TypeAdapter[list[dict[str, object]]] = TypeAdapter(list[dict[str, object]])
_TEXT_LIST: TypeAdapter[list[str]] = TypeAdapter(list[str])
_TEXT_MAP: TypeAdapter[dict[str, str]] = TypeAdapter(dict[str, str])


def _text_field(document: dict[str, object], key: str) -> str:
    value = document[key]
    if not isinstance(value, str) or not value:
        raise RuntimeError(f"acceptance field {key!r} requires nonempty text")
    return value


def _cli_result(stdout: str, stage: str) -> dict[str, object]:
    envelope = _JSON_OBJECT.validate_json(stdout, strict=True)
    if envelope.get("active_profile") != _PROFILE:
        raise RuntimeError(f"{stage}: unexpected CLI profile/result")
    return _JSON_OBJECT.validate_python(envelope["result"], strict=True)


def _write_json(path: Path, value: dict[str, object]) -> None:
    temporary = path.with_suffix(".pending")
    temporary.write_text(json.dumps(value, indent=2, ensure_ascii=False), encoding="utf-8")
    temporary.replace(path)


def run_protected_live_cli(root: Path, logs: Path, stage: str, arguments: Sequence[str]) -> dict[str, object]:
    """Run only the CLI child, with one allowlisted secret HANDLE and safe logs."""
    if sys.platform != "win32":
        raise RuntimeError("protected live CLI requires Windows")
    import msvcrt

    import win32crypt

    secret = win32crypt.CryptUnprotectData(
        (root / "session02-driver/synthetic-profile-input.dpapi").read_bytes(), None, None, None, 0
    )[1].decode()
    read_fd, write_fd = os.pipe()
    try:
        os.write(write_fd, json.dumps({"profile_passphrase": secret}).encode())
    finally:
        os.close(write_fd)
    handle = msvcrt.get_osfhandle(read_fd)
    os.set_handle_inheritable(handle, True)
    startup = subprocess.STARTUPINFO()
    startup.lpAttributeList = {"handle_list": [handle]}
    environment = {
        **os.environ,
        "CADRUMO_LOCAL_STORAGE_ROOT": str(root),
        "CADRUMO_SECRET_STORE_DIR": str(root / "secrets"),
        "CADRUMO_OUTPUT_LANGUAGE": "en",
        "CADRUMO_CLI_REVEAL_IDENTIFIERS": "1",
        "CADRUMO_LOG_DIR": str(root / "logs"),
    }
    environment.pop("CADRUMO_DEV_RUNTIME_SESSION_OVERRIDE", None)
    try:
        completed = subprocess.run(  # noqa: S603 -- fixed CLI bootstrap, no shell; caller owns runtime.
            [
                sys.executable,
                "-m",
                "cadrumo.entrypoints.cli._windows_profile_secret_bootstrap",
                "--profile-handle",
                str(handle),
                "--",
                "--format",
                "json",
                "--profile",
                _PROFILE,
                *arguments,
            ],
            startupinfo=startup,
            close_fds=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            env=environment,
            timeout=None,
            check=False,
        )
    finally:
        os.close(read_fd)
    leaked = secret in completed.stdout or secret in completed.stderr
    stdout = completed.stdout.replace(secret, "<redacted>")
    stderr = completed.stderr.replace(secret, "<redacted>")
    del secret
    (logs / f"{stage}.json").write_text(stdout, encoding="utf-8")
    (logs / f"{stage}.stderr").write_text(stderr, encoding="utf-8")
    if leaked:
        raise RuntimeError(f"{stage}: protected-input output check failed")
    if completed.returncode:
        raise RuntimeError(f"{stage}: CLI refused (exit {completed.returncode}); inspect acceptance logs")
    return _cli_result(stdout, stage)


def run_live_export_acceptance(workspace: Path, storage_root: Path) -> dict[str, object]:
    """Save a fresh synthetic revision, export XLSX, and publish its native review.

    The journal is created before any effect and reused on retries. Completed
    calls return their existing receipt, so an interrupted caller cannot create
    another spreadsheet by rerunning this helper.
    """
    if sys.platform != "win32":
        raise RuntimeError("live acceptance requires the approved Windows descriptor channel")
    workspace, root = workspace.resolve(), storage_root.resolve()
    if root != (workspace / "var/storage/session02-google-live").resolve():
        raise RuntimeError("live acceptance is restricted to the existing synthetic storage")
    logs = root / "logs/live-export-acceptance"
    logs.mkdir(parents=True, exist_ok=True)
    journal_path = logs / "acceptance-journal.json"
    journal: dict[str, object]
    if journal_path.exists():
        journal = _JSON_OBJECT.validate_json(journal_path.read_text(encoding="utf-8"), strict=True)
    else:
        run_id = uuid4()
        source_envelope = _JSON_OBJECT.validate_json(
            (root / "session03-driver/work-create-130.json").read_text(encoding="utf-8"), strict=True
        )
        source = _JSON_OBJECT.validate_python(source_envelope["result"], strict=True)
        period = _JSON_OBJECT.validate_python(source["period"], strict=True)
        if source["modelo"] != "130" or source["filing_year"] != 2025 or period["code"] != "4T":
            raise RuntimeError("synthetic work selection no longer matches its original coordinates")
        journal = {
            "version": 1,
            "run_id": str(run_id),
            "publication_id": str(uuid4()),
            "work_unit_id": _text_field(source, "work_unit_id"),
            "synthetic_input": str(Decimal(100000 + run_id.int % 90000000) / 100),
        }
        _write_json(journal_path, journal)
    publication_id = _text_field(journal, "publication_id")
    UUID(publication_id)
    evidence = logs / _text_field(journal, "run_id")
    evidence.mkdir(exist_ok=True)
    work_id = _text_field(journal, "work_unit_id")
    if "calculation_revision_id" not in journal:
        if "previous_revision_ids" not in journal:
            previous = run_protected_live_cli(
                root, evidence, "previous-revisions", ["app", "modelo", "work", "revisions", work_id]
            )
            previous_rows = _JSON_OBJECT_ROWS.validate_python(previous["revisions"], strict=True)
            journal["previous_revision_ids"] = [_text_field(row, "calculation_revision_id") for row in previous_rows]
            _write_json(journal_path, journal)
        old_ids = set(_TEXT_LIST.validate_python(journal["previous_revision_ids"], strict=True))
        calculated = run_protected_live_cli(
            root,
            evidence,
            "calculate",
            [
                "app",
                "modelo",
                "work",
                "calculate",
                work_id,
                "--casilla",
                f"05={_text_field(journal, 'synthetic_input')}",
                "--casilla",
                "06=0.00",
                "--binding",
                "irpf.previous_year_economic_activity_net_income=13000",
                "--binding",
                "modelo-130-resultados-negativos-anteriores=0",
            ],
        )
        revision_id = _text_field(calculated, "calculation_revision_id")
        if revision_id in old_ids or calculated.get("saved") is not True:
            raise RuntimeError("calculation did not save a new synthetic revision")
        journal["calculation_revision_id"] = revision_id
        _write_json(journal_path, journal)
    revision_id = _text_field(journal, "calculation_revision_id")
    baseline = run_protected_live_cli(
        root, evidence, "saved-revision", ["app", "modelo", "work", "revision", revision_id]
    )
    if baseline["calculation_revision_id"] != revision_id or baseline["work_unit_id"] != work_id:
        raise RuntimeError("saved revision differs from the selected synthetic calculation")
    inputs = _TEXT_MAP.validate_python(baseline["input_values_by_casilla_id"], strict=True)
    if not any(Decimal(value) == Decimal(_text_field(journal, "synthetic_input")) for value in inputs.values()):
        raise RuntimeError("retained synthetic input did not round-trip")
    output = evidence / "saved-calculation-review.xlsx"
    if "xlsx" not in journal:
        local_result = run_protected_live_cli(
            root,
            evidence,
            "local-export",
            [
                "app",
                "modelo",
                "spreadsheet",
                "review",
                "--calculation-revision-id",
                revision_id,
                "--output",
                str(output),
                "--replace",
            ],
        )
        local = _JSON_OBJECT.validate_python(local_result["publication"], strict=True)
        if local["calculation_revision_id"] != revision_id:
            raise RuntimeError("local export selected a different revision")
        journal["xlsx"] = local
        _write_json(journal_path, journal)
    from openpyxl import load_workbook

    workbook = load_workbook(output, data_only=False)
    try:
        values = [str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value is not None]
        formula_count = sum(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
        if revision_id not in values or formula_count or len(workbook.sheetnames) < 3:
            raise RuntimeError("actual workbook lacks retained revision identity or immutable review structure")
        sheet_count = len(workbook.sheetnames)
    finally:
        workbook.close()
    xlsx = _JSON_OBJECT.validate_python(journal["xlsx"], strict=True)
    if hashlib.sha256(output.read_bytes()).hexdigest() != _text_field(xlsx, "file_sha256"):
        raise RuntimeError("published XLSX bytes differ from the registered receipt")
    if "google" not in journal:
        native_result = run_protected_live_cli(
            root,
            evidence,
            "native-publish",
            [
                "app",
                "modelo",
                "spreadsheet",
                "publish",
                "--calculation-revision-id",
                revision_id,
                "--publication-id",
                publication_id,
                "--accept-readable-export",
            ],
        )
        native = _JSON_OBJECT.validate_python(native_result["publication"], strict=True)
        if native["publication_id"] != publication_id:
            raise RuntimeError("native publication returned another idempotency identity")
        if _text_field(native, "snapshot_digest") != _text_field(xlsx, "snapshot_digest"):
            raise RuntimeError("local and native exports used different retained snapshots")
        journal["google"] = native
        _write_json(journal_path, journal)
    package = importlib.util.find_spec("cadrumo")
    if package is None or package.origin is None:
        raise RuntimeError("CLI package origin cannot be established")
    google = _JSON_OBJECT.validate_python(journal["google"], strict=True)
    result = {
        "python_executable": sys.executable,
        "package_origin": package.origin,
        "profile_label": _PROFILE,
        "work_unit_id": work_id,
        "calculation_revision_id": revision_id,
        "publication_id": publication_id,
        "spreadsheet_url": _text_field(google, "spreadsheet_url"),
        "spreadsheet_id": _text_field(google, "spreadsheet_id"),
        "xlsx_path": str(output),
        "evidence_directory": str(evidence),
        "checks": {
            "fresh_saved_synthetic_input": True,
            "xlsx_sheet_count": sheet_count,
            "xlsx_formula_count": formula_count,
            "xlsx_digest_verified": True,
            "same_local_native_snapshot": True,
            "protected_input_not_output": True,
        },
    }
    _write_json(evidence / "verified-result.json", result)
    return result
