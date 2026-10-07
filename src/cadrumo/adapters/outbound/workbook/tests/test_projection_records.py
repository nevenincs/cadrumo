"""The XLSX transport retains saved activity values and absent record slots."""

import pytest

from cadrumo.application.modelo.tests.test_work_form_projection_records import (
    annual_saved as annual_saved,
)
from cadrumo.application.modelo.tests.test_work_form_projection_records import (
    projection_block as projection_block,
)
from cadrumo.application.modelo.tests.test_work_form_projection_records import (
    saved as saved,
)
from cadrumo.application.modelo.tests.test_work_form_projection_records import (
    snapshot as snapshot,
)
from cadrumo.domain.calculations.registry.schema_form_layouts import (
    FormContextFieldBlock,
    FormLayoutDefinition,
    FormPageDefinition,
    FormPlacementDefinition,
    FormSectionDefinition,
)

pytestmark = [pytest.mark.unit, pytest.mark.hex_outbound_adapter]


def test_workbook_renders_saved_projection_values_without_recalculating_them(snapshot, saved, projection_block):
    from io import BytesIO

    from openpyxl import load_workbook

    from cadrumo.application.storage.calc_sheets.engine import build_export_plan
    from cadrumo.application.storage.calc_sheets.form_workbook import add_form_workbook
    from cadrumo.application.storage.calc_sheets.records import TabName

    from ..calc_sheets_xlsx import materialize_export_plan

    block = projection_block
    # Synthetic placement exercises the consumer without publishing authoring
    # source into the runtime authority used by the fixture.
    layout = FormLayoutDefinition(
        id="projection-test",
        revision_id=snapshot.revision.id,
        seed_source="authored",
        generator_version=1,
        source_state_digest="a" * 64,
        pages=(
            FormPageDefinition(
                id="page",
                heading_key="test.page",
                sections=(FormSectionDefinition(id="activity", heading_key="test.activity", blocks=(block,)),),
            ),
        ),
        placements=tuple(
            FormPlacementDefinition(casilla_id=c.id, kind="working_figure") for c in snapshot.revision.casillas
        ),
    )
    snapshot = snapshot.model_copy(
        update={"revision": snapshot.revision.model_copy(update={"form_layouts": (layout,)})}
    )
    plan = add_form_workbook(build_export_plan(snapshot), snapshot, revision=saved)
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)))
    form = workbook[TabName.FORM.value]
    quantity_cells = [form.cell(cell.row, 9) for row in form for cell in row if cell.value == "Dato 24"]
    assert [cell.value for cell in quantity_cells] == [0, 2]
    assert all(cell.data_type != "f" for cell in quantity_cells)
    # The third activity occupies the first slot of the second record. Its
    # absent second slot must not borrow an activity or become a spreadsheet 0.
    other_slots = [form.cell(cell.row, 9).value for row in form for cell in row if cell.value == "Dato 52"]
    assert other_slots == [1, "Sin dato"]
    workbook.close()


def test_workbook_renders_annual_saved_context_as_human_values(annual_saved):
    from io import BytesIO

    from openpyxl import load_workbook

    from cadrumo.application.storage.calc_sheets.engine import build_export_plan
    from cadrumo.application.storage.calc_sheets.form_workbook import add_form_workbook
    from cadrumo.application.storage.calc_sheets.records import TabName

    from ..calc_sheets_xlsx import materialize_export_plan

    snapshot, saved = annual_saved
    labels = (
        (6, "Código principal"),
        (7, "IAE principal"),
        (8, "Código segunda actividad"),
        (18, "Declaración de terceros"),
    )
    blocks = tuple(
        FormContextFieldBlock(
            id=f"annual-{number}",
            export_layout_id=snapshot.revision.export_layouts[0].id,
            export_record_id="m303-exonerado-390",
            export_field_id=f"m303-2026.dp30304.f{number:03}",
            heading_key=f"test.annual.{number}",
            official_heading=label,
        )
        for number, label in labels
    )
    layout = FormLayoutDefinition(
        id="annual-test",
        revision_id=snapshot.revision.id,
        seed_source="authored",
        generator_version=1,
        source_state_digest="a" * 64,
        pages=(
            FormPageDefinition(
                id="annual",
                heading_key="test.annual",
                sections=(FormSectionDefinition(id="activities", heading_key="test.activities", blocks=blocks),),
            ),
        ),
        placements=tuple(
            FormPlacementDefinition(casilla_id=c.id, kind="working_figure") for c in snapshot.revision.casillas
        ),
    )
    snapshot = snapshot.model_copy(
        update={"revision": snapshot.revision.model_copy(update={"form_layouts": (layout,)})}
    )
    plan = add_form_workbook(build_export_plan(snapshot), snapshot, revision=saved)
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)))
    form = workbook[TabName.FORM.value]
    rendered = {
        cell.value: form.cell(row_number, 9).value
        for row_number, row in enumerate(form, 1)
        for cell in row
        if cell.value in {label for _, label in labels}
    }
    assert rendered == {
        "Código principal": "A03",
        "IAE principal": "6732",
        "Código segunda actividad": "Sin dato",
        "Declaración de terceros": "No",
    }
    workbook.close()
