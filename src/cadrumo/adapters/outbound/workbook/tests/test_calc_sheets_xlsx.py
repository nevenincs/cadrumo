"""Offline workbook materialization of real published export plans.

The subject is the ``.xlsx`` payload itself: every assertion reads the bytes back
through ``openpyxl`` rather than inspecting the writer's intentions, because a
facet the writer resolves but never serializes is a facet the operator never
gets. Two published modelos are exercised -- 130 (pagos fraccionados, one small
grid) and 303 (IVA trimestral, a wider grid with constrained inputs) -- so a
behaviour that only holds for the smallest plan is not mistaken for the contract.
"""

from __future__ import annotations

from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Protocol, cast

import pytest
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.packaging.custom import StringProperty
from openpyxl.workbook.workbook import Workbook

from cadrumo.domain.calculations.registry.tests.published_authority import published_snapshot

from .....application.storage.calc_sheets.engine import build_export_plan
from .....application.storage.calc_sheets.export_tables import IDENTITY_STAMP_KEYS, export_identity_stamps
from .....application.storage.calc_sheets.records import (
    OperatorInput,
    OperatorInputs,
    SheetCellAddress,
    SheetExportMetadata,
    SheetExportPlan,
    SheetGuideContent,
    SheetTemplatePreviewMetadata,
    SheetValueCell,
    TabName,
)
from .....application.storage.calc_sheets.theme import ROLE_STYLES, WORKBOOK_FONT_FAMILY, StyleRole
from .....core.errors.hierarchy import InternalInvariantError
from .....core.period import Period
from .....domain.calculations.registry.schema import RegistrySnapshot
from ..calc_sheets_xlsx import materialize_export_plan, unrendered_plan_facets

pytestmark = [pytest.mark.unit, pytest.mark.hex_outbound_adapter, pytest.mark.usefixtures("operation")]


class _CustomPropertyReader(Protocol):
    """Read side of the custom document properties the installed stubs omit."""

    @property
    def props(self) -> list[StringProperty]:
        """Every custom document property the workbook carries."""
        ...


class _WorkbookCustomProperties(Protocol):
    """A loaded workbook exposing its custom document properties."""

    @property
    def custom_doc_props(self) -> _CustomPropertyReader:
        """The workbook's custom document properties."""
        ...


def _m130_snapshot() -> RegistrySnapshot:
    return published_snapshot("130", filing_year=2025, period="1T", on=date(2025, 4, 1))


def test_fictional_preview_materializes_without_filing_identity_or_technical_cells() -> None:
    plan = SheetExportPlan[SheetTemplatePreviewMetadata](
        metadata=SheetTemplatePreviewMetadata(
            kind="template_preview",
            modelo_id="232",
            revision_id="2016-2017",
            preview_year=2017,
            preview_period="0A",
            template_digest="a" * 64,
            engine_version="test",
            title="Modelo 232 · Ejemplo ficticio 2017",
            exported_at=datetime(2026, 10, 6, tzinfo=UTC),
        ),
        human_presentation=True,
        tabs=(TabName.FORM,),
        value_cells=(
            SheetValueCell(address=SheetCellAddress.at(TabName.FORM, 1, 1), value="Ejemplo ficticio", role="label"),
        ),
        guide=SheetGuideContent(title="Ejemplo ficticio", paragraphs=("No válido para presentar.",)),
    )
    book = _loaded(materialize_export_plan(plan))
    assert book.properties.title == plan.metadata.title
    assert book.sheetnames == [TabName.FORM.value]
    assert book[TabName.FORM.value]["A1"].value == "Ejemplo ficticio"
    assert not cast("_WorkbookCustomProperties", book).custom_doc_props.props


def _m303_snapshot() -> RegistrySnapshot:
    return published_snapshot("303", filing_year=2025, period="1T", on=date(2025, 4, 1))


def _m130_plan() -> SheetExportPlan:
    return build_export_plan(_m130_snapshot())


def _m303_plan() -> SheetExportPlan:
    return build_export_plan(_m303_snapshot())


def _loaded(payload: bytes) -> Workbook:
    return load_workbook(BytesIO(payload))


def _cell_at(book: Workbook, address: SheetCellAddress) -> Cell:
    cell = book[address.tab.value].cell(row=address.row, column=address.column)
    assert isinstance(cell, Cell)
    return cell


def test_workbook_carries_every_plan_tab_in_plan_order() -> None:
    plan = _m130_plan()
    book = _loaded(materialize_export_plan(plan))

    assert book.sheetnames == [tab.value for tab in plan.tabs]


def test_computed_casilla_arrives_as_a_live_formula_with_its_number_format() -> None:
    plan = _m130_plan()
    formula_cell = next(cell for cell in plan.formula_cells if cell.address.tab is TabName.CALCULOS)
    pattern = next(
        number_format.pattern for number_format in plan.number_formats if number_format.address == formula_cell.address
    )

    written = _cell_at(_loaded(materialize_export_plan(plan)), formula_cell.address)

    # A formula, not a figure: the workbook recomputes rather than restating what
    # the engine already computed, which is what makes it a live workbook.
    assert written.data_type == "f"
    assert written.value == f"={formula_cell.formula}"
    assert written.number_format == pattern


def test_seeded_money_input_arrives_as_a_number_under_its_declared_format() -> None:
    # A money input written as text would defeat both its number format and every
    # formula reading it, so the offline transport writes the Decimal itself
    # rather than the fixed-point text the online transport has to send.
    snapshot = _m303_snapshot()
    plan = build_export_plan(snapshot)
    money_addresses = {
        number_format.address for number_format in plan.number_formats if number_format.data_type == "money"
    }
    seeded_cell = next(
        cell
        for cell in plan.value_cells
        if cell.role == "operator_input" and cell.casilla_id is not None and cell.address in money_addresses
    )
    assert seeded_cell.casilla_id is not None

    seeded_plan = build_export_plan(
        snapshot,
        operator_inputs=OperatorInputs(
            values=(OperatorInput(casilla_id=seeded_cell.casilla_id, value=Decimal("1234.56")),),
        ),
    )
    written = _cell_at(_loaded(materialize_export_plan(seeded_plan)), seeded_cell.address)

    # The payload stores the exact decimal text; the reader hands numbers back as
    # floats, so the comparison goes through Decimal rather than asserting a type.
    assert written.data_type == "n"
    assert Decimal(str(written.value)) == Decimal("1234.56")
    assert written.number_format == "#,##0.00"


def test_header_band_renders_the_shared_palette_and_font() -> None:
    plan = _m130_plan()
    header = next(
        styled
        for styled in plan.styled_ranges
        if styled.role is StyleRole.HEADER and styled.tab is TabName.ENTRADAS and styled.start_row == 1
    )
    role = ROLE_STYLES[StyleRole.HEADER]
    assert role.fill_hex is not None

    book = _loaded(materialize_export_plan(plan))
    cell = _cell_at(book, SheetCellAddress.at(header.tab, header.start_row, header.start_column))

    assert cell.fill.start_color.rgb == f"FF{role.fill_hex}"
    assert cell.font.bold is True
    assert cell.font.name == WORKBOOK_FONT_FAMILY
    assert plan.font_family == WORKBOOK_FONT_FAMILY


def test_declared_frozen_view_filter_and_width_reach_the_worksheet() -> None:
    plan = _m303_plan()
    frozen = next(view for view in plan.frozen_views if view.tab is TabName.ENTRADAS)
    auto_filter = next(item for item in plan.auto_filters if item.tab is TabName.ENTRADAS)
    width = next(item for item in plan.column_widths if item.tab is TabName.ENTRADAS and item.column == 1)

    sheet = _loaded(materialize_export_plan(plan))[TabName.ENTRADAS.value]

    assert sheet.freeze_panes == f"A{frozen.frozen_rows + 1}"
    assert sheet.auto_filter.ref == f"A{auto_filter.start_row}:D{auto_filter.end_row}"
    assert sheet.column_dimensions["A"].width == width.width
    # The declared header band is repeated on every printed page.
    assert sheet.print_title_rows == f"$1:${frozen.frozen_rows}"


def test_only_the_declared_tabs_are_protected_and_inputs_stay_editable() -> None:
    plan = _m303_plan()
    declared = {region.tab for region in plan.protected_ranges}

    book = _loaded(materialize_export_plan(plan))
    protected = {tab for tab in plan.tabs if book[tab.value].protection.sheet}

    assert protected == declared
    # Entradas is the operator's own surface: the plan declares no protected
    # range there, so the transport must not invent one.
    assert TabName.ENTRADAS not in protected
    # The protection carries no password: nothing secret belongs in a filing
    # artefact, and the online protection holds no shared secret either.
    assert book[TabName.CALCULOS.value].protection.password is None


def test_constraint_reaches_its_cell_as_validation_and_legal_grounding() -> None:
    plan = _m303_plan()
    constraint = next(item for item in plan.cell_constraints if item.resolved_bounds() != (None, None))
    lower, _upper = constraint.resolved_bounds()
    assert lower is not None

    book = _loaded(materialize_export_plan(plan))
    sheet = book[constraint.address.tab.value]
    cell = _cell_at(book, constraint.address)
    validations = [
        validation for validation in sheet.data_validations.dataValidation if constraint.address.a1 in validation.sqref
    ]

    assert cell.comment is not None
    assert constraint.grounding_message() in cell.comment.text
    assert len(validations) == 1
    assert validations[0].type == "decimal"
    assert validations[0].formula1 == format(lower, "f")


def test_identity_stamps_are_carried_as_custom_document_properties() -> None:
    plan = _m130_plan()

    book = _loaded(materialize_export_plan(plan))
    properties = cast("_WorkbookCustomProperties", book).custom_doc_props
    stamps = {prop.name: prop.value for prop in properties.props}

    assert stamps == dict(export_identity_stamps(plan))
    assert set(IDENTITY_STAMP_KEYS) <= set(stamps)
    assert stamps["cadrumo_registry_sha"] == plan.metadata.registry_sha


@pytest.mark.parametrize("modelo", ["130", "303"])
def test_materializing_one_plan_twice_yields_identical_bytes(modelo: str) -> None:
    # Byte determinism is what lets a caller content-address the artefact: the
    # workbook's timestamps come from the plan, never from the clock.
    plan = build_export_plan(published_snapshot(modelo, filing_year=2025, period="1T", on=date(2025, 4, 1)))

    first = materialize_export_plan(plan)
    second = materialize_export_plan(plan)

    assert first == second


def test_export_timestamp_rather_than_the_clock_dates_the_workbook() -> None:
    plan = _m130_plan()

    book = _loaded(materialize_export_plan(plan))

    # The document properties hold whole seconds, so the plan's instant is
    # compared at that granularity rather than to the microsecond.
    exported_at = plan.metadata.exported_at.astimezone(UTC).replace(tzinfo=None, microsecond=0)
    assert book.properties.created == exported_at
    assert book.properties.modified == exported_at


def test_unrendered_plan_facets_names_only_facets_outside_the_rendered_set() -> None:
    # The live plan schema is fully rendered, and an unknown facet name is
    # reported: the judgment the materializer refuses on, exercised on its own.
    assert unrendered_plan_facets(SheetExportPlan.model_fields) == ()
    assert unrendered_plan_facets(["value_cells", "signature_block"]) == ("signature_block",)


def test_materializer_refuses_a_plan_declaring_a_facet_it_cannot_render() -> None:
    class _PlanWithAnExtraFacet(SheetExportPlan):
        """A plan whose schema grew a facet this materializer never learned."""

        signature_block: tuple[str, ...] = ()

    plan = _PlanWithAnExtraFacet(
        metadata=SheetExportMetadata(
            modelo_id="303",
            revision_id="2022",
            filing_year=2025,
            period=Period.from_year_and_code(2025, "1T"),
            engine_version="test",
            registry_sha="abcd1234",
            exported_at=datetime(2025, 4, 1, 12, 0, tzinfo=UTC),
        ),
        guide=SheetGuideContent(title="Modelo 303", paragraphs=("Use Entradas.",)),
        signature_block=("firma",),
    )

    with pytest.raises(InternalInvariantError) as raised:
        materialize_export_plan(plan)

    assert raised.value.context == {"unrendered_facets": "signature_block"}


def test_a_plan_whose_facets_are_all_rendered_is_not_refused() -> None:
    # The refusal above must not come from refusing every plan: the live schema
    # passes the same gate in the same suite.
    assert materialize_export_plan(_m130_plan())
