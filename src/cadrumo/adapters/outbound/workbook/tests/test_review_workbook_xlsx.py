"""Real workbook serialization of the shared saved-review plan."""

from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID

import pytest
from openpyxl import load_workbook

from .....application.storage.calc_sheets.review_workbook import build_review_workbook
from .....application.storage.calc_sheets.tests.review_fixture import review_label, review_snapshot
from ..calc_sheets_xlsx import materialize_export_plan

pytestmark = [pytest.mark.unit, pytest.mark.hex_outbound_adapter]


def test_review_serialization_preserves_literals_dates_values_and_editable_notes() -> None:
    plan = build_review_workbook(
        review_snapshot(), publication_id=UUID(int=2), exported_at=datetime(2026, 10, 5, tzinfo=UTC), label=review_label
    )
    payload = materialize_export_plan(plan)
    workbook = load_workbook(BytesIO(payload))
    assert workbook.sheetnames == ["Guía", "Cálculos", "Detalle", "Procedencia", "Evidencia", "Entradas"]
    assert workbook["Cálculos"]["D5"].value == 25.2
    assert workbook["Cálculos"]["E5"].value == "25.20"
    assert workbook["Detalle"]["K5"].data_type == "s"
    assert workbook["Detalle"]["K5"].value == '=IMPORTXML("https://example.invalid", "x")'
    assert workbook["Detalle"]["B5"].value == datetime(2026, 1, 12)
    assert workbook["Detalle"].freeze_panes == "B5"
    assert workbook["Entradas"].protection.sheet is False
    assert workbook["Entradas"]["B5"].fill.fgColor.rgb == "FFFFF8E1"
    assert materialize_export_plan(plan) == payload


def test_ledger_materialization_has_no_fake_modelo_metadata() -> None:
    plan = build_review_workbook(
        review_snapshot(ledger_only=True),
        publication_id=UUID(int=2),
        exported_at=datetime(2026, 10, 5, tzinfo=UTC),
        label=review_label,
    )
    workbook = load_workbook(BytesIO(materialize_export_plan(plan)))
    assert "Cálculos" not in workbook.sheetnames
    assert workbook.properties.title == "Overview"
