"""The application export entry composed with the real offline transport.

The handoff is the subject: no stand-in materializer, no injected plan builder. The
entry resolves the published authority, builds the plan, and the offline
materializer turns it into a workbook the assertions open. The case lives beside
the adapter because an application-resident test may not import one.
"""

from __future__ import annotations

import hashlib
from io import BytesIO

import pytest
from openpyxl import load_workbook

from .....application.storage.calc_sheets.records import TabName
from .....application.storage.calc_sheets.workbook_export import export_modelo_workbook
from .....core.period import Period
from ..calc_sheets_xlsx import materialize_export_plan

pytestmark = [pytest.mark.integration, pytest.mark.hex_application, pytest.mark.usefixtures("operation")]


def test_entry_returns_an_openable_workbook_and_facts_that_identify_it() -> None:
    export = export_modelo_workbook(
        modelo="130",
        period=Period.from_year_and_code(2025, "1T"),
        materializer=materialize_export_plan,
    )

    book = load_workbook(BytesIO(export.payload))

    assert book.sheetnames == list(export.tab_names)
    assert book.sheetnames == [tab.value for tab in TabName]
    assert export.modelo == "130"
    assert export.period == "1T"
    assert export.filing_year == 2025
    assert export.byte_size == len(export.payload)
    assert export.sha256 == hashlib.sha256(export.payload).hexdigest()
    # Every casilla the workbook carries is counted, and a real modelo carries some.
    assert export.casilla_count > 0
    # The identity the caller reports is the revision the authority selected.
    assert export.revision
