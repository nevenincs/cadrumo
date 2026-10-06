"""One plan, two transports, the same cells: online payload against offline workbook.

The export decision allows two transports and forbids either of them adding
structure the plan does not declare. This gate holds them to that offline, with no
Google call: the online ``values.batchUpdate`` payload is assembled in memory and
compared, cell by cell, against the ``.xlsx`` payload the offline materializer
produces from the same plan.

It lives beside the online adapter because the payload assembly it reads is that
adapter's own, and it consumes the offline materializer through its public
function. The comparison itself is the online adapter's
:func:`changed_cell_addresses`, so this gate does not invent a second notion of
"the same value" -- a ``Decimal`` sent to Sheets as fixed-point text and written
offline as a number are the same figure, and an online blank matches a cell the
offline workbook left empty.
"""

from __future__ import annotations

from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from .....application.storage.calc_sheets.engine import build_export_plan
from .....application.storage.calc_sheets.records import (
    SheetCellAddress,
    SheetExportPlan,
    SheetGuideContent,
    SheetTemplatePreviewMetadata,
    SheetValueCell,
    TabName,
)
from .....domain.calculations.registry.tests.published_authority import published_snapshot
from ...workbook.calc_sheets_xlsx import materialize_export_plan
from .._calc_sheets_apply_values import (
    build_formula_data,
    changed_cell_addresses,
    payload_written_addresses,
    written_cell_values,
)
from ..calc_sheets_apply import _developer_metadata_pairs, _plan_value_payload

pytestmark = [pytest.mark.unit, pytest.mark.hex_outbound_adapter, pytest.mark.usefixtures("operation")]

_COVERED = [("130", 2025, "1T", date(2025, 4, 1)), ("303", 2025, "1T", date(2025, 4, 1))]


def test_human_template_preview_preserves_zero_and_blank_without_technical_metadata() -> None:
    plan = SheetExportPlan[SheetTemplatePreviewMetadata](
        metadata=SheetTemplatePreviewMetadata(
            kind="template_preview",
            modelo_id="232",
            revision_id="2016-2017",
            preview_year=2017,
            preview_period="0A",
            template_digest="a" * 64,
            engine_version="test",
            title="Modelo 232 · Ejemplo ficticio 2017",
            exported_at=datetime(2026, 10, 6, tzinfo=UTC),
        ),
        human_presentation=True,
        tabs=(TabName.FORM,),
        value_cells=tuple(
            SheetValueCell(address=SheetCellAddress.at(TabName.FORM, row, 1), value=value, role="operator_input")
            for row, value in enumerate(("Ejemplo ficticio", Decimal(0), None), start=1)
        ),
        guide=SheetGuideContent(title="Ejemplo ficticio", paragraphs=("No válido para presentar.",)),
    )
    online_values = written_cell_values(_plan_value_payload(plan))
    offline_values = _offline_cell_values(materialize_export_plan(plan))
    assert changed_cell_addresses(target=online_values, current=offline_values) == ()
    assert offline_values == {"'Modelo'!A1": "Ejemplo ficticio", "'Modelo'!A2": 0}
    assert _developer_metadata_pairs(plan) == []


def _plan(modelo: str, year: int, period: str, on: date) -> SheetExportPlan:
    return build_export_plan(published_snapshot(modelo, filing_year=year, period=period, on=on))


def _offline_cell_values(payload: bytes) -> dict[str, object]:
    """Read every filled cell of an offline workbook, keyed by qualified address."""
    book = load_workbook(BytesIO(payload))
    values: dict[str, object] = {}
    for title in book.sheetnames:
        sheet = book[title]
        tab = TabName(title)
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                values[SheetCellAddress.at(tab, cell.row, cell.column).qualified()] = cell.value
    return values


@pytest.mark.parametrize(("modelo", "year", "period", "on"), _COVERED)
def test_offline_workbook_writes_exactly_the_online_payload_cells(
    modelo: str,
    year: int,
    period: str,
    on: date,
) -> None:
    plan = _plan(modelo, year, period, on)
    online_payload = _plan_value_payload(plan) + build_formula_data(plan.formula_cells)
    online_values = written_cell_values(online_payload)
    online_addresses = payload_written_addresses(online_payload)

    offline_values = _offline_cell_values(materialize_export_plan(plan))

    # Nothing the online transport would write is missing or different offline.
    assert changed_cell_addresses(target=online_values, current=offline_values) == ()
    # And nothing the plan does not declare appears offline: no extra tab, no
    # extra metadata row, no cell the online workbook would not have.
    assert set(offline_values) <= online_addresses


@pytest.mark.parametrize(("modelo", "year", "period", "on"), _COVERED)
def test_both_transports_carry_the_same_formula_at_the_same_address(
    modelo: str,
    year: int,
    period: str,
    on: date,
) -> None:
    plan = _plan(modelo, year, period, on)
    online_formulas = written_cell_values(build_formula_data(plan.formula_cells))

    offline_values = _offline_cell_values(materialize_export_plan(plan))
    offline_formulas = {
        address: value for address, value in offline_values.items() if isinstance(value, str) and value.startswith("=")
    }

    assert online_formulas
    assert offline_formulas == online_formulas


@pytest.mark.parametrize(("modelo", "year", "period", "on"), _COVERED)
def test_both_transports_create_the_same_tabs(modelo: str, year: int, period: str, on: date) -> None:
    plan = _plan(modelo, year, period, on)

    book = load_workbook(BytesIO(materialize_export_plan(plan)))

    # The online adapter creates its spreadsheet with exactly this tab list.
    assert book.sheetnames == [tab.value for tab in plan.tabs]
