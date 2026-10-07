"""Form geometry survives both workbook transports; destructive merges fail."""

from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from uuid import UUID

import pytest
from openpyxl import load_workbook

from .....application.storage.calc_sheets.number_formats import numeric_format
from .....application.storage.calc_sheets.records import (
    SheetCellAddress,
    SheetColumnWidth,
    SheetExportMetadata,
    SheetExportPlan,
    SheetFormulaCell,
    SheetGuideContent,
    SheetHiddenRow,
    SheetMergedRange,
    SheetNumberFormat,
    SheetProtectedRange,
    SheetReviewMetadata,
    SheetRoundingRule,
    SheetRowHeight,
    SheetStyledRange,
    SheetValueCell,
    TabName,
)
from .....application.storage.calc_sheets.theme import StyleRole
from .....core.period import Period
from ...workbook.calc_sheets_xlsx import materialize_export_plan
from .._calc_sheets_apply_formatting import (
    build_base_font_requests,
    build_form_geometry_requests,
    build_grid_resize_requests,
    build_number_format_requests,
    build_styled_range_requests,
)
from .._calc_sheets_apply_values import build_value_data

pytestmark = [pytest.mark.unit, pytest.mark.hex_outbound_adapter]


def _plan() -> SheetExportPlan[SheetReviewMetadata]:
    return SheetExportPlan[SheetReviewMetadata](
        metadata=SheetReviewMetadata(
            kind="calculation",
            snapshot_digest="a" * 64,
            publication_id=UUID(int=1),
            title="Ejemplo",
            exported_at=datetime(2026, 10, 5, tzinfo=UTC),
        ),
        tabs=(TabName.FORM, TabName.GUIDE),
        guide=SheetGuideContent(title="Guía", paragraphs=("Ejemplo de diseño.",)),
        value_cells=(
            SheetValueCell(address=SheetCellAddress.at(TabName.FORM, 2, 2), value="Liquidación", role="label"),
        ),
        merged_ranges=(SheetMergedRange(tab=TabName.FORM, start_row=2, end_row=2, start_column=2, end_column=12),),
        row_heights=(SheetRowHeight(tab=TabName.FORM, row=2, height_pixels=48),),
        styled_ranges=(
            SheetStyledRange(
                tab=TabName.FORM,
                start_row=2,
                end_row=2,
                start_column=2,
                end_column=12,
                role=StyleRole.FORM_SECTION,
                boxed=True,
                wrap=True,
            ),
        ),
    )


def test_both_transports_preserve_form_geometry() -> None:
    plan = _plan()
    book = load_workbook(BytesIO(materialize_export_plan(plan)))
    form = book["Modelo"]
    assert str(next(iter(form.merged_cells.ranges))) == "B2:L2"
    assert form["B2"].value == "Liquidación"
    assert form.row_dimensions[2].height == 36
    assert form.sheet_view.showGridLines is False
    assert form["B2"].border.top.style == "thin"
    requests = build_form_geometry_requests(plan, sheet_id_by_tab={"Modelo": 42})
    assert requests[0]["mergeCells"]["range"] == {
        "sheetId": 42,
        "startRowIndex": 1,
        "endRowIndex": 2,
        "startColumnIndex": 1,
        "endColumnIndex": 12,
    }
    assert requests[1]["updateDimensionProperties"]["properties"]["pixelSize"] == 48
    assert requests[2]["updateSheetProperties"]["properties"]["gridProperties"]["hideGridlines"] is True
    styled = build_styled_range_requests(plan, sheet_id_by_tab={"Modelo": 42})
    assert styled[0]["repeatCell"]["cell"]["userEnteredFormat"]["borders"]["top"]["style"] == "SOLID"


def test_merge_must_not_discard_another_value() -> None:
    plan = _plan()
    data = dict(plan)
    data["value_cells"] = (
        *plan.value_cells,
        SheetValueCell(address=SheetCellAddress.at(TabName.FORM, 2, 3), value="Do not lose", role="label"),
    )
    with pytest.raises(ValueError, match="discard"):
        SheetExportPlan[SheetReviewMetadata].model_validate(data)


def test_internal_row_visibility_preserves_formula_dependencies_in_both_transports() -> None:
    data = dict(_plan())
    data["tabs"] = (*data["tabs"], TabName.CALCULOS)
    data["hidden_rows"] = (SheetHiddenRow(tab=TabName.CALCULOS, row=3),)
    data["formula_cells"] = (
        SheetFormulaCell(
            address=SheetCellAddress.at(TabName.CALCULOS, 3, 4),
            casilla_id="test.control",
            formula="1",
            rounding_rule=SheetRoundingRule.NONE,
        ),
        SheetFormulaCell(
            address=SheetCellAddress.at(TabName.FORM, 4, 2),
            casilla_id="test.resultado",
            formula="'Cálculos'!D3+1",
            rounding_rule=SheetRoundingRule.NONE,
        ),
    )
    plan = SheetExportPlan[SheetReviewMetadata].model_validate(data)
    restored = SheetExportPlan[SheetReviewMetadata].model_validate_json(plan.model_dump_json())
    assert restored.hidden_rows == plan.hidden_rows
    book = load_workbook(BytesIO(materialize_export_plan(restored)))
    assert book["Cálculos"].row_dimensions[3].hidden
    assert not book["Modelo"].row_dimensions[4].hidden
    assert book["Cálculos"]["D3"].value == "=1"
    assert book["Modelo"]["B4"].value == "='Cálculos'!D3+1"
    requests = build_form_geometry_requests(restored, sheet_id_by_tab={"Modelo": 42, "Cálculos": 43})
    assert {
        "updateDimensionProperties": {
            "range": {"sheetId": 43, "dimension": "ROWS", "startIndex": 2, "endIndex": 3},
            "properties": {"hiddenByUser": True},
            "fields": "hiddenByUser",
        }
    } in requests
    book.close()


def test_hidden_row_cannot_target_an_undeclared_tab() -> None:
    data = dict(_plan())
    data["hidden_rows"] = (SheetHiddenRow(tab=TabName.CALCULOS, row=3),)
    with pytest.raises(ValueError, match="hidden row targets an undeclared tab"):
        SheetExportPlan[SheetReviewMetadata].model_validate(data)


def test_overlapping_merge_is_refused() -> None:
    plan = _plan()
    data = dict(plan)
    data["merged_ranges"] = (
        *plan.merged_ranges,
        SheetMergedRange(tab=TabName.FORM, start_row=2, end_row=3, start_column=10, end_column=12),
    )
    with pytest.raises(ValueError, match="overlap"):
        SheetExportPlan[SheetReviewMetadata].model_validate(data)


def test_geometry_alone_extends_the_native_grid() -> None:
    data = dict(_plan())
    data["merged_ranges"] = (
        SheetMergedRange(tab=TabName.FORM, start_row=1100, end_row=1101, start_column=1, end_column=30),
    )
    plan = SheetExportPlan[SheetReviewMetadata].model_validate(data)
    requests = build_grid_resize_requests(plan, sheet_id_by_tab={"Modelo": 42})
    grid = requests[0]["updateSheetProperties"]["properties"]["gridProperties"]
    assert grid["rowCount"] >= 1101
    assert grid["columnCount"] >= 30


def test_form_grid_trims_only_beyond_declared_padding() -> None:
    data = dict(_plan())
    data["column_widths"] = (SheetColumnWidth(tab=TabName.FORM, column=13, width=3),)
    data["protected_ranges"] = (
        SheetProtectedRange(
            tab=TabName.FORM,
            start_row=1,
            end_row=8,
            start_column=1,
            end_column=13,
            description="Includes empty outer padding",
        ),
    )
    plan = SheetExportPlan[SheetReviewMetadata].model_validate(data)
    requests = build_grid_resize_requests(plan, sheet_id_by_tab={"Modelo": 42, "Guía": 43})
    grids = {
        r["updateSheetProperties"]["properties"]["sheetId"]: r["updateSheetProperties"]["properties"]["gridProperties"]
        for r in requests
    }
    assert grids[42] == {"rowCount": 8, "columnCount": 13}
    assert grids[43]["rowCount"] >= 1000
    assert grids[43]["columnCount"] >= 26


def test_form_grid_keeps_blank_dimensions_outside_content() -> None:
    data = dict(_plan())
    data["column_widths"] = (SheetColumnWidth(tab=TabName.FORM, column=17, width=3),)
    data["row_heights"] = (SheetRowHeight(tab=TabName.FORM, row=21, height_pixels=12),)
    plan = SheetExportPlan[SheetReviewMetadata].model_validate(data)
    request = build_grid_resize_requests(plan, sheet_id_by_tab={"Modelo": 42})[0]
    assert request["updateSheetProperties"]["properties"]["gridProperties"] == {"rowCount": 21, "columnCount": 17}


def test_derived_guide_content_cannot_be_erased_by_a_merge() -> None:
    data = dict(_plan())
    data["metadata"] = SheetExportMetadata(
        modelo_id="130",
        revision_id="2019-y-siguientes",
        filing_year=2025,
        period=Period.from_year_and_code(2025, "4T"),
        engine_version="test",
        registry_sha="a" * 64,
        exported_at=datetime(2026, 10, 5, tzinfo=UTC),
    )
    data["merged_ranges"] = (SheetMergedRange(tab=TabName.GUIDE, start_row=1, end_row=3, start_column=1, end_column=2),)
    plan = SheetExportPlan.model_validate(data)
    with pytest.raises(ValueError, match="discard"):
        materialize_export_plan(plan)
    with pytest.raises(ValueError, match="discard"):
        build_form_geometry_requests(plan, sheet_id_by_tab={"Guía": 7})


def test_numeric_looking_box_label_remains_text_in_both_transports() -> None:
    data = dict(_plan())
    address = SheetCellAddress.at(TabName.FORM, 3, 9)
    data["value_cells"] = (*data["value_cells"], SheetValueCell(address=address, value="01", role="label"))
    data["number_formats"] = (SheetNumberFormat(address=address, data_type="text", pattern="@"),)
    plan = SheetExportPlan[SheetReviewMetadata].model_validate(data)
    cell = load_workbook(BytesIO(materialize_export_plan(plan)))["Modelo"]["I3"]
    assert cell.value == "01"
    assert cell.data_type == "s"
    assert cell.number_format == "@"
    request = build_number_format_requests(plan, sheet_id_by_tab={"Modelo": 42})[0]
    assert request["repeatCell"]["cell"]["userEnteredFormat"]["numberFormat"] == {"type": "TEXT", "pattern": "@"}


@pytest.mark.parametrize(
    ("role", "horizontal", "bold"),
    [
        (StyleRole.CASILLA, "center", True),
        (StyleRole.FORM_INPUT, "right", False),
        (StyleRole.FORM_COMPUTED, "right", False),
        (StyleRole.FORM_RESULT, "right", True),
    ],
)
def test_box_styles_match_between_transports(role, horizontal, bold) -> None:
    plan = _plan()
    plan = plan.model_copy(update={"styled_ranges": (plan.styled_ranges[0].model_copy(update={"role": role}),)})
    cell = load_workbook(BytesIO(materialize_export_plan(plan)))["Modelo"]["B2"]
    assert cell.alignment.vertical == "center"
    assert cell.alignment.horizontal == horizontal
    assert cell.font.bold is bold
    assert all(getattr(cell.border, edge).style == "thin" for edge in ("top", "bottom", "left", "right"))
    native = build_styled_range_requests(plan, sheet_id_by_tab={"Modelo": 42})[0]["repeatCell"]["cell"][
        "userEnteredFormat"
    ]
    assert native["verticalAlignment"] == "MIDDLE"
    assert native["horizontalAlignment"] == horizontal.upper()
    assert native["textFormat"]["bold"] is bold
    assert all(native["borders"][edge]["style"] == "SOLID" for edge in ("top", "bottom", "left", "right"))


@pytest.mark.parametrize(
    ("kind", "currency", "pattern", "native_pattern"),
    [
        ("money", "EUR", '#,##0.00" €"', '#,##0.00" €"'),
        ("money", "USD", "#,##0.00", "#,##0.00"),
        ("money", None, "#,##0.00", "#,##0.00"),
        ("integer", None, "#,##0", "#,##0"),
        ("decimal", None, "#,##0.############", "#,##0.0###########"),
        ("float", None, "#,##0.############", "#,##0.0###########"),
        ("ratio", None, "0.00####", "0.00####"),
        ("percentage", None, "0.00####%", "0.00####%"),
    ],
)
@pytest.mark.parametrize(
    "value",
    [Decimal("1234.56"), Decimal("-1234.56"), Decimal(0), Decimal(1234), Decimal("0.123456789012"), None],
)
def test_spanish_numeric_formats_preserve_values_and_missing_cells(kind, currency, pattern, native_pattern, value) -> None:
    declared = numeric_format(kind, currency=currency)
    assert declared is not None
    assert declared[1] == pattern
    address = SheetCellAddress.at(TabName.FORM, 3, 2)
    plan = _plan().model_copy(
        update={
            "value_cells": (SheetValueCell(address=address, value=value, role="operator_input"),),
            "number_formats": (SheetNumberFormat(address=address, data_type=declared[0], pattern=declared[1]),),
        }
    )
    cell = load_workbook(BytesIO(materialize_export_plan(plan)))["Modelo"]["B3"]
    assert cell.value is None if value is None else Decimal(str(cell.value)) == value
    assert cell.number_format == "[$-C0A]" + pattern
    native = build_number_format_requests(plan, sheet_id_by_tab={"Modelo": 42})[0]["repeatCell"]["cell"][
        "userEnteredFormat"
    ]
    assert native["numberFormat"] == {
        "type": "PERCENT" if kind == "percentage" else "NUMBER",
        "pattern": native_pattern,
    }
    assert build_value_data(plan.value_cells)[0]["values"] == [["" if value is None else float(value)]]
    assert build_base_font_requests(plan, sheet_id_by_tab={"Modelo": 42})[0] == {
        "updateSpreadsheetProperties": {"properties": {"locale": "es_ES"}, "fields": "locale"}
    }


def test_title_size_is_shared_between_transports() -> None:
    plan = _plan()
    plan = plan.model_copy(
        update={"styled_ranges": (plan.styled_ranges[0].model_copy(update={"role": StyleRole.TITLE}),)}
    )
    cell = load_workbook(BytesIO(materialize_export_plan(plan)))["Modelo"]["B2"]
    assert cell.font.sz == 16
    native = build_styled_range_requests(plan, sheet_id_by_tab={"Modelo": 42})[0]["repeatCell"]["cell"][
        "userEnteredFormat"
    ]
    assert native["textFormat"]["fontSize"] == 16
