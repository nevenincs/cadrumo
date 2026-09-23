"""Unit tests for the legacy ``.xls`` financial ingestion provider.

The parity test copies every cell value of the synthetic ``.xlsx`` corpus
statement into a real BIFF8 workbook and requires the ``.xls`` provider to
produce the same parsed rows the independently implemented ``.xlsx`` provider
produces from the original.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ......domain.transactions.enums import TransactionDirection
from ......domain.transactions.raw_transaction import SourceFormat
from ......tests.inventory import FIXTURES_DIR
from ......tests.xls_fixtures import XlsFormula, xls_workbook_bytes
from ..base import InvalidFinancialSourceError, ParsedLedgerRow
from ..detection import detect_provider
from ..xls import XlsProvider
from ..xlsx import XlsxProvider

pytestmark = [pytest.mark.unit, pytest.mark.hex_inbound_adapter]

_XLSX_FIXTURE = FIXTURES_DIR / "financial" / "synthetic-transactions.xlsx"
_N26_HEADER = ["Date", "Payee", "Payment reference", "Amount (EUR)", "Currency", "Transaction ID"]


def _xlsx_cell_values(path: Path) -> list[tuple[str, list[list[object]]]]:
    workbook = load_workbook(filename=path, read_only=True, data_only=False)
    try:
        return [
            (worksheet.title, [list(row) for row in worksheet.iter_rows(values_only=True)])
            for worksheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _comparable(row: ParsedLedgerRow) -> tuple[object, ...]:
    raw = row.raw
    return (
        raw.booked_date,
        raw.value_date,
        raw.amount,
        raw.currency,
        raw.counterparty,
        raw.description,
        row.direction,
        raw.provenance.source_row_index,
    )


def test_xls_statement_parses_exactly_like_the_same_xlsx_statement(tmp_path: Path) -> None:
    legacy = tmp_path / "synthetic-transactions.xls"
    legacy.write_bytes(xls_workbook_bytes(_xlsx_cell_values(_XLSX_FIXTURE)))

    xlsx_rows = tuple(XlsxProvider().ingest(_XLSX_FIXTURE))
    xls_rows = tuple(XlsProvider().ingest(legacy))

    assert xlsx_rows
    assert [_comparable(row) for row in xls_rows] == [_comparable(row) for row in xlsx_rows]
    assert {row.raw.provenance.source_format for row in xls_rows} == {SourceFormat.XLS}
    assert {row.raw.provenance.source_path.name for row in xls_rows} == {legacy.name}


def test_xls_statement_reads_typed_date_and_number_cells_with_worksheet_row_provenance(tmp_path: Path) -> None:
    source = tmp_path / "statement.xls"
    source.write_bytes(
        xls_workbook_bytes(
            [
                ("Notes", [["exported for synthetic testing"]]),
                (
                    "Statement",
                    [
                        ["Synthetic bank statement"],
                        _N26_HEADER,
                        [date(2025, 3, 11), "Synthetic payee", "first", -12.5, "EUR", "xls-1"],
                        [None],
                        [date(2025, 3, 12), "Synthetic payer", "second", 30, "EUR", "xls-2"],
                    ],
                ),
            ]
        )
    )
    provider = XlsProvider()

    validation = provider.validate_source(source)
    rows = tuple(provider.ingest(source))

    assert validation.is_valid, validation.warnings
    assert validation.detected_dialect == "worksheet=Statement,header_row=2"
    assert [(row.raw.booked_date, row.raw.amount, row.direction) for row in rows] == [
        (date(2025, 3, 11), Decimal("12.5"), TransactionDirection.OUTGOING),
        (date(2025, 3, 12), Decimal("30"), TransactionDirection.INCOMING),
    ]
    assert [row.raw.provenance.source_row_index for row in rows] == [3, 5]


def test_xls_formula_cell_below_the_header_refuses_the_source(tmp_path: Path) -> None:
    source = tmp_path / "formula.xls"
    source.write_bytes(
        xls_workbook_bytes(
            [("Statement", [_N26_HEADER, ["2025-03-11", "Synthetic payee", "first", XlsFormula(9000), "EUR", "x"]])]
        )
    )

    validation = XlsProvider().validate_source(source)

    assert not validation.is_valid
    assert "formula cell at row 2, column 4" in validation.warnings[0]
    with pytest.raises(InvalidFinancialSourceError, match="formula cached values are not accepted"):
        tuple(XlsProvider().ingest(source))


def test_xls_formula_cell_above_the_header_is_outside_the_statement(tmp_path: Path) -> None:
    source = tmp_path / "preamble-formula.xls"
    source.write_bytes(
        xls_workbook_bytes(
            [
                (
                    "Statement",
                    [[XlsFormula(1)], _N26_HEADER, ["2025-03-11", "Synthetic payee", "first", "10.00", "EUR", "x"]],
                )
            ]
        )
    )

    rows = tuple(XlsProvider().ingest(source))

    assert [row.raw.provenance.source_row_index for row in rows] == [3]


def test_xls_without_a_bank_layout_or_readable_workbook_is_invalid(tmp_path: Path) -> None:
    unrelated = tmp_path / "unrelated.xls"
    unrelated.write_bytes(xls_workbook_bytes([("Sheet1", [["alpha", "beta"], ["1", "2"]])]))
    corrupt = tmp_path / "corrupt.xls"
    corrupt.write_bytes(b"not a workbook")

    assert "supported bank-statement header row" in XlsProvider().validate_source(unrelated).warnings[0]
    assert "could not open legacy workbook" in XlsProvider().validate_source(corrupt).warnings[0]


def test_auto_detection_reads_an_xls_statement_by_suffix_and_by_signature(tmp_path: Path) -> None:
    workbook = xls_workbook_bytes([("Statement", [_N26_HEADER, ["2025-03-11", "Payee", "ref", "10.00", "EUR", "a"]])])
    named = tmp_path / "statement.xls"
    unnamed = tmp_path / "statement.download"
    named.write_bytes(workbook)
    unnamed.write_bytes(workbook)

    assert isinstance(detect_provider(named), XlsProvider)
    assert isinstance(detect_provider(unnamed), XlsProvider)
